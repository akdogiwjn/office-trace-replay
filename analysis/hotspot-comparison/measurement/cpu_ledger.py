#!/usr/bin/env python3
"""CPU-time interval ledger for openpyxl load / lazy-read / save.

Activated only when the environment variable CPU_LEDGER_OUT points to a
writable directory (the measurement launcher sets it; normal runs never do).
When active, every Python process that imports this module through
``sitecustomize`` installs transparent wrappers around:

  * ``openpyxl.load_workbook``          -> ledger bucket "load_workbook"
  * ``Workbook.save``                   -> ledger bucket "workbook_save"
  * read-only worksheet ``_cells_by_row`` -> "readonly_lazy_parse", covering
    iter_rows, rows/values, slices and direct cell access

Interval accounting rules:

  * CPU time is ``time.thread_time_ns()`` (user+system of the executing
    thread), accumulated over threads. Other threads' unrelated work and
    wall-clock blocking never accrue to an operation.
  * The read-only bucket measures the *generator's own* CPU: the clock runs
    while the generator is executing (XML parse/row construction) and stops
    while it is suspended at a yield, so consumer-side business logic between
    iterations is excluded.  Business code that never pulls a row is not
    counted as read time at all.
  * Nested instrumented intervals de-duplicate via an explicit stack: while an
    inner interval runs, the outer interval's accumulator is paused, so CPU is
    attributed to exactly one bucket at any moment.
  * Wrappers never alter arguments, return values or exceptions, so business
    behaviour is unchanged; the ledger records its own module hash so a run
    can prove which hook version was actually loaded.

On interpreter exit (atexit) each process writes one JSON ledger file named
``ledger-<pid>.json`` into $CPU_LEDGER_OUT.  Processes that never touch
openpyxl still write an identity-only ledger, which is how a run proves the
hook was loaded process-wide.

Known limits (kept honest in reports):

  * The hook adds interval bookkeeping per row and at generator cleanup;
    the validation program estimates its cost on a synthetic generator.
  * Only Python-side openpyxl CPU is measured; LibreOffice's own store CPU is
    intentionally outside the numerator (it belongs to the denominator).
"""
import atexit
import hashlib
import json
import os
import sys
import threading
import time
from pathlib import Path

BUCKETS = ('load_workbook', 'workbook_save', 'readonly_lazy_parse')
LEDGER_OUT = os.environ.get('CPU_LEDGER_OUT')


class IntervalLedger:
    """Accrue exclusive thread CPU with independent stacks for each thread."""

    def __init__(self):
        self.cpu_ns = {bucket: 0 for bucket in BUCKETS}
        self.calls = {bucket: 0 for bucket in BUCKETS}
        self.yields = 0
        self._threads = threading.local()
        self._lock = threading.Lock()

    def _stack(self):
        if not hasattr(self._threads, 'stack'):
            self._threads.stack = []
        return self._threads.stack

    def _now(self):
        return time.thread_time_ns()

    def enter(self, bucket):
        now = self._now()
        stack = self._stack()
        if stack:
            stack[-1][1] += now - stack[-1][2]
        stack.append([bucket, 0, now])

    def exit(self, bucket):
        now = self._now()
        stack = self._stack()
        if not stack or stack[-1][0] != bucket:
            raise RuntimeError('CPU ledger interval stack mismatch: ' + bucket)
        frame = stack.pop()
        elapsed = frame[1] + now - frame[2]
        with self._lock:
            self.cpu_ns[bucket] += elapsed
            self.calls[bucket] += 1
        if stack:
            # Exclude ledger aggregation overhead from the resumed parent.
            stack[-1][2] = self._now()
        return elapsed

    def note_yield(self):
        with self._lock:
            self.yields += 1


LEDGER = IntervalLedger()


def _wrap_load_workbook():
    import openpyxl
    from openpyxl.reader import excel
    original = openpyxl.load_workbook

    def timed_load_workbook(*args, **kwargs):
        LEDGER.enter('load_workbook')
        try:
            return original(*args, **kwargs)
        finally:
            LEDGER.exit('load_workbook')

    timed_load_workbook.__wrapped__ = original
    openpyxl.load_workbook = timed_load_workbook
    excel.load_workbook = timed_load_workbook


def _wrap_workbook_save():
    from openpyxl.workbook.workbook import Workbook
    original = Workbook.save

    def timed_save(self, *args, **kwargs):
        LEDGER.enter('workbook_save')
        try:
            return original(self, *args, **kwargs)
        finally:
            LEDGER.exit('workbook_save')

    timed_save.__wrapped__ = original
    Workbook.save = timed_save


def _timed_rows(generator):
    """Time each resume/cleanup, leaving no interval open across a yield."""
    try:
        while True:
            LEDGER.enter('readonly_lazy_parse')
            try:
                row = next(generator)
            except StopIteration:
                return
            finally:
                LEDGER.exit('readonly_lazy_parse')
            LEDGER.note_yield()
            yield row
    finally:
        # Early break, direct _get_cell and explicit close must release the
        # XML stream, without charging the consumer's intervening CPU time.
        LEDGER.enter('readonly_lazy_parse')
        try:
            generator.close()
        finally:
            LEDGER.exit('readonly_lazy_parse')


def _wrap_readonly_rows():
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    original = ReadOnlyWorksheet._cells_by_row

    def timed_cells_by_row(self, *args, **kwargs):
        return _timed_rows(original(self, *args, **kwargs))

    timed_cells_by_row.__wrapped__ = original
    ReadOnlyWorksheet._cells_by_row = timed_cells_by_row


def install():
    """Install all wrappers; safe to call only once per process."""
    if getattr(install, '_installed', False):
        return
    install._installed = True
    _wrap_load_workbook()
    _wrap_workbook_save()
    _wrap_readonly_rows()


def _module_identity():
    digest = hashlib.sha256()
    with open(__file__, 'rb') as stream:
        for block in iter(lambda: stream.read(65536), b''):
            digest.update(block)
    return {'module': __file__, 'module_sha256': digest.hexdigest()}


def _write_ledger():
    if not LEDGER_OUT:
        return
    try:
        out = Path(LEDGER_OUT)
        out.mkdir(parents=True, exist_ok=True)
        payload = {
            'schema_version': 'cpu-ledger-v2',
            'operation_cpu_clock': 'time.thread_time_ns (user+system, summed across threads)',
            'lazy_read_boundary': 'ReadOnlyWorksheet._cells_by_row resume and close; consumer excluded',
            'pid': os.getpid(),
            'cmdline': sys.argv[:],
            'identity': _module_identity(),
            'cpu_total_process_time_ns': time.process_time_ns(),
            'buckets': {bucket: {'cpu_ns': LEDGER.cpu_ns[bucket],
                                 'calls': LEDGER.calls[bucket]}
                        for bucket in BUCKETS},
            'readonly_rows_yielded': LEDGER.yields,
        }
        (out / f'ledger-{os.getpid()}.json').write_text(
            json.dumps(payload, indent=1) + '\n', encoding='utf-8')
    except Exception:  # pragma: no cover - never break the measured process
        pass


if LEDGER_OUT:
    install()
    atexit.register(_write_ledger)
