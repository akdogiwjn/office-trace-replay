"""Small synthetic validation of the CPU ledger measurement semantics.

These tests run the real ledger code on small synthetic workbooks inside
temporary directories.  They verify measurement semantics only (CPU accrual,
sleep exclusion, lazy-read/consumer separation, nested de-duplication,
sitecustomize injection, business-output neutrality) and produce no
performance claims about real workloads.
"""
import json
import re
import subprocess
import sys
import tempfile
import unittest
import zipfile
import hashlib
import importlib.util
import threading
from unittest.mock import patch
from pathlib import Path

MEASUREMENT_DIR = Path(__file__).resolve().parents[1] / 'measurement'


class LedgerRegressionTests(unittest.TestCase):
    """Deterministic CPU clocks catch accounting errors without busy waits."""

    def setUp(self):
        spec = importlib.util.spec_from_file_location(
            'cpu_ledger_regression_fixture', MEASUREMENT_DIR / 'cpu_ledger.py')
        self.module = importlib.util.module_from_spec(spec)
        with patch.dict('os.environ', {}, clear=True):
            spec.loader.exec_module(self.module)
        self.ledger = self.module.LEDGER
        self.clock = 0
        self.ledger._now = lambda: self.clock

    def rows(self, costs, cleanup=0, fail=False):
        try:
            for index, cost in enumerate(costs):
                self.clock += cost
                yield index
            if fail:
                raise ValueError('parser failure')
        finally:
            self.clock += cleanup

    def test_early_close_excludes_consumer_cpu_and_counts_cleanup(self):
        rows = self.module._timed_rows(self.rows([3, 9], cleanup=2))
        self.assertEqual(next(rows), 0)
        self.assertEqual(self.ledger._stack(), [])
        self.clock += 1_000_000_000  # consumer CPU, rather than a sleep
        rows.close()
        self.assertEqual(self.ledger.cpu_ns['readonly_lazy_parse'], 5)
        self.assertEqual(self.ledger._stack(), [])

    def test_exhaustion_excludes_cpu_between_rows(self):
        rows = self.module._timed_rows(self.rows([3, 4], cleanup=2))
        self.assertEqual(next(rows), 0)
        self.clock += 1_000_000_000
        self.assertEqual(next(rows), 1)
        self.clock += 1_000_000_000
        self.assertEqual(list(rows), [])
        self.assertEqual(self.ledger.cpu_ns['readonly_lazy_parse'], 9)
        self.assertEqual(self.ledger.yields, 2)

    def test_interleaved_generators_and_consumer_save_do_not_overlap(self):
        left = self.module._timed_rows(self.rows([3, 4], cleanup=1))
        right = self.module._timed_rows(self.rows([5], cleanup=2))
        next(left)
        self.clock += 100
        next(right)
        self.ledger.enter('workbook_save')
        self.clock += 7
        self.ledger.exit('workbook_save')
        self.clock += 100
        next(left)
        right.close()
        left.close()
        self.assertEqual(self.ledger.cpu_ns['readonly_lazy_parse'], 15)
        self.assertEqual(self.ledger.cpu_ns['workbook_save'], 7)
        self.assertEqual(self.ledger._stack(), [])

    def test_parser_exception_closes_interval_and_preserves_exception(self):
        rows = self.module._timed_rows(self.rows([3], cleanup=2, fail=True))
        next(rows)
        self.clock += 100
        with self.assertRaisesRegex(ValueError, 'parser failure'):
            next(rows)
        self.assertEqual(self.ledger.cpu_ns['readonly_lazy_parse'], 5)
        self.assertEqual(self.ledger._stack(), [])

    def test_nested_operations_count_each_cpu_unit_once(self):
        self.ledger.enter('load_workbook')
        self.clock += 2
        self.ledger.enter('readonly_lazy_parse')
        self.clock += 5
        self.ledger.exit('readonly_lazy_parse')
        self.clock += 3
        self.ledger.exit('load_workbook')
        self.assertEqual(self.ledger.cpu_ns['load_workbook'], 5)
        self.assertEqual(self.ledger.cpu_ns['readonly_lazy_parse'], 5)
        self.assertEqual(sum(self.ledger.cpu_ns.values()), 10)

    def test_threads_have_independent_interval_stacks(self):
        clock = threading.local()
        self.ledger._now = lambda: clock.value
        barrier = threading.Barrier(2)
        errors = []

        def worker(bucket, cost):
            try:
                clock.value = 0
                self.ledger.enter(bucket)
                barrier.wait(timeout=5)
                clock.value += cost
                self.ledger.exit(bucket)
            except BaseException as exc:
                errors.append(exc)

        threads = [threading.Thread(target=worker, args=args) for args in
                   [('load_workbook', 7), ('workbook_save', 11)]]
        for thread in threads:
            thread.start()
        for thread in threads:
            thread.join(timeout=6)
        self.assertFalse(any(thread.is_alive() for thread in threads))
        self.assertEqual(errors, [])
        self.assertEqual(self.ledger.cpu_ns['load_workbook'], 7)
        self.assertEqual(self.ledger.cpu_ns['workbook_save'], 11)

    def test_readonly_cell_slice_rows_and_values_are_all_instrumented(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'small.xlsx'
            book = Workbook()
            book.active.append([1, 2])
            book.active.append([3, 4])
            book.save(path)
            book.close()
            # Restore the actual class after this fixture; don't leak hooks.
            with patch.object(ReadOnlyWorksheet, '_cells_by_row',
                              ReadOnlyWorksheet._cells_by_row):
                self.module._wrap_readonly_rows()
                for read in (lambda ws: ws['A2'].value,
                             lambda ws: [[c.value for c in r] for r in ws['A1:B2']],
                             lambda ws: list(ws.iter_rows(values_only=True)),
                             lambda ws: list(ws.rows),
                             lambda ws: list(ws.values)):
                    before = self.ledger.calls['readonly_lazy_parse']
                    book = load_workbook(path, read_only=True)
                    try:
                        read(book.active)
                    finally:
                        book.close()
                    self.assertGreater(self.ledger.calls['readonly_lazy_parse'], before)
                    self.assertEqual(self.ledger._stack(), [])


class CpuLedgerSemanticsTests(unittest.TestCase):
    def test_validation_program_passes_all_checks(self):
        with tempfile.TemporaryDirectory() as tmp:
            environment = dict(__import__('os').environ,
                               CPU_LEDGER_OUT=str(Path(tmp) / 'ledgers'))
            completed = subprocess.run(
                [sys.executable, str(MEASUREMENT_DIR / 'validate_cpu_ledger.py'), '--json'],
                capture_output=True, text=True, env=environment, timeout=300)
            results = json.loads(completed.stdout)
        self.assertEqual(completed.returncode, 0, results)
        for key in ('check_cpu_accrues', 'check_sleep_excluded',
                    'check_consumer_excluded', 'check_rows_counted',
                    'check_subprocess_cpu_visible', 'check_nesting_dedup'):
            self.assertTrue(results[key], key)
        # the overhead stays quantified and small relative to real loops
        self.assertLess(results['hook_overhead_ns_per_yield'], 10_000)


class SitecustomizeInjectionTests(unittest.TestCase):
    BUSINESS = '''"""Synthetic business program: load, lazy-read, save."""
import sys, hashlib, json
from pathlib import Path
from openpyxl import Workbook, load_workbook
out = Path(sys.argv[1])
rows = [(i, i * 3, f'text-{i}') for i in range(3000)]
book = Workbook()
sheet = book.active
sheet.title = 'S'
for row in rows:
    sheet.append(row)
book.save(out)
book.close()
book = load_workbook(out, read_only=True)
try:
    total = sum((r[0].value or 0) + (r[1].value or 0) for r in book['S'].iter_rows())
finally:
    book.close()
print(json.dumps({'rows_total': total}))
'''

    def _business_fingerprint(self, path):
        with zipfile.ZipFile(path) as archive:
            parts = {}
            for name in archive.namelist():
                data = archive.read(name)
                if name.startswith('docProps/'):
                    data = re.sub(rb'\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z',
                                  b'TIMESTAMP', data)
                parts[name] = hashlib.sha256(data).hexdigest()
        return parts

    def test_hook_loads_via_sitecustomize_and_leaves_business_content_identical(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            hook_dir = root / 'hook'
            hook_dir.mkdir()
            for name in ('cpu_ledger.py', 'sitecustomize.py'):
                (hook_dir / name).write_bytes((MEASUREMENT_DIR / name).read_bytes())
            program = root / 'business.py'
            program.write_text(self.BUSINESS)
            ledger_out = root / 'ledgers'
            baseline = root / 'baseline.xlsx'
            measured = root / 'measured.xlsx'
            run = lambda out, env: subprocess.run(
                [sys.executable, str(program), str(out)],
                capture_output=True, text=True, env=env, check=True)
            import os
            run(baseline, dict(os.environ))
            hooked_env = dict(os.environ,
                              CPU_LEDGER_OUT=str(ledger_out),
                              PYTHONPATH=str(hook_dir))
            run(measured, hooked_env)
            ledgers = sorted(ledger_out.glob('ledger-*.json'))
            self.assertEqual(len(ledgers), 1, 'exactly the business process wrote a ledger')
            ledger = json.loads(ledgers[0].read_text())
            self.assertIn('cpu_ledger', ledger['identity']['module'])
            self.assertTrue(ledger['identity']['module_sha256'])
            for bucket in ('load_workbook', 'workbook_save', 'readonly_lazy_parse'):
                self.assertGreater(ledger['buckets'][bucket]['cpu_ns'], 0, bucket)
            self.assertEqual(ledger['readonly_rows_yielded'], 3000)
            # business semantics identical, workbook content identical modulo
            # openpyxl's volatile save timestamps
            self.assertEqual(self._business_fingerprint(baseline),
                             self._business_fingerprint(measured))

    def test_no_ledger_written_without_environment_variable(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            hook_dir = root / 'hook'
            hook_dir.mkdir()
            for name in ('cpu_ledger.py', 'sitecustomize.py'):
                (hook_dir / name).write_bytes((MEASUREMENT_DIR / name).read_bytes())
            program = root / 'business.py'
            program.write_text(self.BUSINESS)
            import os
            completed = subprocess.run(
                [sys.executable, str(program), str(root / 'out.xlsx')],
                capture_output=True, text=True,
                env=dict(os.environ, PYTHONPATH=str(hook_dir)), check=True)
            self.assertEqual(json.loads(completed.stdout)['rows_total'],
                             sum(i + i * 3 for i in range(3000)))
            self.assertEqual(list(root.glob('ledger-*.json')), [])
            self.assertEqual(list(root.glob('ledgers')), [])


if __name__ == '__main__':
    unittest.main()
