"""Targeted positive/negative cases for the replay-private semantic verifier.

The fixtures use openpyxl only; `original_content` reads formulas/values with
data_only=False so cache-clearing by openpyxl.save cannot mask these checks.
String literals, formula types, defined names, empty-string policy and unit
expression are exercised directly.
"""
import json
import re
import zipfile
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from compatibility.verification.verify import Checks
from compatibility.verification.spreadsheet import (
    compare_number_formats, format_meaning, original_content,
)
from compatibility.verification.workbooks import normalized_formula


def _sheet_part(book_path, sheet_name):
    with zipfile.ZipFile(book_path) as z:
        wbxml = z.read('xl/workbook.xml').decode()
        rels = z.read('xl/_rels/workbook.xml.rels').decode()
        rid2target = {}
        for tag in re.findall(r'<Relationship\b[^>]*/>', rels):
            mid = re.search(r'Id="([^"]+)"', tag)
            mtgt = re.search(r'Target="([^"]+)"', tag)
            if mid and mtgt:
                rid2target[mid.group(1)] = mtgt.group(1)
        match = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="([^"]+)"' % re.escape(sheet_name), wbxml)
        target = rid2target[match.group(1)].lstrip('/')
        return target if target.startswith('xl/') else 'xl/' + target


def _inject_empty_string(book_path, sheet_name, coord):
    """Add a real `<is><t></t></is>` empty-string cell like Census workbooks."""
    part = _sheet_part(book_path, sheet_name)
    with zipfile.ZipFile(book_path) as z:
        xml = z.read(part).decode()
    row = int(re.search(r'(\d+)$', coord).group(1))
    cell = '<c r="%s" t="inlineStr"><is><t></t></is></c>' % coord
    marker = '<row r="%d">' % row
    if marker in xml:
        xml = xml.replace(marker, marker + cell, 1)
    else:
        xml = xml.replace('</sheetData>', cell + '</sheetData>', 1)
    tmp = book_path.with_suffix('.tmp.xlsx')
    with zipfile.ZipFile(book_path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                data = xml.encode()
            zout.writestr(item, data)
    import os
    os.replace(tmp, book_path)


def build_source(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'S1'
    ws['A1'] = '=1+2'
    ws['B1'] = '=IF(A1="hello world",1,0)'
    ws['C1'] = 5
    ws['E1'] = 123
    ws['E1'].number_format = '$#,##0'
    ws.merge_cells('A5:B5')
    wb.defined_names['myname'] = DefinedName(name='myname', attr_text="S1!$A$1")
    ws2 = wb.create_sheet('S2')
    ws2['B1'] = 'x'  # ensure the row exists so the injected empty cell is placed in it
    wb.save(path)
    wb.close()
    _inject_empty_string(path, 'S1', 'D1')
    _inject_empty_string(path, 'S2', 'A1')


def roundtrip(src, dst):
    wb = load_workbook(src, data_only=False)
    wb.save(dst)
    wb.close()


class NormalizedFormulaTests(unittest.TestCase):
    def test_supported_serialization_equivalences_are_equal(self):
        self.assertEqual(normalized_formula('=SUM(A1:A2)'), normalized_formula('=sum(a1:a2)'))
        self.assertEqual(normalized_formula("='S1'!A1"), normalized_formula("=S1!$A$1"))
        self.assertEqual(normalized_formula('=A1 + B1'), normalized_formula('=A1+B1'))

    def test_sheet_name_whitespace_is_significant(self):
        self.assertNotEqual(normalized_formula("='A B'!A1"), normalized_formula('=AB!A1'))
        self.assertEqual(normalized_formula("='A B'!A1"), normalized_formula("='A B'!A1"))

    def test_escaped_single_quote_in_sheet_name(self):
        self.assertEqual(normalized_formula("='O''Brien'!A1"), normalized_formula("='O''Brien'!A1"))
        self.assertNotEqual(normalized_formula("='O''Brien'!A1"), normalized_formula("=O'Brien!A1"))

    def test_intersection_whitespace_is_significant(self):
        self.assertNotEqual(normalized_formula('=A1:B1 A1:C1'), normalized_formula('=A1:B1A1:C1'))
        self.assertEqual(normalized_formula('=A1:B1 A1:C1'), normalized_formula('=A1:B1 A1:C1'))

    def test_string_literal_content_is_verbatim(self):
        self.assertNotEqual(normalized_formula('=IF(A1="hello world",1,0)'),
                            normalized_formula('=IF(A1="helloworld",1,0)'))
        self.assertNotEqual(normalized_formula('=IF(A1="Hello",1,0)'),
                            normalized_formula('=IF(A1="hello",1,0)'))
        self.assertEqual(normalized_formula('=IF(A1="hello world",1,0)'),
                         normalized_formula('=IF(A1="hello world",1,0)'))


def make_names_book(path, names):
    """names: list of (name, scope, target); scope None=global or 'S1'/'S2'."""
    wb = Workbook()
    wb.active.title = 'S1'
    wb.active['A1'] = 1
    wb.active['A2'] = 2
    ws2 = wb.create_sheet('S2')
    ws2['A1'] = 3
    ws2['A2'] = 4
    idx = {'S1': 0, 'S2': 1}
    for name, scope, target in names:
        item = DefinedName(name=name, attr_text=target,
                           localSheetId=(idx[scope] if scope else None))
        if scope:
            wb[scope].defined_names[name] = item
        else:
            wb.defined_names[name] = item
    wb.save(path)
    wb.close()


def names_pair(tmp, src_names, out_names):
    src = tmp / 'names_src.xlsx'
    dst = tmp / 'names_dst.xlsx'
    make_names_book(src, src_names)
    make_names_book(dst, out_names)
    return src, dst


class OriginalContentPositiveNegativeTests(unittest.TestCase):
    def _checks_for(self, src, dst, policy=None):
        checks = Checks({'dataset_id': 'fixture', 'kind': 'xlsx'})
        original_content(src, dst, checks, empty_string_policy=policy)
        return checks

    def _make_pair(self, tmp, tamper=None, policy=None):
        src = tmp / 'src.xlsx'
        dst = tmp / 'dst.xlsx'
        build_source(src)
        roundtrip(src, dst)
        if tamper:
            tamper(dst)
        return src, dst

    def test_clean_roundtrip_passes_with_declared_empty_string_policy(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = self._make_pair(tmp, policy={'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'success', checks.data['failures'])
            data = checks.data['checks']['source:data:S1']
            self.assertNotIn('D1', data.get('detail', {}).get('missing', []))

    def test_empty_string_loss_on_unlisted_sheet_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = self._make_pair(tmp)  # policy lists nothing
            checks = self._checks_for(src, dst, None)
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:data:S2', checks.data['failures'])

    def test_data_value_damage_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb['S1']['C1'] = 6
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:data:S1', checks.data['failures'])

    def test_formula_replaced_by_text_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb['S1']['A1'] = '1+2'  # same visual text but no longer a formula
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:formulas:S1', checks.data['failures'])

    def test_formula_string_literal_change_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb['S1']['B1'] = '=IF(A1="helloworld",1,0)'
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:formulas:S1', checks.data['failures'])

    def test_defined_name_missing_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                del wb.defined_names['myname']
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names', checks.data['failures'])

    def test_defined_name_ref_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb.defined_names['myname'] = DefinedName(name='myname', attr_text='#REF!')
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names_damaged', checks.data['failures'])

    def test_defined_name_target_sheet_removed_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb.defined_names['myname'] = DefinedName(name='myname', attr_text="GONE!$A$1")
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names_damaged', checks.data['failures'])

    def test_defined_name_target_cell_change_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = names_pair(tmp, [('G', None, 'S1!$A$1')], [('G', None, 'S1!$A$2')])
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names_damaged', checks.data['failures'])

    def test_defined_name_nonexistent_target_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = names_pair(tmp, [('G', None, 'S1!$A$1')], [('G', None, 'MissingAlias')])
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names_damaged', checks.data['failures'])

    def test_defined_name_local_deleted_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src_names = [('G', None, 'S1!$A$1'), ('Loc', 'S1', 'S1!$A$1')]
            out_names = [('G', None, 'S1!$A$1')]
            src, dst = names_pair(tmp, src_names, out_names)
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names', checks.data['failures'])

    def test_defined_name_global_and_local_scope_change_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            # global X and local X on S2; the local moves to S1 in the output.
            src_names = [('X', None, 'S1!$A$1'), ('X', 'S2', 'S2!$A$1')]
            out_names = [('X', None, 'S1!$A$1'), ('X', 'S1', 'S1!$A$1')]
            src, dst = names_pair(tmp, src_names, out_names)
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names', checks.data['failures'])

    def test_added_local_name_cannot_shadow_used_global(self):
        for formula, local_name in [('=Chosen', 'Chosen'), ('=SUM(Chosen)', 'chosen'),
                                    ("='S2'!Chosen", 'Chosen')]:
            with self.subTest(formula=formula, local_name=local_name), tempfile.TemporaryDirectory() as tmp:
                tmp = Path(tmp)
                src, dst = tmp / 'src.xlsx', tmp / 'dst.xlsx'
                wb = Workbook()
                wb.active.title = 'S1'
                wb['S1']['A1'], wb['S1']['A2'] = 10, 20
                wb.create_sheet('S2')
                wb['S1']['B1'] = formula
                wb.defined_names['Chosen'] = DefinedName('Chosen', attr_text='S1!$A$1')
                wb.save(src)
                scope = 'S2' if 'S2' in formula else 'S1'
                wb[scope].defined_names[local_name] = DefinedName(local_name, attr_text='S1!$A$2')
                wb.save(dst)
                wb.close()
                report = self._checks_for(src, dst).finish()
                self.assertEqual(report['status'], 'failed')
                self.assertIn('source:defined_name_shadowing', report['failures'])

    def test_unused_added_local_names_are_allowed(self):
        for formula, scope in [('="Chosen"', 'S1'), ('=Chosen', 'S2')]:
            with self.subTest(formula=formula, scope=scope), tempfile.TemporaryDirectory() as tmp:
                tmp = Path(tmp)
                src, dst = tmp / 'src.xlsx', tmp / 'dst.xlsx'
                wb = Workbook()
                wb.active.title = 'S1'
                wb['S1']['A1'], wb['S1']['A2'] = 10, 20
                wb.create_sheet('S2')
                wb['S1']['B1'] = formula
                wb.defined_names['Chosen'] = DefinedName('Chosen', attr_text='S1!$A$1')
                wb.save(src)
                wb[scope].defined_names['Chosen'] = DefinedName('Chosen', attr_text='S1!$A$2')
                wb.save(dst)
                wb.close()
                report = self._checks_for(src, dst).finish()
                self.assertEqual(report['status'], 'success', report['failures'])
                self.assertTrue(report['diagnostics']['source:defined_names_added']['detail'])

    def test_added_local_name_with_equivalent_target_is_allowed(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = tmp / 'src.xlsx', tmp / 'dst.xlsx'
            wb = Workbook()
            wb.active.title = 'S1'
            wb['S1']['A1'], wb['S1']['B1'] = 10, '=Chosen'
            wb.defined_names['Chosen'] = DefinedName('Chosen', attr_text='S1!$A$1')
            wb.save(src)
            wb['S1'].defined_names['Chosen'] = DefinedName('Chosen', attr_text="'S1'!A1")
            wb.save(dst)
            wb.close()
            report = self._checks_for(src, dst).finish()
            self.assertEqual(report['status'], 'success', report['failures'])

    def test_defined_name_indirect_missing_reference_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src_names = [('G', None, '=Other'), ('Other', None, 'S1!$A$1')]
            out_names = [('G', None, '=Other')]  # Other removed -> unresolvable chain
            src, dst = names_pair(tmp, src_names, out_names)
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:defined_names_damaged', checks.data['failures'])

    def test_defined_name_cycle_does_not_hang(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            cycle = [('A', None, '=B'), ('B', None, '=A')]
            src, dst = names_pair(tmp, cycle, cycle)
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            # both sides resolve to the same cycle; clean round-trip passes.
            self.assertEqual(checks.finish()['status'], 'success', checks.data['failures'])

    def test_defined_name_serialization_only_is_diagnostic(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = names_pair(tmp, [('G', None, 'S1!$A$1')], [('G', None, "'S1'!A1")])
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            report = checks.finish()
            self.assertEqual(report['status'], 'success', report['failures'])
            self.assertIn('source:defined_names_serialization', report['diagnostics'])

    def test_chart_reference_change_fails(self):
        from openpyxl.chart import LineChart, Reference
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def build(path, rows):
                wb = Workbook()
                ws = wb.active
                ws.title = 'S1'
                for r in range(1, 6):
                    ws.cell(r, 1, value=r)
                ch = LineChart()
                ch.add_data(Reference(ws, min_col=1, min_row=1, max_row=rows))
                ws.add_chart(ch, 'E1')
                wb.save(path)
                wb.close()
            src = tmp / 'chart_src.xlsx'
            dst = tmp / 'chart_dst.xlsx'
            build(src, 3)
            build(dst, 4)  # data range widened
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:charts:S1', checks.data['failures'])

    def test_unit_expression_format_change_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb['S1']['E1'].number_format = '0'  # currency marker dropped
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:number_formats_meaning:S1', checks.data['failures'])

    def test_merged_region_removed_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            def tamper(dst):
                wb = load_workbook(dst, data_only=False)
                wb['S1'].unmerge_cells('A5:B5')
                wb.save(dst)
            src, dst = self._make_pair(tmp, tamper, {'sheets': ['S1', 'S2']})
            checks = self._checks_for(src, dst, {'sheets': ['S1', 'S2']})
            self.assertEqual(checks.finish()['status'], 'failed')
            self.assertIn('source:merged_regions_removed:S1', checks.data['failures'])

    def test_style_serialization_is_only_diagnostic(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src = tmp / 'src.xlsx'
            dst = tmp / 'dst.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'S1'
            ws['A1'] = 1
            ws['A1'].font = ws['A1'].font.copy(name='Arial')
            wb.save(src)
            out = load_workbook(src, data_only=False)
            out['S1']['A1'].font = out['S1']['A1'].font.copy(bold=True)
            out.save(dst)
            checks = self._checks_for(src, dst, {'sheets': ['S1']})
            report = checks.finish()
            self.assertEqual(report['status'], 'success')
            self.assertIn('source:styles:S1', report['diagnostics'])
            self.assertGreater(report['diagnostic_diff_count'], 0)


class NumberFormatMeaningTests(unittest.TestCase):
    def test_confirmed_date_serializations_are_equivalent(self):
        pairs = (
            ('m/d/yy h:mm', 'm/d/yyyy\\ h:mm'),
            ('mm/dd/yyyy', 'm/d/yyyy'),
            ('M/D/YY H:MM', 'm/d/yyyy h:mm'),
        )
        for source, output in pairs:
            with self.subTest(source=source, output=output):
                result = compare_number_formats(source, output)
                self.assertTrue(result['equivalent'], result)
                self.assertEqual(result['classification'], 'allowed_representation_change')

    def test_date_fields_order_and_duration_units_are_preserved(self):
        pairs = (
            ('mm/dd/yyyy', 'dd/mm/yyyy'),
            ('m/d/yy h:mm', 'm/d/yy'),
            ('[h]', '[m]'),
            ('[h]', '0'),
        )
        for source, output in pairs:
            with self.subTest(source=source, output=output):
                self.assertFalse(compare_number_formats(source, output)['equivalent'])

    def test_currency_percent_literals_and_scaling_are_preserved(self):
        pairs = (
            ('$#,##0.00', '#,##0.00'),
            ('[$$-409]#,##0.00', '#,##0.00'),
            ('[$€-407]#,##0.00', '#,##0.00'),
            ('0.00%', '0.00'),
            ('0.00" kg"', '0.00" KG"'),
            ('0,,', '0'),  # trailing commas scale; grouping commas do not.
            ('#,##0', '0'),
        )
        for source, output in pairs:
            with self.subTest(source=source, output=output):
                self.assertFalse(compare_number_formats(source, output)['equivalent'])

    def test_quoted_and_escaped_units_and_locale_colors(self):
        self.assertTrue(compare_number_formats('0.00" kg"', r'0.00\ \k\g')['equivalent'])
        self.assertTrue(compare_number_formats('$#,##0.00', '"$"#,##0.00')['equivalent'])
        self.assertTrue(compare_number_formats('[$$-409]#,##0.00', '[$$-407]#,##0.00')['equivalent'])
        self.assertTrue(compare_number_formats('[Red]0.00', '0.00')['equivalent'])
        # Literal text remains case-sensitive; unsupported changed syntax is
        # not silently treated as equivalent.
        self.assertFalse(compare_number_formats('0.0"USD"', '0.0"usd"')['equivalent'])
        self.assertFalse(compare_number_formats('[DBNum1]0', '[DBNum2]0')['equivalent'])
        self.assertEqual(format_meaning('[DBNum1]0'), format_meaning('[DBNum1]0'))
        self.assertEqual(compare_number_formats('0,,', '0')['reason'], 'numeric_scaling_changed')

    def test_original_content_accepts_retail_date_format_and_reports_all_diffs(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = tmp / 'src.xlsx', tmp / 'dst.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'S1'
            ws['A1'] = __import__('datetime').datetime(2010, 12, 1, 8, 26)
            ws['A1'].number_format = 'm/d/yy h:mm'
            ws['B1'] = 24.5
            ws['B1'].number_format = '[$$-409]#,##0.00'
            wb.save(src)
            wb.close()
            out = load_workbook(src, data_only=False)
            out['S1']['A1'].number_format = 'm/d/yyyy\\ h:mm'
            out['S1']['B1'].number_format = '#,##0.00'
            out.save(dst)
            out.close()

            checks = Checks({'dataset_id': 'fixture', 'kind': 'xlsx'})
            original_content(src, dst, checks)
            self.assertIn('source:number_formats_meaning:S1', checks.data['failures'])
            stats = checks.data['number_format_statistics']['S1']
            self.assertEqual(stats['difference_count'], 2)
            self.assertEqual(stats['distinct_format_pair_count'], 2)
            self.assertEqual(stats['allowed_representation_change_count'], 1)
            self.assertEqual(stats['strict_failure_count'], 1)
            self.assertEqual(stats['cell_samples'][0]['cell'], 'A1')

    def test_original_content_accepts_real_retail_date_representation(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = tmp / 'src.xlsx', tmp / 'dst.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'S1'
            ws['E2'] = __import__('datetime').datetime(2010, 12, 1, 8, 26)
            ws['E2'].number_format = 'm/d/yy h:mm'
            wb.save(src)
            wb.close()
            out = load_workbook(src, data_only=False)
            out['S1']['E2'].number_format = 'm/d/yyyy\\ h:mm'
            out.save(dst)
            out.close()
            checks = Checks({'dataset_id': 'fixture', 'kind': 'xlsx'})
            original_content(src, dst, checks)
            self.assertNotIn('source:number_formats_meaning:S1', checks.data['failures'])
            stats = checks.data['number_format_statistics']['S1']
            self.assertEqual(stats['difference_count'], 1)
            self.assertEqual(stats['allowed_representation_change_count'], 1)
            self.assertEqual(stats['strict_failure_count'], 0)

    def test_original_content_counts_every_difference_but_bounds_samples(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            src, dst = tmp / 'src.xlsx', tmp / 'dst.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.title = 'S1'
            for row in range(1, 31):
                ws.cell(row, 1, row).number_format = '0.00" kg"'
            wb.save(src)
            wb.close()
            out = load_workbook(src, data_only=False)
            for row in range(1, 31):
                out['S1'].cell(row, 1).number_format = r'0.00\ \k\g'
            out.save(dst)
            out.close()
            checks = Checks({'dataset_id': 'fixture', 'kind': 'xlsx'})
            original_content(src, dst, checks)
            stats = checks.data['number_format_statistics']['S1']
            self.assertEqual(stats['difference_count'], 30)
            self.assertEqual(stats['format_pairs'][0]['count'], 30)
            self.assertEqual(stats['distinct_format_pair_count'], 1)
            self.assertEqual(len(stats['cell_samples']), 24)
            self.assertEqual(stats['strict_failure_count'], 0)

class DiagnosticsCountingTests(unittest.TestCase):
    def test_total_slots_vs_nonempty_diffs(self):
        checks = Checks({'dataset_id': 'x', 'kind': 'xlsx'})
        checks.diag('a', [])
        checks.diag('b', [1, 2])
        checks.diag('c', ['x'])
        report = checks.finish()
        self.assertEqual(report['diagnostics_total'], 3)
        self.assertEqual(report['diagnostic_diff_count'], 2)


if __name__ == '__main__':
    unittest.main()
