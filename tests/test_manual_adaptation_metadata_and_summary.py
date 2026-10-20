"""Small fixtures for the manually adapted XLSX dataset bindings."""
from datetime import datetime
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

from compatibility.helpers.hr_metro_build import build as build_hr_metro
from compatibility.helpers.retail_online_build import build as build_retail_online
from compatibility.verification.verify import Checks
from compatibility.verification.spreadsheet import original_content


ROOT = Path(__file__).resolve().parents[1]
NS = {
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'pr': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'ct': 'http://schemas.openxmlformats.org/package/2006/content-types',
}


def shared_string_empty_fixture(path):
    book = Workbook()
    sheet = book.active
    sheet.title = 'MSA_M2025_dl'
    sheet['A1'] = 'header'
    sheet['A2'] = ''
    sheet['A3'] = '*'
    sheet['A4'] = '=1+1'
    book.save(path)
    book.close()
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet_xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
    for cell in sheet_xml.findall('.//m:c', NS):
        address = cell.attrib.get('r')
        if address == 'A2':
            cell.clear()
            cell.attrib.update(r='A2', t='s')
            ET.SubElement(cell, '{%s}v' % NS['m']).text = '0'
        elif address == 'A3':
            cell.attrib['t'] = 's'
            for child in list(cell):
                cell.remove(child)
            ET.SubElement(cell, '{%s}v' % NS['m']).text = '1'
    parts['xl/worksheets/sheet1.xml'] = ET.tostring(
        sheet_xml, encoding='utf-8', xml_declaration=True)

    rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
    ET.SubElement(rels, '{%s}Relationship' % NS['pr'], {
        'Id': 'rIdSharedStrings',
        'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings',
        'Target': 'sharedStrings.xml'})
    parts['xl/_rels/workbook.xml.rels'] = ET.tostring(
        rels, encoding='utf-8', xml_declaration=True)
    types = ET.fromstring(parts['[Content_Types].xml'])
    ET.SubElement(types, '{%s}Override' % NS['ct'], {
        'PartName': '/xl/sharedStrings.xml',
        'ContentType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml'})
    parts['[Content_Types].xml'] = ET.tostring(
        types, encoding='utf-8', xml_declaration=True)
    strings = ET.Element('{%s}sst' % NS['m'], {'count': '2', 'uniqueCount': '2'})
    ET.SubElement(ET.SubElement(strings, '{%s}si' % NS['m']), '{%s}t' % NS['m'])
    ET.SubElement(ET.SubElement(strings, '{%s}si' % NS['m']), '{%s}t' % NS['m']).text = '*'
    parts['xl/sharedStrings.xml'] = ET.tostring(
        strings, encoding='utf-8', xml_declaration=True)
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)


def _assert_original_content(source, output, authorized):
    checks = Checks({'dataset_id': 'hr_metro', 'kind': 'xlsx'})
    original_content(source, output, checks, {'sheets': authorized})
    return checks.data['checks']


def _write_summary_inputs(workspace, dataset_id, workbook_name, output_book):
    (workspace / 'input').mkdir(parents=True, exist_ok=True)
    (workspace / 'output').mkdir()
    source = workspace / 'input' / workbook_name
    target = workspace / 'output' / workbook_name
    target.write_bytes(source.read_bytes())
    # The helpers consume a frozen preparation artifact; generate one from this
    # small fixture exactly like the out-of-loop preparation stage would.
    from scripts import prepare_xlsx_selection as preparation
    preparation.cli(['--dataset', dataset_id, '--input', str(source),
                     '--output', str(workspace / 'input/prepared_selection.json'),
                     '--fixture'])
    (workspace / 'input/dataset_manifest.json').write_text(json.dumps({
        'kind': 'xlsx', 'dataset_id': dataset_id, 'workbook': workbook_name,
        'output_workbook': output_book}))
    (workspace / 'output/formula_recalc.json').write_text(json.dumps({
        'status': 'success', 'total_errors': 0, 'total_formulas': 8}))
    (workspace / 'output/business_verification.json').write_text(json.dumps({
        'status': 'success', 'failures': []}))
    return source, workspace / 'output' / workbook_name


class ManualAdaptationTests(unittest.TestCase):
    def test_only_declared_hr_table_allows_shared_empty_string_normalization(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / 'source.xlsx'
            shared_string_empty_fixture(source)
            loaded = load_workbook(source, read_only=True, data_only=False)
            self.assertEqual(loaded['MSA_M2025_dl']['A2'].value, '')
            loaded.close()

            output = root / 'output.xlsx'
            book = load_workbook(source)
            book.save(output)  # openpyxl serializes the shared empty string as a blank.
            book.close()
            ok = _assert_original_content(source, output, ['MSA_M2025_dl'])
            self.assertTrue(ok['source:data:MSA_M2025_dl']['ok'])

            broken = root / 'broken.xlsx'
            book = load_workbook(output)
            book['MSA_M2025_dl']['A3'] = None  # suppression marker must remain.
            book.save(broken)
            book.close()
            failed = _assert_original_content(source, broken, ['MSA_M2025_dl'])
            self.assertFalse(failed['source:data:MSA_M2025_dl']['ok'])
            self.assertTrue(failed['source:formulas:MSA_M2025_dl']['ok'])

            formula_broken = root / 'formula-broken.xlsx'
            book = load_workbook(output)
            book['MSA_M2025_dl']['A4'] = None
            book.save(formula_broken)
            book.close()
            formula_failure = _assert_original_content(
                source, formula_broken, ['MSA_M2025_dl'])
            self.assertFalse(formula_failure['source:data:MSA_M2025_dl']['ok'])

            unauthorized = _assert_original_content(source, output, [])
            self.assertFalse(unauthorized['source:data:MSA_M2025_dl']['ok'])

    def test_retail_helper_diagnostics_survive_generic_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp)
            (ws / 'input').mkdir()
            source = ws / 'input/Online Retail.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = 'Online Retail'
            sheet.append(['InvoiceNo', 'Quantity', 'UnitPrice', 'InvoiceDate', 'CustomerID'])
            sheet.append(['C1', 1, 10, datetime(2010, 1, 1), 1])
            sheet.append(['2', 2, 3, datetime(2010, 2, 1), None])
            book.save(source)
            book.close()
            _source, output = _write_summary_inputs(ws, 'retail_online', source.name,
                                                     'retail_online_enhanced.xlsx')
            build_retail_online(output)
            subprocess.run([sys.executable, str(ROOT / 'compatibility/tools/summary.py')],
                           cwd=ws, check=True, capture_output=True, text=True)
            diag = json.loads((ws / 'output/xlsx_business_diagnostics.json').read_text())
            summary = json.loads((ws / 'output/xlsx_enhancement_summary.json').read_text())
            self.assertEqual(diag['exclusion_counts']['cancelled_invoice'], 1)
            self.assertEqual(diag['eligible_lines_missing_customer_id'], 1)
            self.assertEqual(diag['currency'], 'GBP')
            self.assertEqual(summary['final_verifier_status'], 'success')
            self.assertEqual(summary['recalculation']['status'], 'success')
            self.assertEqual(summary['chart_count'], 1)

    def test_hr_metro_business_scope_survives_generic_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp)
            (ws / 'input').mkdir()
            source = ws / 'input/MSA_M2025_dl.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = 'MSA_M2025_dl'
            sheet.append(['AREA', 'AREA_TITLE', 'AREA_TYPE', 'NAICS', 'I_GROUP',
                          'OWN_CODE', 'O_GROUP', 'OCC_CODE', 'OCC_TITLE', 'TOT_EMP',
                          'A_MEAN', 'H_MEDIAN', 'H_MEAN', 'A_MEDIAN'])
            scope = ['35620', 'New York-Newark-Jersey City, NY-NJ', '4',
                     '000000', 'cross-industry', '1235']
            sheet.append(scope + ['total', '00-0000', 'All occupations', 1000,
                                  80000, 30, 40, 60000])
            for i in range(5):
                sheet.append(scope + ['major', f'{i:02d}-0000', f'Group {i}',
                                      500-i, 70000, 25, 35, 50000])
            book.save(source)
            book.close()
            _source, output = _write_summary_inputs(ws, 'hr_metro', source.name,
                                                     'hr_metro_enhanced.xlsx')
            build_hr_metro(output)
            subprocess.run([sys.executable, str(ROOT / 'compatibility/tools/summary.py')],
                           cwd=ws, check=True, capture_output=True, text=True)
            diag = json.loads((ws / 'output/xlsx_business_diagnostics.json').read_text())
            summary = json.loads((ws / 'output/xlsx_enhancement_summary.json').read_text())
            self.assertEqual(diag['area_code'], '35620')
            self.assertIn('New York-Newark-Jersey City', diag['area_title'])
            self.assertIn('scope', diag)
            self.assertEqual(summary['final_verifier_status'], 'success')
            self.assertEqual(summary['recalculation']['status'], 'success')


if __name__ == '__main__':
    unittest.main()
