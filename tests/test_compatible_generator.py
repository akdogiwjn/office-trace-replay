"""Small-fixture tests for compatibility recipe/source generation."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from scripts.compile_compatible import (
    _copy_current_verifier_package, _freeze_base_sheet_dimensions,
    _update_industry_productivity_recipe, main as compile_main,
)
from replay.common import digest, write
from replay import freeze


class CompatibleGeneratorTests(unittest.TestCase):
    def test_pdf_generator_installs_current_loop_and_independent_verifiers(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            package = root / 'compatibility/verification'
            package.mkdir(parents=True)
            for name in ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py', 'pdf.py', 'pdf_loop.py'):
                (package / name).write_text(name + '\n')
            source = root / 'case/source'
            _copy_current_verifier_package(root, source, pdf=True)
            generated = source / 'office_trace_bench'
            for name in ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py'):
                self.assertEqual((generated / name).read_text(), name + '\n')
            self.assertEqual((generated / 'pdf_loop.py').read_text(), 'pdf_loop.py\n')
            self.assertEqual((generated / 'compatible_pdf.py').read_text(), 'pdf.py\n')

    def test_industry_generator_uses_shared_summary_and_preserves_parent_binding(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            recipes = root / 'compatibility/recipes'
            recipes.mkdir(parents=True)
            verification = root / 'compatibility/verification'
            verification.mkdir(parents=True)
            for name in ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py'):
                (verification / name).write_text(name + '\n')
            source_package = root / 'compatibility/cases/industry/source/office_trace_bench'
            source_package.mkdir(parents=True)
            parent_path = root / 'compatibility/cases/industry/canonical/parent-run.json'
            parent_path.parent.mkdir(parents=True)
            parent = {'run_id': 'accepted-run', 'canonical_sha256': 'a' * 64}
            write(parent_path, parent)
            recipe_path = recipes / 'industry_productivity.json'
            seed = root / 'compatibility/cases/industry/seed/input'
            seed.mkdir(parents=True)
            (seed / 'verify_office.py').write_text('print("fixture verifier")\n')
            write(recipe_path, {
                'dataset_id': 'industry_productivity',
                'seed_path': 'compatibility/cases/industry/seed',
                'input_hashes': {},
                'parent': {'run_id': parent['run_id'], 'canonical_sha256': parent['canonical_sha256'],
                           'recipe_path': str(parent_path.relative_to(root))},
                'key_operations': [{'source_tool_calls': [{
                    'tool_call': {'function_name': 'exec', 'arguments': {
                        'command': 'python3 compatibility/tools/industry_productivity_summary.py'}}}]}],
                'adaptations': ['P04 uses metadata-only summary.'],
                'derivation': {'adaptations': ['metadata-only summary follows verifier.']},
                'source_path': 'compatibility/cases/industry/source',
            })
            index = {'entries': []}
            _update_industry_productivity_recipe(root, index)
            updated = json.loads(recipe_path.read_text())
            self.assertEqual(updated['parent']['run_id'], parent['run_id'])
            self.assertEqual(updated['parent']['canonical_sha256'], parent['canonical_sha256'])
            command = updated['key_operations'][0]['source_tool_calls'][0]['tool_call']['arguments']['command']
            self.assertIn('compatibility/tools/write_xlsx_summary.py', command)
            self.assertNotIn('industry_productivity_summary.py', command)
            self.assertIn('shared workbook-reading summary', updated['adaptations'][0])
            self.assertEqual(index['entries'][0]['recipe_sha256'], digest(recipe_path))

    def test_generator_main_emits_loop_entry_and_xlsx_base_dimensions(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for relative in ('compatibility/verification', 'compatibility/prompts',
                             'compatibility/helpers', 'compatibility/cases/retail-parent/seed/input',
                             'compatibility/cases/retail-parent/source/datasets/xlsx/retail',
                             'compatibility/cases/retail-parent/source/office_trace_bench',
                             'legacy/document-oe', 'recipes'):
                (root / relative).mkdir(parents=True, exist_ok=True)
            for name in ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py'):
                (root / 'compatibility/verification' / name).write_text(name + '\n')
            (root / 'compatibility/contract.json').write_text('{}\n')
            (root / 'legacy/document-oe/lock.json').write_text('{}\n')
            (root / 'compatibility/prompts/xlsx.txt').write_text('fixture prompt\n')
            (root / 'compatibility/helpers/retail_build_enhance.py').write_text(
                "TARGET = '/fixture/output.xlsx'\n")

            source_book = Workbook()
            source_book.active.title = 'Data'
            source_book.active['A1'] = 'value'
            source_book.active['A2'] = 7
            input_path = root / 'compatibility/cases/retail-parent/seed/input/tiny.xlsx'
            source_book.save(input_path)
            source_book.close()
            old_verifier = root / 'compatibility/cases/retail-parent/seed/input/verify_office.py'
            old_verifier.write_text('print("old independent verifier")\n')
            manifest = {
                'schema_version': 'office-dataset-v2', 'kind': 'xlsx', 'dataset_id': 'retail',
                'domain': 'retail', 'task': 'tiny fixture', 'runtime_case': 'RETAIL-FIXTURE',
                'workbook': 'tiny.xlsx', 'output_workbook': 'retail_enhanced.xlsx',
                'expected_file': 'expected.json', 'trace_rules': [],
                'requirements': {'metrics': [], 'scenario': {'default': 'Base', 'rows': [], 'projections': []},
                                 'charts': []},
            }
            old_manifest = root / 'compatibility/cases/retail-parent/seed/input/dataset_manifest.json'
            old_manifest.write_text(json.dumps(manifest))
            source_manifest = root / 'compatibility/cases/retail-parent/source/datasets/xlsx/retail/manifest.json'
            source_manifest.write_text(json.dumps(manifest))
            (source_manifest.parent / 'expected.json').write_text('{}\n')
            parent_recipe = {
                'seed_path': 'compatibility/cases/retail-parent/seed',
                'source_path': 'compatibility/cases/retail-parent/source',
                'workspace': '/workspace/RETAIL-FIXTURE',
                'source': {'original_run_id': 'fixture-parent', 'source_sha256': '1' * 64},
                'key_operations': [],
            }
            (root / 'recipes/retail.json').write_text(json.dumps(parent_recipe))

            compile_main(root=root, datasets=['retail'])

            case = root / 'compatibility/cases/retail'
            generated = json.loads((case / 'seed/input/dataset_manifest.json').read_text())
            self.assertEqual(generated['requirements']['base_sheets'], ['Data'])
            self.assertEqual(generated['requirements']['base_dimensions'], {'Data': [2, 1]})
            self.assertEqual(generated['requirements']['base_dimension_details']['Data']['effective_data_range'], [2, 1])
            self.assertEqual(generated['requirements']['base_dimension_details']['Data']['serialized_cell_range'], [2, 1])
            self.assertEqual(generated['loop_verification_contract'],
                             'document-oe-compatible-xlsx-loop-v1')
            wrapper = (case / 'seed/input/verify_office.py').read_text()
            self.assertIn("scope='loop'", wrapper)
            self.assertNotIn('old independent verifier', wrapper)
            self.assertEqual(wrapper, (case / 'source/runtime/verify_office.py').read_text())
            recipe = json.loads((root / 'compatibility/recipes/retail.json').read_text())
            verify_calls = [call['tool_call']['arguments']['command']
                            for step in recipe['key_operations']
                            for call in step['source_tool_calls']
                            if call['role'] == 'semantic_verification']
            self.assertEqual(len(verify_calls), 1)
            self.assertIn('input/verify_office.py', verify_calls[0])
            self.assertIn("scope='loop'", wrapper)
            self.assertEqual(recipe['input_hashes']['input/verify_office.py'],
                             digest(case / 'seed/input/verify_office.py'))
            suite = json.loads((root / 'compatibility/suite.json').read_text())
            self.assertEqual(suite['entries'][0]['recipe_sha256'], digest(root / suite['entries'][0]['recipe_path']))

    def test_large_complexity_metadata_does_not_trigger_preparation_value_scan(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            book_path = root / 'input.xlsx'
            wb = Workbook()
            wb.active['A1'] = 'x'
            wb.save(book_path)
            wb.close()
            complexity_path = root / 'complexity.json'
            complexity_path.write_text(json.dumps({
                'populated_cell_count': 566659,
                'sheets': [{'name': 'Sheet', 'row_count': 566659, 'column_count': 9}],
            }))
            from openpyxl.worksheet._read_only import ReadOnlyWorksheet
            with patch.object(ReadOnlyWorksheet, 'iter_rows',
                              side_effect=AssertionError('large workbook value scan entered')):
                dimensions, details = _freeze_base_sheet_dimensions(book_path, complexity_path)
            self.assertEqual(dimensions['Sheet'], [566659, 9])
            self.assertIn('not rescanned', details['Sheet']['effective_range_method'])

    def test_shared_summary_reads_a_small_workbook(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace = Path(tmp)
            (workspace / 'input').mkdir()
            output = workspace / 'output'
            output.mkdir()
            source = Workbook()
            source.active.title = 'Annual'
            source.save(workspace / 'input/source.xlsx')
            source.close()
            result = Workbook()
            result.active.title = 'Summary'
            result.create_sheet('Annual')
            result.save(output / 'result.xlsx')
            result.close()
            (workspace / 'input/dataset_manifest.json').write_text(json.dumps({
                'kind': 'xlsx', 'workbook': 'source.xlsx', 'output_workbook': 'result.xlsx'}))
            (output / 'workbook_features.json').write_text(json.dumps({
                'summary_sheet': 'Summary', 'metrics': {}, 'scenario': {}, 'charts': {}}))
            (output / 'formula_recalc.json').write_text(json.dumps({'status': 'success'}))
            (output / 'business_verification.json').write_text(json.dumps({'status': 'success'}))
            script = Path(__file__).parents[1] / 'compatibility/tools/summary.py'
            subprocess.run([sys.executable, str(script)], cwd=workspace, check=True,
                           capture_output=True, text=True)
            summary = json.loads((output / 'xlsx_enhancement_summary.json').read_text())
            self.assertEqual(summary['added_sheet_names'], ['Summary'])
            self.assertEqual(summary['chart_count'], 0)

    def test_runner_preflight_accepts_bound_xlsx_and_rejects_tampered_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for relative in ('compatibility/verification', 'compatibility/cases/retail/source/office_trace_bench',
                             'compatibility/cases/retail/seed/input'):
                (root / relative).mkdir(parents=True, exist_ok=True)
            for name in ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py'):
                central = root / 'compatibility/verification' / name
                central.write_text(name)
                (root / 'compatibility/cases/retail/source/office_trace_bench' / name).write_text(name)
            contract = root / 'compatibility/contract.json'
            prompt = root / 'compatibility/prompt.txt'
            parent = root / 'recipes/parent.json'
            for path in (contract, prompt, parent):
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_text(path.name)
            asset = root / 'compatibility/cases/retail/seed/input/source.xlsx'
            asset.write_bytes(b'workbook-fixture')
            recipe_path = root / 'compatibility/recipes/retail.json'
            recipe_path.parent.mkdir(parents=True)
            recipe = {
                'parent': {'recipe_path': 'recipes/parent.json', 'recipe_sha256': digest(parent)},
                'classification': 'canonicalized_deterministic_replay_derived_from_successful_agent_executions',
                'prompt_path': 'compatibility/prompt.txt', 'prompt_sha256': digest(prompt),
                'seed_path': 'compatibility/cases/retail/seed',
                'source_path': 'compatibility/cases/retail/source',
                'input_hashes': {'input/source.xlsx': digest(asset)}, 'helper_hashes': {},
            }
            write(recipe_path, recipe)
            suite = {'contract_sha256': digest(contract), 'entries': [{
                'dataset_id': 'retail', 'record_count': None,
                'recipe_path': 'compatibility/recipes/retail.json',
                'recipe_sha256': digest(recipe_path)}]}
            write(root / 'compatibility/suite.json', suite)
            with patch.object(freeze, 'ROOT', root), patch('replay.legacy.validate'), \
                 patch('replay.run.validate_project'), \
                 patch('replay.freeze.validate_source_assets') as source_assets:
                freeze.validate_dataset('xlsx', 'retail', {'source': 'bound'}, None,
                                        check_installed=False)
                source_assets.assert_called_once()
                asset.write_bytes(b'tampered')
                with self.assertRaisesRegex(ValueError, 'input hash mismatch'):
                    freeze.validate_dataset('xlsx', 'retail', {'source': 'bound'}, None,
                                            check_installed=False)


if __name__ == '__main__':
    unittest.main()
