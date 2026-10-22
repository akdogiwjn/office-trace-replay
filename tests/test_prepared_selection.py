"""Small self-contained fixtures for the frozen XLSX preparation stage.

These tests exercise the real behaviour of the preparation entry, both
enhancement helpers and the loop/independent verifiers on small workbooks:

  * HR selection: unique total, top-5 major ranking, suppression markers.
  * Retail aggregation: hand-computed fixture covering cancellations,
    nonpositive/nonnumeric quantity and price, date boundaries, missing
    CustomerID and repeated InvoiceNo; full-period KPIs vs 2010-only monthly.
  * Rejections: missing/stale/mis-bound preparation, tampered derived values,
    hardcoded KPIs, disconnected charts, formulas smuggled into derived sheets.
  * Helpers reference frozen positions instead of rescanning the business
    table (verified with a cell-access spy, not by grepping source strings).

These are fixture-level checks only; they are not LibreOffice acceptance and
they do not touch the real large workbooks.
"""
import json
import uuid
from datetime import date, datetime
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.workbook.workbook import Workbook as OpenpyxlWorkbook

from compatibility.helpers.hr_metro_build import build as hr_build
from compatibility.helpers.hr_metro_build import TARGET as HR_TARGET
from compatibility.helpers.retail_online_build import build as retail_build
from compatibility.helpers.retail_online_build import TARGET as RETAIL_TARGET
from compatibility.verification.spreadsheet import (
    verify_semantic_xlsx, verify_semantic_xlsx_loop)
from compatibility.verification.verify import Checks
from scripts import prepare_xlsx_selection as preparation
from scripts.prepare_manual_xlsx_assets import select_hr_metro, select_retail_rows

HR_HEADERS = ['AREA', 'AREA_TITLE', 'AREA_TYPE', 'NAICS', 'I_GROUP', 'OWN_CODE',
              'O_GROUP', 'OCC_CODE', 'OCC_TITLE', 'TOT_EMP', 'A_MEAN',
              'A_MEDIAN', 'H_MEAN', 'H_MEDIAN']
RETAIL_HEADERS = ['InvoiceNo', 'StockCode', 'Description', 'Quantity',
                  'InvoiceDate', 'UnitPrice', 'CustomerID', 'Country']


def hr_row(area, area_title, area_type, naics, i_group, own_code, o_group,
           occ_code, occ_title, tot_emp, a_mean, a_median, h_mean, h_median):
    return [area, area_title, area_type, naics, i_group, own_code, o_group,
            occ_code, occ_title, tot_emp, a_mean, a_median, h_mean, h_median]


def make_hr_workbook(path):
    metro = ('35620', 'New York-Newark-Jersey City, NY-NJ', '4', '000000',
             'cross-industry', '1235')
    rows = [
        # row 2: national total, outside the metro scope
        hr_row('00000', 'United States', '1', '000000', 'cross-industry', '1235',
               'total', '00-0000', 'National total', 99999999, 1, 1, 1, 1),
        # row 3: the unique selected metro total
        hr_row(*metro, 'total', '00-0000', 'All occupations', 1000000, 80000, 60000, 40, 30),
        # rows 4..8: five metro major groups
        hr_row(*metro, 'major', '11-0000', 'Management', 500, 1, 1, 1, 20),
        hr_row(*metro, 'major', '13-0000', 'Business', 400, 1, 1, 1, 20),
        hr_row(*metro, 'major', '15-0000', 'Computers', 300, 1, 1, 1, 'N'),
        hr_row(*metro, 'major', '17-0000', 'Architecture', 200, 1, 1, 1, 20),
        hr_row(*metro, 'major', '19-0000', 'Life sciences', 100, 1, 1, 1, 20),
        # row 9: suppressed major employment (excluded from ranking, counted)
        hr_row(*metro, 'major', '21-0000', 'Community service', 'N', 1, 1, 1, 20),
        # row 10: detailed row, never mixed with major groups
        hr_row(*metro, 'detailed', '11-1011', 'Detailed row', 999999, 1, 1, 1, 1),
        # row 11: sixth major group below the top five
        hr_row(*metro, 'major', '23-0000', 'Legal', 50, 1, 1, 1, 20),
        # row 12: different metro, outside the scope
        hr_row('31080', 'Los Angeles', '4', '000000', 'cross-industry', '1235',
               'total', '00-0000', 'LA total', 123, 1, 1, 1, 1),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = 'MSA_M2025_dl'
    ws.append(HR_HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)
    wb.close()
    return len(rows) + 1  # total rows including header


# Hand-computed retail fixture: eligible rows are 100/2x3 (Jan 2010, missing
# customer), 100/1x4 (Feb 2010, repeated invoice, distinct customer), 108/1x1.5
# (31 Dec 2010 as a numeric Excel serial, missing customer), 109/1x2.5
# (1 Jan 2011: full-period only, outside the 2010 monthly chart) and 110/3x2
# (Jun 2010).  Everything else is excluded for exactly one reason each.
RETAIL_ROWS = [
    ('100', 'S1', 'd', 2, datetime(2010, 1, 4), 3.0, None, 'UK'),
    ('100', 'S1', 'd', 1, datetime(2010, 2, 4), 4.0, 12, 'UK'),
    ('C100', 'S2', 'd', 5, datetime(2010, 2, 5), 2.0, 12, 'UK'),
    ('c200', 'S2', 'd', 1, datetime(2010, 2, 5), 2.0, 12, 'UK'),
    ('101', 'S3', 'd', -1, datetime(2010, 2, 5), 2.0, 12, 'UK'),
    ('102', 'S3', 'd', 1, datetime(2010, 2, 5), 0.0, 12, 'UK'),
    ('103', 'S4', 'd', 1, None, 2.0, 12, 'UK'),
    (None, 'S4', 'd', 1, datetime(2010, 2, 5), 2.0, 12, 'UK'),
    ('   ', 'S4', 'd', 1, datetime(2010, 3, 2), 10.0, 12, 'UK'),
    ('104', 'S4', 'd', 1, '2010-03-01', 2.0, 12, 'UK'),
    ('105', 'S4', 'd', None, datetime(2010, 3, 2), 2.0, 12, 'UK'),
    ('106', 'S4', 'd', '2', datetime(2010, 3, 2), 2.0, 12, 'UK'),
    ('107', 'S4', 'd', 1, datetime(2010, 3, 2), None, 12, 'UK'),
    ('108', 'S5', 'd', 1, to_excel(datetime(2010, 12, 31)), 1.5, None, 'UK'),
    ('109', 'S5', 'd', 1, datetime(2011, 1, 1), 2.5, 12, 'UK'),
    ('110', 'S5', 'd', 3, datetime(2010, 6, 15), 2.0, 12, 'UK'),
]
RETAIL_HAND_TOTALS = {
    'gross_eligible_sales_gbp': 20.0,   # 6 + 4 + 1.5 + 2.5 + 6
    'eligible_line_count': 5,
    'eligible_quantity_units': 8.0,     # 2 + 1 + 1 + 1 + 3
    'eligible_lines_missing_customer_id': 2,
    'exclusion_counts': {
        'cancelled_invoice': 2, 'missing_invoice': 2,
        'nonpositive_quantity': 1, 'nonpositive_unit_price': 1,
        'missing_or_invalid_quantity': 2, 'missing_or_invalid_unit_price': 1,
        'missing_or_invalid_invoice_date': 2,
    },
    'source_rows': 16,
}
RETAIL_HAND_MONTHLY = [6.0, 4.0, 0.0, 0.0, 0.0, 6.0, 0.0, 0.0, 0.0, 0.0, 0.0, 1.5]


def make_retail_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Online Retail'
    ws.append(RETAIL_HEADERS)
    for row in RETAIL_ROWS:
        ws.append(row)
    wb.save(path)
    wb.close()
    return len(RETAIL_ROWS) + 1


def run_preparation_cli(argv, root=None):
    return preparation.cli(argv, root=root)


class AccessSpy:
    """Count worksheet cell reads before wb.save() to prove no full-table scan."""

    def __init__(self):
        self.enabled = True
        self.rows = {}

    def note(self, ws, row):
        if self.enabled:
            self.rows.setdefault(ws.title, set()).add(row)

    def source_rows(self, title):
        return self.rows.get(title, set())


class PreparationCliTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)

    def test_help_has_no_side_effects(self):
        with patch('sys.argv', ['prepare_xlsx_selection.py', '--help']):
            with self.assertRaises(SystemExit) as caught:
                run_preparation_cli(['--help'])
        self.assertEqual(caught.exception.code, 0)
        self.assertFalse(list(self.root.rglob('*.json')))

    def test_unknown_dataset_is_rejected_before_writes(self):
        with self.assertRaises(SystemExit) as caught:
            run_preparation_cli(['--dataset', 'nope', '--input', 'x.xlsx',
                                 '--output', str(self.root / 'a.json')])
        self.assertEqual(caught.exception.code, 2)
        self.assertFalse(list(self.root.rglob('*.json')))

    def test_fixture_mode_requires_output_and_rejects_assets_paths(self):
        workbook = self.root / 'in.xlsx'
        make_retail_workbook(workbook)
        with self.assertRaises(SystemExit):
            run_preparation_cli(['--dataset', 'retail_online', '--input', str(workbook)])
        assets_out = self.root / 'assets/cases/retail_online/.../prepared_selection.json'
        assets_out.parent.mkdir(parents=True)
        with self.assertRaisesRegex(ValueError, 'assets'):
            run_preparation_cli(['--dataset', 'retail_online', '--input', str(workbook),
                                 '--output', str(assets_out), '--fixture'], root=self.root)
        self.assertFalse(assets_out.exists())

    def test_existing_artifact_is_protected_and_force_keeps_superseded(self):
        workbook = self.root / 'in.xlsx'
        make_retail_workbook(workbook)
        out = self.root / 'prepared_selection.json'
        run_preparation_cli(['--dataset', 'retail_online', '--input', str(workbook),
                             '--output', str(out), '--fixture'])
        first = json.loads(out.read_text())
        with self.assertRaises(FileExistsError):
            run_preparation_cli(['--dataset', 'retail_online', '--input', str(workbook),
                                 '--output', str(out), '--fixture'])
        self.assertEqual(json.loads(out.read_text()), first)
        run_preparation_cli(['--dataset', 'retail_online', '--input', str(workbook),
                             '--output', str(out), '--fixture', '--force'])
        self.assertTrue(list(self.root.glob('prepared_selection.superseded-*.json')))
        self.assertNotEqual(json.loads(out.read_text())['prepared_by']['script_sha256'],
                            None)

    def test_formal_mode_verifies_input_hash_and_expected_before_writing(self):
        root = self.root
        directory = preparation.asset_directory(root, 'retail_online')
        (directory / 'input').mkdir(parents=True)
        workbook = directory / 'input/Online Retail.xlsx'
        make_retail_workbook(workbook)
        expected = {
            'metric_values': {'gross_eligible_sales_gbp': 999.0,
                              'eligible_line_count': 5,
                              'eligible_quantity_units': 8.0},
            'chart_values': {'monthly_2010_sales': RETAIL_HAND_MONTHLY},
        }
        (directory / 'expected.json').write_text(json.dumps(expected))
        manifest = {
            'schema_version': 'office-dataset-v2', 'kind': 'xlsx',
            'dataset_id': 'retail_online', 'workbook': 'Online Retail.xlsx',
            'expected_file': 'expected.json',
            'input_files': {'input/Online Retail.xlsx': '0' * 64},
            'requirements': {}, 'origin_type': 'manual_workflow_adaptation',
        }
        (directory / 'manifest.json').write_text(json.dumps(manifest))
        # input SHA-256 disagrees with the manifest binding -> refuse, write nothing
        with self.assertRaisesRegex(ValueError, 'SHA-256'):
            run_preparation_cli(['--dataset', 'retail_online'], root=root)
        self.assertFalse((directory / 'input/prepared_selection.json').exists())
        # correct hash but the frozen oracle disagrees -> refuse, write nothing
        manifest['input_files']['input/Online Retail.xlsx'] = preparation._sha256(workbook)
        (directory / 'manifest.json').write_text(json.dumps(manifest))
        with self.assertRaisesRegex(ValueError, 'expected oracle'):
            run_preparation_cli(['--dataset', 'retail_online'], root=root)
        self.assertFalse((directory / 'input/prepared_selection.json').exists())
        # matching oracle -> publish artifact and bind it in the manifest
        expected['metric_values']['gross_eligible_sales_gbp'] = 20.0
        (directory / 'expected.json').write_text(json.dumps(expected))
        run_preparation_cli(['--dataset', 'retail_online'], root=root)
        artifact_path = directory / 'input/prepared_selection.json'
        self.assertTrue(artifact_path.is_file())
        updated = json.loads((directory / 'manifest.json').read_text())
        self.assertEqual(updated['input_files']['input/prepared_selection.json'],
                         preparation._sha256(artifact_path))
        self.assertEqual(updated['preparation']['workload_version'],
                         preparation.WORKLOAD_VERSION)
        declared = updated['requirements']['derived_sheets']['Derived Totals']
        self.assertEqual(declared['prepared_selection'], 'input/prepared_selection.json')
        self.assertEqual(declared['metric_anchors'], preparation.RETAIL_DERIVED_LAYOUT['metric_anchors'])
        self.assertEqual(declared['metric_value_paths'], preparation.RETAIL_DERIVED_LAYOUT['metric_value_paths'])
        self.assertEqual(declared['chart_series'], preparation.RETAIL_DERIVED_LAYOUT['chart_series'])
        self.assertEqual(declared['frozen_cells'], preparation.RETAIL_DERIVED_LAYOUT['frozen_cells'])
        self.assertEqual(declared['max_rows'], 40)
        self.assertEqual(updated['requirements']['diagnostics_bindings'],
                         preparation.RETAIL_DIAGNOSTICS_BINDINGS)


class HrPreparationTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)
        self.workbook = self.root / 'MSA_M2025_dl.xlsx'
        self.total_rows = make_hr_workbook(self.workbook)
        self.out = self.root / 'prepared_selection.json'

    def test_freezes_unique_total_and_top_five_majors(self):
        run_preparation_cli(['--dataset', 'hr_metro', '--input', str(self.workbook),
                             '--output', str(self.out), '--fixture'])
        artifact = json.loads(self.out.read_text())
        self.assertEqual(artifact['schema_version'], 'xlsx-prepared-selection-v1')
        self.assertEqual(artifact['dataset_id'], 'hr_metro')
        self.assertEqual(artifact['input_sha256'], preparation._sha256(self.workbook))
        selection = artifact['selection']
        self.assertEqual(selection['total']['row'], 3)
        self.assertEqual(selection['total']['occ_code'], '00-0000')
        self.assertEqual(selection['total']['kpis']['total_employment'],
                         {'cell': 'J3', 'value': 1000000})
        self.assertEqual(selection['total']['kpis']['annual_mean_wage'],
                         {'cell': 'K3', 'value': 80000})
        self.assertEqual(selection['total']['kpis']['hourly_median_wage'],
                         {'cell': 'N3', 'value': 30})
        self.assertEqual([row['occupation_code'] for row in selection['chart_rows']],
                         ['11-0000', '13-0000', '15-0000', '17-0000', '19-0000'])
        self.assertEqual([row['employment'] for row in selection['chart_rows']],
                         [500, 400, 300, 200, 100])
        self.assertTrue(all(row['occupation_title_cell'].startswith('I')
                            for row in selection['chart_rows']))
        self.assertEqual(selection['diagnostics']['major_group_count'], 7)
        self.assertEqual(selection['diagnostics']['major_group_suppressed_employment_count'], 1)
        self.assertEqual(selection['diagnostics']['suppression_marker_counts']['H_MEDIAN'], 1)

    def test_duplicate_or_missing_total_is_rejected(self):
        rows = load_workbook(self.workbook)['MSA_M2025_dl']
        wb = load_workbook(self.workbook)
        ws = wb['MSA_M2025_dl']
        ws.append(hr_row('35620', 'New York-Newark-Jersey City, NY-NJ', '4', '000000',
                         'cross-industry', '1235', 'total', '00-0000', 'Second total', 5, 5, 5, 5, 5))
        duplicate = self.root / 'duplicate.xlsx'
        wb.save(duplicate)
        wb.close()
        with self.assertRaisesRegex(ValueError, 'one selected MSA total row'):
            run_preparation_cli(['--dataset', 'hr_metro', '--input', str(duplicate),
                                 '--output', str(self.root / 'x.json'), '--fixture'])
        # fewer than five unsuppressed majors -> rejected
        headers = HR_HEADERS
        thin = [[
            '35620', 'New York-Newark-Jersey City, NY-NJ', '4', '000000',
            'cross-industry', '1235', 'total', '00-0000', 'All occupations',
            10, 10, 10, 10, 10]] + [
            ['35620', 'New York-Newark-Jersey City, NY-NJ', '4', '000000',
             'cross-industry', '1235', 'major', f'{i:02d}-0000', f'G{i}', i, 1, 1, 1, 1]
            for i in range(3)]
        with self.assertRaisesRegex(ValueError, 'five'):
            select_hr_metro(thin, headers)


class RetailPreparationTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)
        self.workbook = self.root / 'Online Retail.xlsx'
        self.total_rows = make_retail_workbook(self.workbook)
        self.out = self.root / 'prepared_selection.json'

    def test_hand_computed_full_period_and_monthly_ranges_are_separate(self):
        run_preparation_cli(['--dataset', 'retail_online', '--input', str(self.workbook),
                             '--output', str(self.out), '--fixture'])
        artifact = json.loads(self.out.read_text())
        totals = artifact['selection']['totals']
        for key, value in RETAIL_HAND_TOTALS.items():
            self.assertEqual(totals[key], value, key)
        # full-period KPIs include the 2011-01-01 row ...
        self.assertEqual(totals['gross_eligible_sales_gbp'], 20.0)
        # ... while the monthly chart covers 2010 only (Dec 31 in, Jan 1 out)
        self.assertEqual(artifact['selection']['monthly_2010_sales_gbp'], RETAIL_HAND_MONTHLY)
        self.assertEqual(artifact['selection']['columns']['invoicedate']['column'], 'E')
        self.assertEqual(artifact['selection']['columns']['unitprice']['header'], 'UnitPrice')
        self.assertEqual(artifact['selection']['derived_layout'],
                         preparation.RETAIL_DERIVED_LAYOUT)

    def test_oracle_cross_check_detects_wrong_values(self):
        artifact = preparation.build_artifact('retail_online', self.workbook)
        wrong = {'metric_values': {'gross_eligible_sales_gbp': 21.0,
                                   'eligible_line_count': 5,
                                   'eligible_quantity_units': 8.0},
                 'chart_values': {'monthly_2010_sales': RETAIL_HAND_MONTHLY}}
        with self.assertRaisesRegex(ValueError, 'gross_eligible_sales_gbp'):
            preparation.check_against_expected('retail_online', artifact['selection'], wrong)
        wrong_monthly = list(RETAIL_HAND_MONTHLY)
        wrong_monthly[0] = 99.0
        wrong = {'metric_values': RETAIL_HAND_TOTALS,
                 'chart_values': {'monthly_2010_sales': wrong_monthly}}
        with self.assertRaisesRegex(ValueError, 'monthly'):
            preparation.check_against_expected('retail_online', artifact['selection'], wrong)


def _hr_workspace(root, name='workspace'):
    workspace = root / name
    (workspace / 'input').mkdir(parents=True)
    (workspace / 'output').mkdir()
    workbook = workspace / 'input/MSA_M2025_dl.xlsx'
    total_rows = make_hr_workbook(workbook)
    shutil.copyfile(workbook, workspace / 'output' / HR_TARGET)
    run_preparation_cli(['--dataset', 'hr_metro', '--input', str(workbook),
                         '--output', str(workspace / 'input/prepared_selection.json'),
                         '--fixture'])
    return workspace, total_rows


def _retail_workspace(root, name='workspace'):
    workspace = root / name
    (workspace / 'input').mkdir(parents=True)
    (workspace / 'output').mkdir()
    workbook = workspace / 'input/Online Retail.xlsx'
    total_rows = make_retail_workbook(workbook)
    shutil.copyfile(workbook, workspace / 'output' / RETAIL_TARGET)
    run_preparation_cli(['--dataset', 'retail_online', '--input', str(workbook),
                         '--output', str(workspace / 'input/prepared_selection.json'),
                         '--fixture'])
    return workspace, total_rows


class HelperBindingRejectionTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)

    def test_hr_helper_rejects_missing_preparation(self):
        workspace, _ = _hr_workspace(self.root)
        (workspace / 'input/prepared_selection.json').unlink()
        with self.assertRaisesRegex(FileNotFoundError, 'preparation artifact missing'):
            hr_build(workspace / 'output' / HR_TARGET)

    def test_hr_helper_rejects_stale_or_misbound_preparation(self):
        workspace, _ = _hr_workspace(self.root)
        artifact_path = workspace / 'input/prepared_selection.json'
        artifact = json.loads(artifact_path.read_text())
        artifact['schema_version'] = 'xlsx-prepared-selection-v0'
        artifact_path.write_text(json.dumps(artifact))
        with self.assertRaisesRegex(ValueError, 'schema'):
            hr_build(workspace / 'output' / HR_TARGET)
        artifact = json.loads(artifact_path.read_text())
        artifact['schema_version'] = 'xlsx-prepared-selection-v1'
        artifact['workload_version'] = 'manual-workflow-adaptation-v1'
        artifact_path.write_text(json.dumps(artifact))
        with self.assertRaisesRegex(ValueError, 'workload version'):
            hr_build(workspace / 'output' / HR_TARGET)

    def test_helpers_do_not_hash_the_complete_workbook(self):
        """The full-file hash binding is preflight work, not per-iteration work.

        A sha256 spy records every byte fed to any hash during a real build;
        neither helper may feed the workbook's bytes through a digest.
        """
        import hashlib
        real_sha256 = hashlib.sha256
        fed = []

        class SpyDigest:
            def __init__(self, data=b''):
                self._hash = real_sha256()
                self.update(data)

            def update(self, data):
                fed.append(bytes(data))
                self._hash.update(data)
                return self

            def hexdigest(self):
                return self._hash.hexdigest()

        for make_workspace, build, target_name, sheet in (
                (_hr_workspace, hr_build, HR_TARGET, 'MSA_M2025_dl'),
                (_retail_workspace, retail_build, RETAIL_TARGET, 'Online Retail')):
            with self.subTest(sheet=sheet):
                workspace, _rows = make_workspace(self.root, 'ws-' + uuid.uuid4().hex[:8])
                target = workspace / 'output' / target_name
                workbook_bytes = target.read_bytes()
                fed.clear()
                with patch('hashlib.sha256', SpyDigest):
                    build(target)
                total = sum(len(chunk) for chunk in fed)
                self.assertLess(total, len(workbook_bytes),
                                'a complete workbook was hashed inside the helper')
                self.assertFalse(any(chunk == workbook_bytes for chunk in fed))

    def test_hr_helper_rejects_wrong_recorded_positions(self):
        workspace, _ = _hr_workspace(self.root)
        artifact_path = workspace / 'input/prepared_selection.json'
        artifact = json.loads(artifact_path.read_text())
        # point the total KPI at a different row with different values
        artifact['selection']['total']['kpis']['total_employment']['cell'] = 'J4'
        artifact_path.write_text(json.dumps(artifact))
        with self.assertRaisesRegex(ValueError, 'does not match this workbook'):
            hr_build(workspace / 'output' / HR_TARGET)

    def test_retail_helper_rejects_missing_stale_and_misbound_preparation(self):
        workspace, _ = _retail_workspace(self.root)
        target = workspace / 'output' / RETAIL_TARGET
        (workspace / 'input/prepared_selection.json').unlink()
        with self.assertRaisesRegex(FileNotFoundError, 'preparation artifact missing'):
            retail_build(target)
        run_preparation_cli(['--dataset', 'retail_online',
                             '--input', str(workspace / 'input/Online Retail.xlsx'),
                             '--output', str(workspace / 'input/prepared_selection.json'),
                             '--fixture'])
        artifact_path = workspace / 'input/prepared_selection.json'
        artifact = json.loads(artifact_path.read_text())
        artifact['dataset_id'] = 'hr_metro'
        artifact_path.write_text(json.dumps(artifact))
        with self.assertRaisesRegex(ValueError, 'dataset'):
            retail_build(target)
        artifact = json.loads(artifact_path.read_text())
        artifact['dataset_id'] = 'retail_online'
        artifact['selection']['columns']['quantity']['column'] = 'H'
        artifact_path.write_text(json.dumps(artifact))
        with self.assertRaisesRegex(ValueError, 'header position'):
            retail_build(target)


class HelperBehaviourTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)

    def _spy_build(self, build, target, source_sheet):
        spy = AccessSpy()
        real_get_cell = Worksheet._get_cell
        real_save = OpenpyxlWorkbook.save

        def spied_get_cell(ws_self, row, column):
            spy.note(ws_self, row)
            return real_get_cell(ws_self, row, column)

        def spied_save(wb_self, output):
            spy.enabled = False
            try:
                return real_save(wb_self, output)
            finally:
                spy.enabled = True

        with patch.object(Worksheet, '_get_cell', spied_get_cell), \
                patch.object(OpenpyxlWorkbook, 'save', spied_save):
            build(target)
        return spy

    def test_hr_helper_reads_frozen_positions_not_the_whole_table(self):
        workspace, total_rows = _hr_workspace(self.root)
        spy = self._spy_build(hr_build, workspace / 'output' / HR_TARGET, 'MSA_M2025_dl')
        touched = spy.source_rows('MSA_M2025_dl')
        # header row + the frozen total/chart rows only; the full table has
        # total_rows rows and the previous helper touched every one of them.
        self.assertLessEqual(len(touched), 10)
        self.assertIn(1, touched)
        self.assertIn(3, touched)
        self.assertTrue(touched.issubset(set(range(1, total_rows + 1))))

    def test_retail_helper_reads_only_the_header_row(self):
        workspace, total_rows = _retail_workspace(self.root)
        spy = self._spy_build(retail_build, workspace / 'output' / RETAIL_TARGET, 'Online Retail')
        touched = spy.source_rows('Online Retail')
        self.assertLessEqual(len(touched), 2)
        self.assertTrue(touched.issubset({1}))

    def test_hr_helper_output_references_and_preserved_source(self):
        workspace, total_rows = _hr_workspace(self.root)
        hr_build(workspace / 'output' / HR_TARGET)
        book = load_workbook(workspace / 'output' / HR_TARGET)
        try:
            self.assertEqual(book['MSA_M2025_dl'].max_row, total_rows)
            self.assertEqual(book['Metro Summary']['B2'].value, "='MSA_M2025_dl'!J3")
            self.assertEqual(book['Metro Summary']['B3'].value, "='MSA_M2025_dl'!K3")
            self.assertEqual(book['Metro Summary']['B9'].value, '=IF(B6="Base",B3*B7,B3*B8)')
            self.assertEqual(book['Chart Data']['A2'].value, "='MSA_M2025_dl'!I4")
            self.assertEqual(book['Chart Data']['B6'].value, "='MSA_M2025_dl'!J8")
            features = json.loads((workspace / 'output/workbook_features.json').read_text())
            self.assertEqual(features['metrics']['total_employment'],
                             {'sheet': 'Metro Summary', 'cell': 'B2'})
            # every formula references at most a few explicit cells, never a
            # whole-table range
            for sheet in ('Metro Summary', 'Chart Data'):
                for row in book[sheet].iter_rows():
                    for cell in row:
                        if cell.data_type == 'f':
                            for token in str(cell.value).split('!'):
                                self.assertNotIn(':', token.replace('B3*B7', ''))
        finally:
            book.close()
        diagnostics = json.loads(
            (workspace / 'output/xlsx_business_diagnostics.json').read_text())
        self.assertEqual(diagnostics['total_row'], 3)
        self.assertEqual(diagnostics['selected_chart_rows'],
                         ['11-0000', '13-0000', '15-0000', '17-0000', '19-0000'])
        self.assertEqual(diagnostics['selection_source'],
                         'frozen preparation artifact (input/prepared_selection.json); '
                         'no in-loop table rescan')

    def test_retail_helper_writes_small_derived_sheet_and_cross_references(self):
        workspace, total_rows = _retail_workspace(self.root)
        retail_build(workspace / 'output' / RETAIL_TARGET)
        book = load_workbook(workspace / 'output' / RETAIL_TARGET, data_only=False)
        try:
            self.assertEqual(book['Online Retail'].max_row, total_rows)
            derived = book['Derived Totals']
            self.assertEqual(derived['A1'].value,
                             'Derived transaction totals (frozen preparation data)')
            self.assertEqual(derived['B4'].value, 20.0)
            self.assertEqual(derived['B5'].value, 5)
            self.assertEqual(derived['B6'].value, 8.0)
            self.assertEqual(derived['B20'].value, 6.0)
            self.assertEqual(derived['B31'].value, 1.5)
            self.assertLessEqual(derived.max_row, 40)
            # the derived sheet holds frozen values only, no formulas at all
            derived_formulas = [c.coordinate for row in derived.iter_rows()
                                for c in row if c.data_type == 'f']
            self.assertEqual(derived_formulas, [])
            summary = book['Sales Overview']
            self.assertEqual(summary['B2'].value, "='Derived Totals'!B4")
            self.assertEqual(summary['B3'].value, "='Derived Totals'!B5")
            self.assertEqual(summary['B4'].value, "='Derived Totals'!B6")
            self.assertEqual(summary['B10'].value,
                             '=IF(B6="Base",B2*B7,IF(B6="Higher",B2*B8,B2*B9))')
            chart = book['Sales Overview']._charts[0]
            series = chart.series[0]
            self.assertEqual(series.val.numRef.f.replace('$', ''),
                             "'Derived Totals'!B20:B31")
            category = series.cat.strRef or series.cat.numRef
            self.assertEqual(category.f.replace('$', ''),
                             "'Derived Totals'!A20:A31")
            # no whole-table array formulas anywhere in the added sheets
            for sheet in ('Sales Overview', 'Derived Totals'):
                for row in book[sheet].iter_rows():
                    for cell in row:
                        if cell.data_type == 'f':
                            formula = str(cell.value)
                            self.assertNotIn('SUMPRODUCT', formula.upper())
                            for part in formula.split('!'):
                                self.assertNotIn('$A$2:$A$', part)
        finally:
            book.close()
        features = json.loads((workspace / 'output/workbook_features.json').read_text())
        self.assertEqual(features['derived_sheet'], 'Derived Totals')
        diagnostics = json.loads(
            (workspace / 'output/xlsx_business_diagnostics.json').read_text())
        self.assertEqual(diagnostics['gross_eligible_sales_gbp'], 20.0)
        self.assertEqual(diagnostics['eligible_line_count'], 5)
        self.assertEqual(diagnostics['currency'], 'GBP')
        self.assertEqual(diagnostics['exclusion_counts']['cancelled_invoice'], 2)


class _CacheInjector:
    """Fixture-only cached formula values (not a LibreOffice recalculation)."""

    @staticmethod
    def set_formula_caches(path, values):
        from zipfile import ZIP_DEFLATED, ZipFile
        import xml.etree.ElementTree as ET
        ns_main = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
                   'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        ET.register_namespace('', ns_main['m'])
        ET.register_namespace('r', ns_main['r'])
        with ZipFile(path) as source:
            workbook = ET.fromstring(source.read('xl/workbook.xml'))
            rels = ET.fromstring(source.read('xl/_rels/workbook.xml.rels'))
            relation_targets = {node.attrib['Id']: node.attrib['Target'] for node in rels}
            sheet_paths = {}
            for node in workbook.findall('m:sheets/m:sheet', ns_main):
                target = relation_targets[node.attrib['{' + ns_main['r'] + '}id']].lstrip('/')
                sheet_paths[node.attrib['name']] = target if target.startswith('xl/') else 'xl/' + target
            patches = {}
            for (sheet_name, address), value in values.items():
                sheet_path = sheet_paths[sheet_name]
                root = patches.get(sheet_path)
                if root is None:
                    root = ET.fromstring(source.read(sheet_path))
                    patches[sheet_path] = root
                cell = next((item for item in root.findall('.//m:c', ns_main)
                             if item.attrib.get('r') == address and
                             item.find('m:f', ns_main) is not None), None)
                if cell is None:
                    raise AssertionError(f'formula cache target missing: {sheet_name}!{address}')
                cached = cell.find('m:v', ns_main)
                if cached is None:
                    cached = ET.SubElement(cell, '{%s}v' % ns_main['m'])
                cached.text = str(value)
                if isinstance(value, str):
                    cell.set('t', 'str')
            patches = {name: ET.tostring(root, encoding='utf-8', xml_declaration=True)
                       for name, root in patches.items()}
            temp = path.with_suffix('.cache-fixture.xlsx')
            with ZipFile(temp, 'w', ZIP_DEFLATED) as target:
                for item in source.infolist():
                    target.writestr(item, patches.get(item.filename, source.read(item.filename)))
        temp.replace(path)


class RetailVerifierTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)

    def _fixture(self, *, tamper_derived=None, hardcode_kpi=False,
                 kpi_wrong_reference=False, disconnect_chart=False,
                 formula_in_derived=False, tamper_month_date=False,
                 swap_month_order=False, tamper_diagnostics_json=None,
                 tamper_category_reference=False):
        workspace, _ = _retail_workspace(self.root, 'ws-' + uuid.uuid4().hex[:8])
        retail_build(workspace / 'output' / RETAIL_TARGET)
        output = workspace / 'output' / RETAIL_TARGET
        book = load_workbook(output)
        if tamper_derived:
            book['Derived Totals'][tamper_derived] = 999.0
        if hardcode_kpi:
            book['Sales Overview']['B2'] = 20.0
        if kpi_wrong_reference:
            book['Sales Overview']['B2'] = "='Derived Totals'!B9"
        if disconnect_chart:
            chart = book['Sales Overview']._charts[0]
            chart.series[0].val.numRef.f = "'Derived Totals'!B4:B15"
        if tamper_category_reference:
            chart = book['Sales Overview']._charts[0]
            chart.series[0].cat.numRef.f = "'Derived Totals'!A2:A13"
        if formula_in_derived:
            book['Derived Totals']['B4'] = '=1*20'
        if tamper_month_date:
            # wrong month in the first category cell (December instead of January)
            book['Derived Totals']['A20'] = date(2010, 12, 1)
        if swap_month_order:
            # swap two month labels: right dates, wrong order
            first, second = book['Derived Totals']['A20'].value, book['Derived Totals']['A21'].value
            book['Derived Totals']['A20'] = second
            book['Derived Totals']['A21'] = first
        book.save(output)
        book.close()
        values = {'gross_eligible_sales_gbp': 20.0, 'eligible_line_count': 5,
                  'eligible_quantity_units': 8.0}
        caches = {
            ('Sales Overview', 'B3'): values['eligible_line_count'],
            ('Sales Overview', 'B4'): values['eligible_quantity_units'],
            ('Sales Overview', 'B10'): values['gross_eligible_sales_gbp'],
        }
        if not hardcode_kpi:
            caches[('Sales Overview', 'B2')] = values['gross_eligible_sales_gbp']
        _CacheInjector.set_formula_caches(output, caches)
        output_dir = workspace / 'output'
        (output_dir / 'formula_recalc.json').write_text(json.dumps(
            {'status': 'success', 'total_errors': 0, 'total_formulas': 4}))
        (output_dir / 'xlsx_enhancement_summary.json').write_text('{}\n')
        (output_dir / 'metrics.csv').write_text(
            'metric,value\n' + ''.join(f'{k},{v}\n' for k, v in values.items()))
        (output_dir / 'scenario.csv').write_text('metric,value\nprojected_sales_gbp,20.0\n')
        # business diagnostics JSON is a formal deliverable; tampering it must fail
        if tamper_diagnostics_json:
            diagnostics = json.loads((output_dir / 'xlsx_business_diagnostics.json').read_text())
            diagnostics[tamper_diagnostics_json] = 999999
            (output_dir / 'xlsx_business_diagnostics.json').write_text(json.dumps(diagnostics))
        ws = load_workbook(workspace / 'input/Online Retail.xlsx', read_only=True).worksheets[0]
        dimensions = [ws.max_row, ws.max_column]
        ws.parent.close()
        manifest = {
            'dataset_id': 'retail_online', 'kind': 'xlsx',
            'workbook': 'Online Retail.xlsx', 'output_workbook': RETAIL_TARGET,
            'required_outputs': [RETAIL_TARGET, 'workbook_features.json',
                                 'formula_recalc.json', 'business_verification.json',
                                 'xlsx_enhancement_summary.json', 'metrics.csv',
                                 'scenario.csv', 'xlsx_business_diagnostics.json'],
            'publish_files': [], 'requirements': {
                'base_sheets': ['Online Retail'],
                'base_dimensions': {'Online Retail': dimensions},
                'base_dimension_details': {'Online Retail': {
                    'declared_dimension': dimensions, 'serialized_cell_range': dimensions,
                    'effective_data_range': dimensions, 'effective_range_method': 'fixture'}},
                'empty_strings_policy': {'sheets': []},
                'raw_sheets': ['Online Retail'], 'summary_sheet': 'Sales Overview',
                'metrics': [{'id': key} for key in values],
                'scenario': {
                    'default': 'Base',
                    'rows': [
                        {'name': 'Base', 'assumptions': {'sales_multiplier': 1.0}},
                        {'name': 'Higher', 'assumptions': {'sales_multiplier': 1.05}},
                        {'name': 'Lower', 'assumptions': {'sales_multiplier': 0.95}},
                    ],
                    'projections': [{'id': 'projected_sales_gbp'}]},
                'charts': [{'id': 'monthly_2010_sales'}],
                'minimum_formula_count': 4, 'minimum_comments': 2,
                'minimum_conditional_formats': 2,
                'derived_sheets': {
                    'Derived Totals': {
                        'provenance': 'fixture frozen preparation artifact',
                        'prepared_selection': 'input/prepared_selection.json',
                        'metric_anchors': preparation.RETAIL_DERIVED_LAYOUT['metric_anchors'],
                        'metric_value_paths': preparation.RETAIL_DERIVED_LAYOUT['metric_value_paths'],
                        'chart_series': preparation.RETAIL_DERIVED_LAYOUT['chart_series'],
                        'frozen_cells': preparation.RETAIL_DERIVED_LAYOUT['frozen_cells'],
                        'max_rows': 40,
                    },
                },
                'diagnostics_bindings': preparation.RETAIL_DIAGNOSTICS_BINDINGS,
            },
        }
        expected = {
            'metric_values': values,
            'scenario_values': {'Base': {'projected_sales_gbp': 20.0}},
            'chart_values': {'monthly_2010_sales': RETAIL_HAND_MONTHLY},
        }
        return workspace, manifest, expected

    def test_derived_layout_passes_loop_and_independent_verification(self):
        workspace, manifest, expected = self._fixture()
        loop = Checks(manifest)
        verify_semantic_xlsx_loop(manifest, expected, workspace, loop,
                                  deferred_outputs=('xlsx_enhancement_summary.json',))
        self.assertTrue(loop.data['checks']['derived:sheet_present:Derived Totals']['ok'])
        self.assertTrue(loop.data['checks']['derived:no_formulas:Derived Totals']['ok'])
        self.assertTrue(loop.data['checks']['derived:anchor_value:Derived Totals!B4']['ok'])
        self.assertTrue(loop.data['checks']['derived:anchor_source_value:Derived Totals!B4']['ok'])
        self.assertTrue(loop.data['checks']['derived:prepared_selection:Derived Totals']['ok'])
        self.assertTrue(loop.data['checks']['derived:chart_categories:Derived Totals:monthly_2010_sales']['ok'])
        self.assertTrue(loop.data['checks']['derived:chart_source_values:Derived Totals:monthly_2010_sales']['ok'])
        self.assertTrue(loop.data['checks']['derived:frozen_cell:Derived Totals!B9']['ok'])
        self.assertTrue(loop.data['checks']['derived:diagnostics:eligible_line_count']['ok'])
        self.assertTrue(loop.data['checks']['semantic:source_dependency:gross_eligible_sales_gbp']['ok'])
        self.assertTrue(loop.data['checks']['derived:chart_series:monthly_2010_sales']['ok'])
        self.assertTrue(loop.data['checks']['semantic:chart:monthly_2010_sales']['ok'])
        self.assertEqual(loop.data['failures'], [], loop.data['checks'])
        independent = Checks(manifest)
        verify_semantic_xlsx(manifest, expected, workspace, independent)
        self.assertTrue(independent.data['checks']['semantic:source_dependency:gross_eligible_sales_gbp']['ok'])
        self.assertTrue(independent.data['checks']['semantic:chart_dependency:monthly_2010_sales']['ok'])
        self.assertEqual(independent.data['failures'], [], independent.data['checks'])

    def test_wrong_month_date_and_swapped_month_order_fail(self):
        for scenario, check in (('date', 'derived:chart_categories:Derived Totals:monthly_2010_sales'),
                                ('order', 'derived:chart_categories:Derived Totals:monthly_2010_sales')):
            with self.subTest(scenario=scenario):
                kwargs = {'tamper_month_date': True} if scenario == 'date' else {'swap_month_order': True}
                workspace, manifest, expected = self._fixture(**kwargs)
                for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
                    checks = Checks(manifest)
                    scope_verifier(manifest, expected, workspace, checks)
                    self.assertFalse(checks.data['checks'][check]['ok'],
                                     checks.data['checks'][check])

    def test_tampered_diagnostics_cells_fail(self):
        # B9 missing CustomerID count, B11 cancellation count, B18 currency
        for cell, check in (('B9', 'derived:frozen_cell:Derived Totals!B9'),
                            ('B11', 'derived:frozen_cell:Derived Totals!B11'),
                            ('B18', 'derived:frozen_cell:Derived Totals!B18')):
            with self.subTest(cell=cell):
                workspace, manifest, expected = self._fixture(tamper_derived=cell)
                for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
                    checks = Checks(manifest)
                    scope_verifier(manifest, expected, workspace, checks)
                    self.assertFalse(checks.data['checks'][check]['ok'],
                                     checks.data['checks'][check])

    def test_tampered_diagnostics_json_fails(self):
        for key in ('eligible_line_count', 'gross_eligible_sales_gbp',
                    'eligible_lines_missing_customer_id'):
            with self.subTest(key=key):
                workspace, manifest, expected = self._fixture(tamper_diagnostics_json=key)
                for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
                    checks = Checks(manifest)
                    scope_verifier(manifest, expected, workspace, checks)
                    self.assertFalse(checks.data['checks']['derived:diagnostics:' + key]['ok'],
                                     checks.data['checks']['derived:diagnostics:' + key])

    def test_category_reference_disconnect_fails(self):
        workspace, manifest, expected = self._fixture(tamper_category_reference=True)
        for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
            checks = Checks(manifest)
            scope_verifier(manifest, expected, workspace, checks)
            self.assertFalse(checks.data['checks']
                             ['derived:chart_categories_reference:monthly_2010_sales']['ok'],
                             checks.data['checks']
                             ['derived:chart_categories_reference:monthly_2010_sales'])

    def test_tampered_derived_value_fails(self):
        workspace, manifest, expected = self._fixture(tamper_derived='B4')
        for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
            checks = Checks(manifest)
            scope_verifier(manifest, expected, workspace, checks)
            self.assertFalse(checks.data['checks']['derived:anchor_value:Derived Totals!B4']['ok'])

    def test_hardcoded_kpi_fails(self):
        workspace, manifest, expected = self._fixture(hardcode_kpi=True)
        loop = Checks(manifest)
        verify_semantic_xlsx_loop(manifest, expected, workspace, loop,
                                  deferred_outputs=('xlsx_enhancement_summary.json',))
        self.assertFalse(loop.data['checks']['feature:formula:Sales Overview!B2']['ok'])

    def test_kpi_referencing_wrong_cell_fails(self):
        workspace, manifest, expected = self._fixture(kpi_wrong_reference=True)
        loop = Checks(manifest)
        verify_semantic_xlsx_loop(manifest, expected, workspace, loop,
                                  deferred_outputs=('xlsx_enhancement_summary.json',))
        self.assertFalse(loop.data['checks']['semantic:source_dependency:gross_eligible_sales_gbp']['ok'])
        self.assertTrue(loop.data['checks']['derived:anchor_value:Derived Totals!B4']['ok'])

    def test_disconnected_chart_series_fails(self):
        workspace, manifest, expected = self._fixture(disconnect_chart=True)
        loop = Checks(manifest)
        verify_semantic_xlsx_loop(manifest, expected, workspace, loop,
                                  deferred_outputs=('xlsx_enhancement_summary.json',))
        self.assertFalse(loop.data['checks']['derived:chart_series:monthly_2010_sales']['ok'])

    def test_formula_inside_derived_sheet_fails(self):
        workspace, manifest, expected = self._fixture(formula_in_derived=True)
        for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
            checks = Checks(manifest)
            scope_verifier(manifest, expected, workspace, checks)
            self.assertFalse(checks.data['checks']['derived:no_formulas:Derived Totals']['ok'])


class HrVerifierTests(unittest.TestCase):
    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.root = Path(self._temp.name)
        self.addCleanup(self._temp.cleanup)

    def _fixture(self, *, wrong_kpi_reference=False, tamper_diagnostics_json=None):
        workspace, total_rows = _hr_workspace(self.root, 'ws-' + uuid.uuid4().hex[:8])
        hr_build(workspace / 'output' / HR_TARGET)
        output = workspace / 'output' / HR_TARGET
        # A wrong KPI reference recalculates (in a real run) to the referenced
        # cell's actual value, so the injected cache mirrors that behaviour.
        total_cache = 500 if wrong_kpi_reference else 1000000
        if wrong_kpi_reference:
            book = load_workbook(output)
            book['Metro Summary']['B2'] = "='MSA_M2025_dl'!J4"  # wrong row
            book.save(output)
            book.close()
        _CacheInjector.set_formula_caches(output, {
            ('Metro Summary', 'B2'): total_cache,
            ('Metro Summary', 'B3'): 80000,
            ('Metro Summary', 'B4'): 30,
            ('Metro Summary', 'B9'): 80000,
            **{('Chart Data', f'A{row}'): title for row, title in enumerate(
                ['Management', 'Business', 'Computers', 'Architecture', 'Life sciences'], 2)},
            **{('Chart Data', f'B{row}'): value for row, value in enumerate(
                [500, 400, 300, 200, 100], 2)},
        })
        output_dir = workspace / 'output'
        (output_dir / 'formula_recalc.json').write_text(json.dumps(
            {'status': 'success', 'total_errors': 0, 'total_formulas': 14}))
        (output_dir / 'xlsx_enhancement_summary.json').write_text('{}\n')
        (output_dir / 'metrics.csv').write_text(
            'metric,value\ntotal_employment,1000000\nannual_mean_wage,80000\nhourly_median_wage,30\n')
        (output_dir / 'scenario.csv').write_text('metric,value\nprojected_mean_wage,80000\n')
        if tamper_diagnostics_json:
            diagnostics = json.loads((output_dir / 'xlsx_business_diagnostics.json').read_text())
            diagnostics[tamper_diagnostics_json] = 999999
            (output_dir / 'xlsx_business_diagnostics.json').write_text(json.dumps(diagnostics))
        ws = load_workbook(workspace / 'input/MSA_M2025_dl.xlsx', read_only=True).worksheets[0]
        dimensions = [ws.max_row, ws.max_column]
        ws.parent.close()
        manifest = {
            'dataset_id': 'hr_metro_fixture', 'kind': 'xlsx',
            'workbook': 'MSA_M2025_dl.xlsx', 'output_workbook': HR_TARGET,
            'required_outputs': [HR_TARGET, 'workbook_features.json',
                                 'formula_recalc.json', 'business_verification.json',
                                 'xlsx_enhancement_summary.json', 'metrics.csv',
                                 'scenario.csv', 'xlsx_business_diagnostics.json'],
            'publish_files': [], 'requirements': {
                'base_sheets': ['MSA_M2025_dl'],
                'base_dimensions': {'MSA_M2025_dl': dimensions},
                'base_dimension_details': {'MSA_M2025_dl': {
                    'declared_dimension': dimensions, 'serialized_cell_range': dimensions,
                    'effective_data_range': dimensions, 'effective_range_method': 'fixture'}},
                'empty_strings_policy': {'sheets': []},
                'raw_sheets': ['MSA_M2025_dl'], 'summary_sheet': 'Metro Summary',
                'metrics': [{'id': 'total_employment'}, {'id': 'annual_mean_wage'},
                            {'id': 'hourly_median_wage'}],
                'scenario': {
                    'default': 'Base',
                    'rows': [{'name': 'Base', 'assumptions': {'wage_multiplier': 1.0}},
                             {'name': 'Increase', 'assumptions': {'wage_multiplier': 1.05}}],
                    'projections': [{'id': 'projected_mean_wage'}]},
                'charts': [{'id': 'major_employment'}],
                'minimum_formula_count': 8, 'minimum_comments': 2,
                'minimum_conditional_formats': 2,
                'diagnostics_bindings': preparation.HR_DIAGNOSTICS_BINDINGS,
            },
        }
        expected = {
            'metric_values': {'total_employment': 1000000, 'annual_mean_wage': 80000,
                              'hourly_median_wage': 30},
            'scenario_values': {'Base': {'projected_mean_wage': 80000}},
            'chart_values': {'major_employment': [500, 400, 300, 200, 100]},
        }
        return workspace, manifest, expected

    def test_frozen_positions_pass_loop_and_independent(self):
        workspace, manifest, expected = self._fixture()
        loop = Checks(manifest)
        verify_semantic_xlsx_loop(manifest, expected, workspace, loop,
                                  deferred_outputs=('xlsx_enhancement_summary.json',))
        self.assertTrue(loop.data['checks']['derived:diagnostics:total_row']['ok'])
        self.assertTrue(loop.data['checks']['derived:diagnostics:selected_chart_rows']['ok'])
        self.assertEqual(loop.data['failures'], [], loop.data['checks'])
        independent = Checks(manifest)
        verify_semantic_xlsx(manifest, expected, workspace, independent)
        self.assertEqual(independent.data['failures'], [], independent.data['checks'])

    def test_tampered_hr_diagnostics_json_fails(self):
        for key in ('total_row', 'area_code', 'selected_chart_rows'):
            with self.subTest(key=key):
                workspace, manifest, expected = self._fixture(tamper_diagnostics_json=key)
                for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
                    checks = Checks(manifest)
                    scope_verifier(manifest, expected, workspace, checks)
                    self.assertFalse(checks.data['checks']['derived:diagnostics:' + key]['ok'],
                                     checks.data['checks']['derived:diagnostics:' + key])

    def test_kpi_referencing_wrong_original_cell_fails(self):
        workspace, manifest, expected = self._fixture(wrong_kpi_reference=True)
        for scope_verifier in (verify_semantic_xlsx_loop, verify_semantic_xlsx):
            checks = Checks(manifest)
            scope_verifier(manifest, expected, workspace, checks)
            self.assertFalse(checks.data['checks']['semantic:metric:total_employment']['ok'])


class CompileBindingTests(unittest.TestCase):
    def test_compile_copies_prepared_selection_only_when_manifest_binds_it(self):
        import scripts.compile_compatible as compiler
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for dataset in ('hr_metro', 'retail_online'):
                (root / 'compatibility/prompts').mkdir(parents=True, exist_ok=True)
                (root / 'compatibility/prompts/xlsx.txt').write_text('prompt')
                manifest_dir = preparation.asset_directory(root, dataset)
                (manifest_dir / 'input').mkdir(parents=True)
                workbook = manifest_dir / 'input/data.xlsx'
                workbook.write_bytes(b'fixture input')
                for name in ('expected.json', 'complexity.json', 'content-summary.json'):
                    (manifest_dir / name).write_text('{}\n')
                provenance = manifest_dir / 'provenance.json'
                provenance.write_text(json.dumps({'references': []}))
                manifest = {
                    'origin_type': 'manual_workflow_adaptation', 'agent_generated': False,
                    'requires_llm': False, 'workbook': 'data.xlsx',
                    'prompt_template': 'xlsx.txt',
                    'prompt_template_sha256': compiler.digest(root / 'compatibility/prompts/xlsx.txt'),
                    'runtime_case': 'FIXTURE', 'kind': 'xlsx',
                    'loop_verification_contract': 'document-oe-compatible-xlsx-loop-v1',
                    'input_files': {'input/data.xlsx': compiler.digest(workbook)},
                }
                if dataset == 'retail_online':
                    artifact = manifest_dir / 'input/prepared_selection.json'
                    artifact.write_text(json.dumps({'schema_version': 'v1'}))
                    manifest['input_files']['input/prepared_selection.json'] = compiler.digest(artifact)
                (manifest_dir / 'manifest.json').write_text(json.dumps(manifest))
                helper = root / 'compatibility/helpers' / f'{dataset}_build.py'
                helper.parent.mkdir(parents=True, exist_ok=True)
                helper.write_text('TARGET = "fixture.xlsx"\n')
                template = root / 'compatibility/cases' / ('hr' if dataset == 'hr_metro' else 'retail') / 'source'
                template.mkdir(parents=True)
            index = {'contract_sha256': 'fixture-contract', 'entries': [
                {'dataset_id': 'hr', 'recipe_path': 'other.json'}]}
            with patch.object(compiler, '_copy_current_verifier_package'):
                for dataset in ('hr_metro', 'retail_online'):
                    compiler._compile_manual_xlsx(root, dataset, index)
            # retail manifest binds the artifact: seed and source copies exist
            retail_seed = root / 'compatibility/cases/retail_online/seed/input'
            retail_source = root / 'compatibility/cases/retail_online/source/datasets/xlsx/retail_online/input'
            self.assertTrue((retail_seed / 'prepared_selection.json').is_file())
            self.assertTrue((retail_source / 'prepared_selection.json').is_file())
            # hr manifest does not bind it: nothing is invented
            hr_seed = root / 'compatibility/cases/hr_metro/seed/input'
            self.assertFalse((hr_seed / 'prepared_selection.json').exists())
            # the unselected dataset entry is preserved
            self.assertIn({'dataset_id': 'hr', 'recipe_path': 'other.json'}, index['entries'])
            recipe = json.loads((root / 'compatibility/recipes/retail_online.json').read_text())
            self.assertIn('input/prepared_selection.json', recipe['input_hashes'])


if __name__ == '__main__':
    unittest.main()
