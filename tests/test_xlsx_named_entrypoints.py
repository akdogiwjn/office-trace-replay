"""Check renamed entrypoints with tiny workbooks; no real replay or recalc."""
from contextlib import redirect_stdout
import io
import json
import os
from pathlib import Path
import runpy
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from replay.common import digest, read
from scripts.compile_compatible import _align_xlsx_entrypoints


ROOT = Path(__file__).resolve().parents[1]
TOOLS = ROOT / 'compatibility/tools'
INSPECTORS = {
    'inspect_xlsx_structure.py': 'source',
    'inspect_xlsx_features.py': 'copy',
    'inspect_xlsx_formulas.py': 'formula',
    'inspect_xlsx_cached_values.py': 'cached',
}


class NamedXlsxEntrypointTests(unittest.TestCase):
    def fixture(self, root):
        (root / 'input').mkdir()
        output = root / 'output'
        output.mkdir()
        book = Workbook()
        book.active.title = 'Data'
        book.active['A1'] = 7
        book.save(root / 'input/source.xlsx')
        sheet = book.create_sheet('Overview')
        sheet['B1'] = 7
        sheet['B2'] = 14
        book.save(output / 'result.xlsx')
        book.close()
        values = {
            root / 'input/dataset_manifest.json': {
                'kind': 'xlsx', 'workbook': 'source.xlsx', 'output_workbook': 'result.xlsx'},
            output / 'workbook_features.json': {
                'summary_sheet': 'Overview',
                'metrics': {'total': {'sheet': 'Overview', 'cell': 'B1'}},
                'scenario': {'projections': {'projected': {'sheet': 'Overview', 'cell': 'B2'}}}},
            output / 'formula_recalc.json': {'status': 'success'},
            output / 'business_verification.json': {'status': 'success'},
        }
        for path, value in values.items():
            path.write_text(json.dumps(value))

    def execute(self, root, name, *args):
        previous = Path.cwd()
        stream = io.StringIO()
        try:
            os.chdir(root)
            with patch.object(sys, 'argv', [name, *args]), redirect_stdout(stream), \
                 patch('openpyxl.load_workbook', wraps=load_workbook) as loads, \
                 patch('subprocess.run', side_effect=AssertionError('unexpected child process')):
                runpy.run_path(str(TOOLS / name), run_name='__main__')
                return stream.getvalue(), loads.call_count
        finally:
            os.chdir(previous)

    def test_four_inspection_names_keep_modes_and_load_counts(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            before = digest(root / 'output/result.xlsx')
            for alias, mode in INSPECTORS.items():
                with self.subTest(alias=alias):
                    expected = self.execute(root, 'inspect_xlsx.py', mode)
                    actual = self.execute(root, alias)
                    self.assertEqual(actual, expected)
                    self.assertEqual(actual[1], 1)
            self.assertEqual(digest(root / 'output/result.xlsx'), before)

    def test_csv_alias_keeps_output_and_one_load(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            self.execute(root, 'export_xlsx.py')
            expected = [(root / 'output' / name).read_bytes()
                        for name in ('metrics.csv', 'scenario.csv')]
            _, count = self.execute(root, 'export_xlsx_csv.py')
            self.assertEqual(count, 1)
            self.assertEqual(expected, [(root / 'output' / name).read_bytes()
                                       for name in ('metrics.csv', 'scenario.csv')])

    def test_summary_alias_keeps_output_and_two_loads(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            self.fixture(root)
            self.execute(root, 'summary.py')
            path = root / 'output/xlsx_enhancement_summary.json'
            expected = path.read_bytes()
            _, count = self.execute(root, 'write_xlsx_summary.py')
            self.assertEqual(count, 2)
            self.assertEqual(path.read_bytes(), expected)

    def test_named_help_does_not_open_workbooks(self):
        for alias in [*INSPECTORS, 'export_xlsx_csv.py', 'write_xlsx_summary.py']:
            with self.subTest(alias=alias), patch.object(sys, 'argv', [alias, '--help']), \
                 redirect_stdout(io.StringIO()), \
                 patch('openpyxl.load_workbook', side_effect=AssertionError('workbook opened')):
                with self.assertRaises(SystemExit) as result:
                    runpy.run_path(str(TOOLS / alias), run_name='__main__')
                self.assertEqual(result.exception.code, 0)

    def test_three_recipes_keep_helper_bytes_parent_and_execution_shape(self):
        record = read(ROOT / 'reports/script-entrypoint-alignment-v1.json')
        for dataset, binding in record['datasets'].items():
            with self.subTest(dataset=dataset):
                before = read(ROOT / binding['previous_recipe_path'])
                after = read(ROOT / binding['recipe_path'])
                self.assertEqual(digest(ROOT / binding['previous_recipe_path']),
                                 binding['previous_recipe_sha256'])
                self.assertEqual(digest(ROOT / binding['recipe_path']), binding['recipe_sha256'])
                self.assertEqual(before.get('parent'), after.get('parent'))
                self.assertEqual(list(before['helper_hashes'].values()), list(after['helper_hashes'].values()))
                old_writes = [item['tool_call']['arguments'] for step in before['key_operations']
                              for item in step['source_tool_calls'] if item.get('role') == 'write_helper']
                new_writes = [item['tool_call']['arguments'] for step in after['key_operations']
                              for item in step['source_tool_calls'] if item.get('role') == 'write_helper']
                self.assertEqual(len(old_writes), 1)
                self.assertEqual(old_writes[0]['content'], new_writes[0]['content'])
                self.assertEqual(Path(old_writes[0]['path']).parent, Path(new_writes[0]['path']).parent)
                self.assertEqual(Path(new_writes[0]['path']).name, 'enhance_workbook.py')
                self.assertEqual([len(s['source_tool_calls']) for s in before['key_operations']],
                                 [len(s['source_tool_calls']) for s in after['key_operations']])
                old_inputs = before['input_hashes']
                self.assertEqual(old_inputs, {k: after['input_hashes'][k] for k in old_inputs})
                seed = ROOT / after['seed_path']
                self.assertEqual(digest(seed / 'input/verify_office.py'),
                                 digest(seed / 'input/verify_xlsx_enhanced.py'))
                duplicate = json.loads(json.dumps(after))
                _align_xlsx_entrypoints(ROOT, duplicate)
                self.assertEqual(duplicate, after)

    def test_unselected_historical_datasets_are_untouched(self):
        recipe = {'dataset_id': 'retail'}
        _align_xlsx_entrypoints(Path('/unused'), recipe)
        self.assertEqual(recipe, {'dataset_id': 'retail'})


if __name__ == '__main__':
    unittest.main()
