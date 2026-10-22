"""Small fixtures for counting effective conditional-formatting regions."""
from datetime import datetime
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.formatting import ConditionalFormatting
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.cell_range import MultiCellRange

from compatibility.helpers.hr_metro_build import build as build_hr_metro
from compatibility.helpers.retail_online_build import build as build_retail_online
from compatibility.verification.spreadsheet import conditional_formatting_region_count


def positive_rule():
    return CellIsRule(operator='greaterThan', formula=['0'],
                      fill=PatternFill('solid', fgColor='E2F0D9'))


class ConditionalFormattingRegionTests(unittest.TestCase):
    def test_split_and_libreoffice_merged_representations_count_identically(self):
        split = Workbook()
        split_sheet = split.active
        split_sheet.conditional_formatting.add('B2:B4', positive_rule())
        split_sheet.conditional_formatting.add('B10', positive_rule())

        merged = Workbook()
        merged_sheet = merged.active
        group = ConditionalFormatting(sqref=MultiCellRange('B2:B4 B10'))
        merged_sheet.conditional_formatting._cf_rules[group] = [positive_rule()]

        self.assertEqual(conditional_formatting_region_count([split_sheet]), 2)
        self.assertEqual(conditional_formatting_region_count([merged_sheet]), 2)
        split.close()
        merged.close()

    def test_one_region_empty_rule_group_and_duplicate_regions_do_not_meet_two(self):
        book = Workbook()
        sheet = book.active
        sheet.conditional_formatting.add('B2:B4', positive_rule())
        self.assertEqual(conditional_formatting_region_count([sheet]), 1)

        no_rule = ConditionalFormatting(sqref=MultiCellRange('C2:C4'))
        sheet.conditional_formatting._cf_rules[no_rule] = []
        duplicate = ConditionalFormatting(sqref=MultiCellRange('B2:B4 B2:B4'))
        sheet.conditional_formatting._cf_rules[duplicate] = [positive_rule()]
        self.assertEqual(conditional_formatting_region_count([sheet]), 1)
        book.close()

    def test_retail_and_hr_metro_helpers_cover_two_regions(self):
        from scripts import prepare_xlsx_selection as preparation
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)

            def staged_workspace(dataset, workbook_name, populate):
                workspace = root / (dataset + '-workspace')
                (workspace / 'input').mkdir(parents=True)
                (workspace / 'output').mkdir()
                source = workspace / 'input' / workbook_name
                book = Workbook()
                sheet = book.active
                populate(sheet)
                book.save(source)
                book.close()
                preparation.cli(['--dataset', dataset, '--input', str(source),
                                 '--output', str(workspace / 'input/prepared_selection.json'),
                                 '--fixture'])
                output = workspace / 'output' / workbook_name
                output.write_bytes(source.read_bytes())
                return output

            def retail_sheet(sheet):
                sheet.title = 'Online Retail'
                sheet.append(['InvoiceNo', 'Quantity', 'UnitPrice', 'InvoiceDate', 'CustomerID'])
                sheet.append(['1001', 2, 3.0, datetime(2010, 1, 15), None])

            def hr_sheet(sheet):
                sheet.title = 'MSA_M2025_dl'
                sheet.append(['AREA', 'AREA_TITLE', 'AREA_TYPE', 'NAICS', 'I_GROUP',
                              'OWN_CODE', 'O_GROUP', 'OCC_CODE', 'OCC_TITLE',
                              'TOT_EMP', 'A_MEAN', 'H_MEDIAN', 'H_MEAN', 'A_MEDIAN'])
                scope = ['35620', 'New York-Newark-Jersey City, NY-NJ', '4',
                         '000000', 'cross-industry', '1235']
                sheet.append(scope + ['total', '00-0000', 'All occupations',
                                      1000, 80000, 30, 40, 60000])
                for index in range(5):
                    sheet.append(scope + ['major', f'{index:02d}-0000', f'Group {index}',
                                          500 - index, 70000, 25, 35, 50000])

            retail_output = build_retail_online(staged_workspace(
                'retail_online', 'Online Retail.xlsx', retail_sheet))
            hr_output = build_hr_metro(staged_workspace(
                'hr_metro', 'MSA_M2025_dl.xlsx', hr_sheet))

            for output, summary_name in ((retail_output, 'Sales Overview'),
                                         (hr_output, 'Metro Summary')):
                generated = load_workbook(output, data_only=False)
                try:
                    self.assertGreaterEqual(
                        conditional_formatting_region_count([generated[summary_name]]), 2)
                finally:
                    generated.close()


if __name__ == '__main__':
    unittest.main()
