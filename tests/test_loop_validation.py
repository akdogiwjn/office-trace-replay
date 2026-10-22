"""Regression tests for the measured loop verifier, distinct from acceptance."""
import json
from pathlib import Path
import shutil
import sys
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile
import xml.etree.ElementTree as ET
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

from compatibility.verification.verify import Checks
from compatibility.verification.verify import verify as verify_entry
from compatibility.verification.verify import verify_pdf as verify_pdf_independent
from compatibility.verification.spreadsheet import verify_semantic_xlsx_loop
from compatibility.verification.pdf_loop import check_representative_changes, selected_records, verify_pdf


class LoopSpreadsheetContractTests(unittest.TestCase):
    @staticmethod
    def _set_formula_cache(path, address, value):
        """Set a fixture-only cached formula value without running a spreadsheet app."""
        ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        ET.register_namespace('', ns['x'])
        temp = path.with_suffix('.patched.xlsx')
        found = False
        with ZipFile(path, 'r') as source, ZipFile(temp, 'w', ZIP_DEFLATED) as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                    root = ET.fromstring(content)
                    for cell in root.findall('.//x:c', ns):
                        if cell.attrib.get('r') == address and cell.find('x:f', ns) is not None:
                            cached = cell.find('x:v', ns)
                            if cached is None:
                                cached = ET.SubElement(cell, '{%s}v' % ns['x'])
                            cached.text = str(value)
                            content = ET.tostring(root, encoding='utf-8', xml_declaration=True)
                            found = True
                            break
                target.writestr(item, content)
        if not found:
            temp.unlink(missing_ok=True)
            raise AssertionError(f'formula cache cell {address} not found in {path}')
        temp.replace(path)

    def _scenario_cache_fixture(self, root, intermediate_formula, cached_value, expected_value):
        workspace, manifest, expected = self._fixture(root)
        target = workspace / 'output/result.xlsx'
        wb = load_workbook(target)
        overview = wb['Overview']
        overview['B3'] = '=B4'
        overview['B4'] = intermediate_formula
        wb.save(target)
        wb.close()
        self._set_formula_cache(target, 'B3', cached_value)
        features = json.loads((workspace / 'output/workbook_features.json').read_text())
        features['scenario']['projections'] = {
            'projection': {'sheet': 'Overview', 'cell': 'B3'}}
        (workspace / 'output/workbook_features.json').write_text(json.dumps(features))
        manifest['requirements']['minimum_formula_count'] = 2
        manifest['requirements']['scenario']['projections'] = [{'id': 'projection'}]
        expected['scenario_values'] = {'Base': {'projection': expected_value}}
        (workspace / 'output/scenario.csv').write_text(
            f'metric,value\nprojection,{cached_value}\n')
        recalc = json.loads((workspace / 'output/formula_recalc.json').read_text())
        recalc.update({'total_formulas': 2, 'total_errors': 0, 'status': 'success'})
        (workspace / 'output/formula_recalc.json').write_text(json.dumps(recalc))
        return workspace, manifest, expected

    def _fixture(self, root):
        workspace = root / 'workspace'
        source = workspace / 'input' / 'source.xlsx'
        output = workspace / 'output'
        source.parent.mkdir(parents=True)
        output.mkdir(parents=True)

        wb = Workbook()
        data = wb.active
        data.title = 'SourceData'
        data['A1'] = 'value'
        data['A2'] = 12
        wb.save(source)
        wb.close()

        wb = load_workbook(source)
        summary = wb.create_sheet('Overview')
        summary['A1'] = 'Scenario'
        summary['B1'] = 'Base'
        summary['B1'].font = Font(color='FF0000FF')
        summary['A2'] = 'Multiplier'
        summary['B2'] = 1.0
        summary['B2'].font = Font(color='FF0000FF')
        summary.freeze_panes = 'B3'
        validation = DataValidation(type='list', formula1='"Base"')
        validation.add(summary['B1'])
        summary.add_data_validation(validation)
        wb.save(output / 'result.xlsx')
        wb.close()

        (output / 'workbook_features.json').write_text(json.dumps({
            'summary_sheet': 'Overview', 'metrics': {},
            'scenario': {'selector': {'sheet': 'Overview', 'cell': 'B1'},
                         'assumptions': {'Base': {'multiplier': {'sheet': 'Overview', 'cell': 'B2'}}},
                         'projections': {}},
            'charts': {},
        }))
        (output / 'formula_recalc.json').write_text(json.dumps({'status': 'success', 'total_errors': 0}))
        (output / 'xlsx_enhancement_summary.json').write_text('{}')
        (output / 'metrics.csv').write_text('metric,value\n')
        (output / 'scenario.csv').write_text('metric,value\n')
        manifest = {
            'dataset_id': 'fixture', 'kind': 'xlsx', 'output_workbook': 'result.xlsx',
            'workbook': 'source.xlsx', 'required_outputs': [
                'result.xlsx', 'workbook_features.json', 'formula_recalc.json',
                'business_verification.json', 'xlsx_enhancement_summary.json'],
            'requirements': {
                'base_sheets': ['SourceData'], 'base_dimensions': {'SourceData': [2, 1]},
                'base_dimension_details': {'SourceData': {
                    'declared_dimension': [3, 1],
                    'serialized_cell_range': [2, 1],
                    'effective_data_range': [2, 1],
                    'effective_range_method': 'fixture populated-cell extent',
                    'boundary_cells': ['A2'],
                }},
                'metrics': [], 'scenario': {
                    'default': 'Base', 'rows': [{'name': 'Base', 'assumptions': {'multiplier': 1.0}}],
                    'projections': []}, 'charts': [], 'minimum_comments': 0,
                'minimum_conditional_formats': 0, 'minimum_formula_count': 0,
            },
            'publish_files': [],
        }
        expected = {'metric_values': {}, 'scenario_values': {}}
        return workspace, manifest, expected

    def test_loop_contract_uses_basic_structure_and_skips_deep_source_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            checks = Checks(manifest)
            with patch('compatibility.verification.spreadsheet.original_content',
                       side_effect=AssertionError('deep source audit entered loop')), \
                 patch('openpyxl.load_workbook', wraps=load_workbook) as loader:
                verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                          deferred_outputs=('xlsx_enhancement_summary.json',))
            self.assertEqual(loader.call_count, 2)
            report = checks.finish()
            self.assertEqual(report['status'], 'success', report['failures'])
            self.assertTrue(report['checks']['base_sheet_scale:SourceData']['ok'])

    def test_verify_loop_entry_skips_asset_rehash_and_recursive_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            case = workspace.parent / 'dataset'
            case.mkdir()
            manifest.update({
                'schema_version': 'office-dataset-v2', 'dataset_id': 'fixture',
                'workload_contract': 'spreadsheet-semantic-v3',
                'task': 'fixture task',
                'input_files': {'input/source.xlsx': '0' * 64},
                'expected_file': 'expected.json', 'expected_sha256': 'f' * 64,
                'trace_rules': [],
            })
            manifest_path = case / 'manifest.json'
            manifest_path.write_text(json.dumps(manifest))
            (case / 'expected.json').write_text(json.dumps(expected))
            # Same manifest must remain strict for independent callers.
            from compatibility.verification.contracts import load_manifest
            with self.assertRaisesRegex(ValueError, 'input hash mismatch'):
                load_manifest(manifest_path)
            with patch('compatibility.verification.verify.load_manifest', wraps=load_manifest) as loader, \
                 patch('compatibility.verification.spreadsheet.dependencies',
                       side_effect=AssertionError('recursive dependency scan entered loop')):
                report = verify_entry(manifest_path, workspace, scope='loop',
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            loader.assert_called_once_with(manifest_path, check_assets=False)
            self.assertEqual(report['status'], 'success', report['failures'])

    def test_loop_contract_fails_if_a_base_sheet_is_removed(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = workspace / 'output/result.xlsx'
            wb = load_workbook(target)
            del wb['SourceData']
            wb.save(target)
            wb.close()
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertIn('base_sheets_present', report['failures'])
            self.assertIn('base_sheet_scale:SourceData', report['failures'])

    def test_declared_empty_tail_does_not_fail_against_serialized_output_extent(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            # The source declaration includes an empty tail row, while frozen
            # effective data extent and output serialized cells both end at 2.
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'success', report['failures'])
            detail = report['checks']['base_sheet_scale:SourceData']['detail']
            self.assertEqual(detail['source_declared_dimension'], [3, 1])
            self.assertEqual(detail['source_effective_data_range'], (2, 1))
            self.assertEqual(detail['output_serialized_cell_range'], (2, 1))

    def test_effective_data_boundary_shrink_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = workspace / 'output/result.xlsx'
            wb = load_workbook(target)
            wb['SourceData']['A2'] = None
            wb.save(target)
            wb.close()
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertIn('base_sheet_scale:SourceData', report['failures'])

    def test_indirect_scenario_formula_is_checked_by_formula_and_cached_result(self):
        """Structural-only branch test; real cache comparison is covered below."""
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = workspace / 'output/result.xlsx'
            wb = load_workbook(target)
            sheet = wb['Overview']
            sheet['B3'] = '=B4*B2'
            sheet['B4'] = '=IF(B1="Base",1,1.1)'
            wb.save(target)
            wb.close()
            features = json.loads((workspace / 'output/workbook_features.json').read_text())
            features['scenario']['projections'] = {'projection': {'sheet': 'Overview', 'cell': 'B3'}}
            (workspace / 'output/workbook_features.json').write_text(json.dumps(features))
            manifest['requirements']['minimum_formula_count'] = 2
            manifest['requirements']['scenario']['projections'] = [{'id': 'projection'}]
            expected['scenario_values'] = {'Base': {'projection': 1}}
            (workspace / 'output/scenario.csv').write_text('metric,value\nprojection,1\n')
            with patch('compatibility.verification.verify.same_number', return_value=True):
                checks = Checks(manifest)
                verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                          deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertTrue(report['checks']['semantic:scenario_dependency:projection']['ok'])
            self.assertIn('scope', report['checks']['semantic:scenario_dependency:projection']['detail'])

    def test_indirect_scenario_uses_real_cached_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._scenario_cache_fixture(
                Path(tmp), '=IF(B1="Base",1,1.1)', 1, 1)
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'success', report['failures'])
            self.assertTrue(report['checks']['semantic:projection:projection']['ok'])
            self.assertTrue(report['checks']['semantic:scenario_dependency:projection']['ok'])

    def test_indirect_scenario_rejects_wrong_real_cached_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._scenario_cache_fixture(
                Path(tmp), '=IF(B1="Base",1,1.1)', 9, 1)
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertEqual(report['failures'], ['semantic:projection:projection'])

    def test_scenario_dependency_rejects_unrelated_intermediate_formula(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._scenario_cache_fixture(
                Path(tmp), '=1+1', 2, 2)
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertIn('semantic:scenario_dependency:projection', report['failures'])
            self.assertTrue(report['checks']['semantic:projection:projection']['ok'])

    def _indirect_chart_fixture(self, workspace, manifest):
        target = workspace / 'output/result.xlsx'
        wb = load_workbook(target)
        raw = wb['SourceData']
        raw['A3'] = 12
        data = wb.create_sheet('Chart Data')
        data['A1'], data['A2'] = 'First', 'Second'
        data['B1'], data['B2'] = '=SourceData!A3', '=B1+1'
        chart = BarChart()
        chart.add_data(Reference(data, min_col=2, min_row=1, max_row=2), titles_from_data=False)
        chart.set_categories(Reference(data, min_col=1, min_row=1, max_row=2))
        data.add_chart(chart, 'D1')
        wb.save(target)
        wb.close()
        features = json.loads((workspace / 'output/workbook_features.json').read_text())
        features['charts'] = {'major_employment': {'sheet': 'Chart Data', 'index': 0}}
        (workspace / 'output/workbook_features.json').write_text(json.dumps(features))
        manifest['requirements']['charts'] = [{'id': 'major_employment'}]
        manifest['requirements']['minimum_formula_count'] = 2
        return target

    def test_chart_dependency_accepts_intermediate_calculation_range(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            self._indirect_chart_fixture(workspace, manifest)
            expected['chart_values'] = {'major_employment': [12, 13]}
            with patch('compatibility.verification.verify.same_number', return_value=True):
                checks = Checks(manifest)
                verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                          deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertTrue(report['checks']['semantic:chart_dependency:major_employment']['ok'],
                            report['checks']['semantic:chart_dependency:major_employment'])

    def test_chart_without_data_reference_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = self._indirect_chart_fixture(workspace, manifest)
            wb = load_workbook(target)
            wb['Chart Data']._charts[0].series[0].val.numRef = None
            wb.save(target)
            wb.close()
            expected['chart_values'] = {'major_employment': [12, 13]}
            with patch('compatibility.verification.verify.same_number', return_value=True):
                checks = Checks(manifest)
                verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                          deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertFalse(report['checks']['semantic:chart:major_employment']['ok'])

    def test_loop_contract_fails_if_a_required_delivery_is_missing(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            (workspace / 'output/metrics.csv').unlink()
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertFalse(report['checks']['csv_matches_workbook:metrics.csv']['ok'])

    def test_loop_contract_rejects_a_required_feature_that_is_not_a_formula(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = workspace / 'output/result.xlsx'
            wb = load_workbook(target)
            wb['Overview']['B3'] = 12
            wb.save(target)
            wb.close()
            (workspace / 'output/workbook_features.json').write_text(json.dumps({
                'summary_sheet': 'Overview', 'metrics': {'answer': {'sheet': 'Overview', 'cell': 'B3'}},
                'scenario': {'selector': {'sheet': 'Overview', 'cell': 'B1'},
                             'assumptions': {'Base': {'multiplier': {'sheet': 'Overview', 'cell': 'B2'}}},
                             'projections': {}}, 'charts': {},
            }))
            manifest['requirements']['metrics'] = [{'id': 'answer'}]
            expected['metric_values'] = {'answer': 12}
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertFalse(report['checks']['feature:formula:Overview!B3']['ok'])

    def test_loop_contract_rejects_a_wrong_cached_metric_value(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = workspace / 'output/result.xlsx'
            wb = load_workbook(target)
            cell = wb['Overview']['B3']
            cell.value = '=SourceData!A2'
            cell.font = Font(color='FF008000')
            wb.save(target)
            wb.close()
            features = json.loads((workspace / 'output/workbook_features.json').read_text())
            features['metrics'] = {'answer': {'sheet': 'Overview', 'cell': 'B3'}}
            (workspace / 'output/workbook_features.json').write_text(json.dumps(features))
            (workspace / 'output/metrics.csv').write_text('metric,value\nanswer,12\n')
            manifest['requirements']['metrics'] = [{'id': 'answer'}]
            manifest['requirements']['minimum_formula_count'] = 1
            expected['metric_values'] = {'answer': 12}
            checks = Checks(manifest)
            verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                      deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertEqual(report['status'], 'failed')
            self.assertFalse(report['checks']['semantic:metric:answer']['ok'])

    def test_loop_scenario_dependency_accepts_middle_cell_in_formula_range(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._fixture(Path(tmp))
            target = workspace / 'output/result.xlsx'
            wb = load_workbook(target)
            sheet = wb['Overview']
            sheet['C7'] = 10
            sheet['C12'] = 'Base'
            sheet['C12'].font = Font(color='FF0000FF')
            validation = DataValidation(type='list', formula1='"Base,Higher,Lower"')
            validation.add(sheet['C12'])
            sheet.add_data_validation(validation)
            sheet['A14'], sheet['A15'], sheet['A16'] = 'Base', 'Higher', 'Lower'
            sheet['B14'], sheet['B15'], sheet['B16'] = 1.0, 1.1, 0.9
            for row in (14, 15, 16):
                sheet[f'B{row}'].font = Font(color='FF0000FF')
            sheet['C18'] = '=C7*INDEX(B14:B16,MATCH(C12,A14:A16,0))'
            wb.save(target)
            wb.close()

            features = {
                'summary_sheet': 'Overview', 'metrics': {}, 'charts': {},
                'scenario': {
                    'selector': {'sheet': 'Overview', 'cell': 'C12'},
                    'assumptions': {
                        name: {'multiplier': {'sheet': 'Overview', 'cell': f'B{row}'}}
                        for name, row in (('Base', 14), ('Higher', 15), ('Lower', 16))},
                    'projections': {'projection': {'sheet': 'Overview', 'cell': 'C18'}},
                },
            }
            (workspace / 'output/workbook_features.json').write_text(json.dumps(features))
            (workspace / 'output/scenario.csv').write_text('metric,value\nprojection,10\n')
            manifest['requirements']['scenario'] = {
                'default': 'Base',
                'rows': [{'name': name, 'assumptions': {'multiplier': value}}
                         for name, value in (('Base', 1.0), ('Higher', 1.1), ('Lower', 0.9))],
                'projections': [{'id': 'projection'}],
            }
            manifest['requirements']['minimum_formula_count'] = 1
            expected['scenario_values'] = {'Base': {'projection': 10}}
            checks = Checks(manifest)
            with patch('compatibility.verification.verify.same_number', return_value=True):
                verify_semantic_xlsx_loop(manifest, expected, workspace, checks,
                                          deferred_outputs=('xlsx_enhancement_summary.json',))
            report = checks.finish()
            self.assertTrue(report['checks']['semantic:scenario_dependency:projection']['ok'],
                            report['checks']['semantic:scenario_dependency:projection'])

    def test_pdf_loop_selection_preserves_n1_and_n10_execution_scopes(self):
        records = [{'id': f'record-{i}'} for i in range(1, 11)]
        for count in (1, 10):
            ids = [row['id'] for row in records[:count]]
            manifest = {'pdf_execution': {'record_ids': ids, 'record_count': count}}
            chosen = selected_records(manifest, records)
            self.assertEqual([row['id'] for row in chosen], ids)
            self.assertEqual(len(chosen), count)

    def test_pdf_loop_rejects_execution_scope_that_is_not_prepared_prefix(self):
        records = [{'id': f'record-{i}'} for i in range(1, 11)]
        manifest = {'pdf_execution': {'record_ids': ['record-2'], 'record_count': 1}}
        with self.assertRaisesRegex(ValueError, 'declared nonempty prefix'):
            selected_records(manifest, records)

    def test_pdf_representative_page_is_checked_for_every_executed_record(self):
        for count in (1, 10):
            with self.subTest(count=count), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                template = root / 'template'
                output = root / 'rendered'
                template.mkdir()
                output.mkdir()
                Image.new('RGB', (20, 20), 'white').save(template / 'page_3.png')
                selected = []
                for index in range(count):
                    ident = f'record-{index + 1}'
                    selected.append({'id': ident})
                    folder = output / ident
                    folder.mkdir()
                    Image.new('RGB', (20, 20), 'black').save(folder / 'page_3.png')
                req = {'representative_visible_change': {
                    'page': 3, 'minimum_changed_pixels': 100, 'pixel_threshold': 12}}
                checks = Checks({'dataset_id': 'fixture', 'kind': 'pdf'})
                check_representative_changes(template, output, selected, req, checks)
                report = checks.finish()
                self.assertEqual(report['status'], 'success', report['failures'])
                self.assertEqual(sum(name.startswith('representative_visible_change:')
                                     for name in report['checks']), count)
                if count == 10:
                    Image.new('RGB', (20, 20), 'white').save(output / 'record-2/page_3.png')
                    checks = Checks({'dataset_id': 'fixture', 'kind': 'pdf'})
                    check_representative_changes(template, output, selected, req, checks)
                    report = checks.finish()
                    self.assertEqual(report['status'], 'failed')
                    self.assertFalse(report['checks']['representative_visible_change:record-2']['ok'])

    def _pdf_fixture(self, root):
        fixture = Path(__file__).parent / 'fixtures/pdf-loop-irs-w4-n1'
        workspace = root / 'workspace'
        shutil.copytree(fixture / 'input', workspace / 'input')
        shutil.copytree(fixture / 'output', workspace / 'output')
        manifest = json.loads((workspace / 'input/dataset_manifest.json').read_text())
        expected = json.loads((fixture / 'expected.json').read_text())
        wheel = (Path(__file__).parents[1] / 'compatibility/cases/irs_w4-n1/source/vendor/wheels'
                 / 'pypdf-6.10.0-py3-none-any.whl')
        sys.path.insert(0, str(wheel))
        sys.path.insert(0, str(Path(__file__).parents[1]
                               / 'compatibility/cases/irs_w4-n1/source/vendor/skills/pdf/scripts'))
        return workspace, manifest, expected

    def test_pdf_loop_accepts_actual_n1_outputs_and_directory_deliverables(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._pdf_fixture(Path(tmp))
            checks = Checks(manifest)
            verify_pdf(manifest, expected, workspace, checks)
            report = checks.finish()
            self.assertEqual(report['status'], 'success', report['failures'])
            self.assertTrue(report['checks']['template_form']['ok'])
            for name in ('field_values', 'filled', 'rendered'):
                self.assertTrue(report['checks']['deliverable:' + name]['ok'])

    def test_pdf_loop_rejects_wrong_mapping_and_protected_field_entry(self):
        for defect in ('wrong_value', 'protected_field'):
            with self.subTest(defect=defect), tempfile.TemporaryDirectory() as tmp:
                workspace, manifest, expected = self._pdf_fixture(Path(tmp))
                values_path = workspace / 'output/field_values/employee_01.json'
                values = json.loads(values_path.read_text())
                if defect == 'wrong_value':
                    values[0]['value'] = 'deliberately incorrect'
                else:
                    values.append({'field_id': manifest['requirements']['protected_blank'][0], 'value': '123'})
                values_path.write_text(json.dumps(values))
                checks = Checks(manifest)
                verify_pdf(manifest, expected, workspace, checks)
                report = checks.finish()
                self.assertEqual(report['status'], 'failed')
                self.assertFalse(report['checks']['mapping:employee_01']['ok'])
                if defect == 'protected_field':
                    self.assertFalse(report['checks']['protected_fields_blank:employee_01']['ok'])

    def test_pdf_independent_acceptance_checks_all_mappings_but_only_executed_pdfs(self):
        with tempfile.TemporaryDirectory() as tmp:
            workspace, manifest, expected = self._pdf_fixture(Path(tmp))
            checks = Checks(manifest)
            verify_pdf_independent(manifest, expected, workspace, checks)
            report = checks.finish()
            self.assertEqual(report['status'], 'success', report['failures'])
            self.assertTrue(report['checks']['mapping:employee_10']['ok'])
            self.assertNotIn('pages:employee_10', report['checks'])


if __name__ == '__main__':
    unittest.main()
