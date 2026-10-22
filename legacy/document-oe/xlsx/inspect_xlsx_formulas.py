#!/usr/bin/env python3
"""Inspect formulas and merged ranges in Executive_Summary."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook = load_workbook(args.workbook, data_only=False)
    try:
        worksheet = workbook["Executive_Summary"]
        for row in range(12, 18):
            for column in range(2, 7):
                cell = worksheet.cell(row=row, column=column)
                if cell.value is not None:
                    print(f"{cell.coordinate}: {repr(cell.value)}")
                else:
                    print(f"{cell.coordinate}: None")
        print()
        print("Merged ranges:", [str(item) for item in worksheet.merged_cells.ranges])
        return 0
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())
