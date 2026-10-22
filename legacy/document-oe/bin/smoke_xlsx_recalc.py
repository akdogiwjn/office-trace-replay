#!/usr/bin/env python3
"""Run a fast, real LibreOffice recalculation check inside the document image."""

from __future__ import annotations

import argparse
import json
import logging
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook


DEFAULT_ROOT = Path("/opt/document-bench")
REQUIRED_CELLS = ("B5", "B6", "B7", "B8", "B9", "B13", "B14", "B15", "B16")
EXPECTED_VALUE = 2
LOGGER = logging.getLogger("document-bench-recalc-smoke")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--timeout", type=int, default=30, help="LibreOffice recalculation timeout in seconds")
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT, help=argparse.SUPPRESS)
    return parser.parse_args()


def create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Executive_Summary"
    for cell in REQUIRED_CELLS:
        sheet[cell] = "=1+1"
    workbook.save(path)
    workbook.close()


def read_cached_values(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["Executive_Summary"]
        return {cell: sheet[cell].value for cell in REQUIRED_CELLS}
    finally:
        workbook.close()


def run_smoke(root: Path, timeout: int) -> dict[str, object]:
    if timeout <= 0:
        raise ValueError("--timeout must be positive")

    recalc_script = root / "skills/xlsx/scripts/recalc.py"
    if not recalc_script.is_file():
        raise FileNotFoundError(f"recalculation script is missing: {recalc_script}")

    with tempfile.TemporaryDirectory(prefix="document-bench-recalc-smoke-") as directory:
        workbook_path = Path(directory) / "minimal_recalc.xlsx"
        create_workbook(workbook_path)
        completed = subprocess.run(
            [sys.executable, str(recalc_script), str(workbook_path), str(timeout)],
            capture_output=True,
            text=True,
            timeout=timeout + 75,
        )
        if completed.returncode != 0:
            detail = completed.stderr.strip() or completed.stdout.strip() or "recalc.py returned no diagnostics"
            raise RuntimeError(f"recalc.py failed with exit code {completed.returncode}: {detail}")

        try:
            recalc_result = json.loads(completed.stdout)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"recalc.py returned invalid JSON: {completed.stdout!r}") from exc

        if recalc_result.get("status") != "success":
            raise RuntimeError(f"recalc.py did not report success: {recalc_result}")
        if recalc_result.get("total_formulas") != len(REQUIRED_CELLS):
            raise RuntimeError(f"unexpected formula count: {recalc_result}")
        if set(recalc_result.get("cached_cells_verified", [])) != set(REQUIRED_CELLS):
            raise RuntimeError(f"required cached cells were not verified: {recalc_result}")
        before_mtime = recalc_result.get("before_mtime_ns", 0)
        after_mtime = recalc_result.get("after_mtime_ns", 0)
        if not isinstance(before_mtime, int) or not isinstance(after_mtime, int) or after_mtime <= before_mtime:
            raise RuntimeError(f"recalculated workbook was not saved: {recalc_result}")

        cached_values = read_cached_values(workbook_path)
        invalid = {cell: value for cell, value in cached_values.items() if value != EXPECTED_VALUE}
        if invalid:
            raise RuntimeError(f"unexpected cached formula values: {invalid}")

    return {
        "status": "success",
        "formula_count": len(REQUIRED_CELLS),
        "cached_values": cached_values,
        "recalc_returncode": recalc_result.get("returncode"),
        "workbook_saved": True,
    }


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    args = parse_args()
    try:
        result = run_smoke(args.root, args.timeout)
    except Exception as exc:
        LOGGER.error("minimal XLSX recalculation smoke failed: %s", exc)
        return 1
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
