#!/usr/bin/env python3
"""Inspect XLSX workbook structure for benchmark P01."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

try:
    # Replace/add this module later with the QEMU-specific marker implementation.
    from document_roi import roi_begin, roi_end
except ImportError:
    def roi_begin(_roi_id: str) -> None:
        pass

    def roi_end(_roi_id: str) -> None:
        pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook = None
    roi_begin("XLSX_P01_STRUCTURE")
    try:
        workbook = load_workbook(args.workbook, data_only=False)
        print("Sheets:", workbook.sheetnames)
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            print(f"\n=== {sheet_name} ===")
            print(f"  Rows: {sheet.max_row}, Cols: {sheet.max_column}")
            print(f"  Freeze: {sheet.freeze_panes}")
            print(f"  Charts: {len(sheet._charts)}")
            print(f"  Cond fmt: {len(sheet.conditional_formatting)}")
            for row_index in range(1, min(sheet.max_row + 1, 25)):
                row_values = []
                for column_index in range(1, min(sheet.max_column + 1, 12)):
                    cell = sheet.cell(row=row_index, column=column_index)
                    if cell.value is not None:
                        row_values.append(f"{cell.coordinate}={cell.value}")
                if row_values:
                    print(f"  Row {row_index}: {row_values}")
        return 0
    finally:
        if workbook is not None:
            workbook.close()
        roi_end("XLSX_P01_STRUCTURE")


if __name__ == "__main__":
    raise SystemExit(main())
