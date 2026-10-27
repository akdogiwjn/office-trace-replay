#!/usr/bin/env python3
"""Freeze explicit old-workflow adaptations; never rewrite parent traces."""
import copy
import argparse
import json
from pathlib import Path
import shlex
import shutil
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from replay.common import digest, read, relative, write

DATASETS = ['irs_w4', 'sba1919', 'hr', 'manufacturing', 'retail', 'industry_productivity',
            'hr_metro', 'retail_online']
HELPERS = {
    'irs_w4': ['build_field_values.py'],
    'sba1919': ['helpers/build_field_values.py'],
    'hr': ['helpers/build.py'],
    'manufacturing': ['build_enhanced.py'],
    'retail': ['helpers/build_enhance.py'],
}
BUILD = {'irs_w4': 'build_field_values.py', 'sba1919': 'helpers/build_field_values.py',
         'hr': 'helpers/build.py', 'manufacturing': 'build_enhanced.py', 'retail': 'helpers/build_enhance.py'}
# New XLSX build helpers live in the replay tree and write the final output
# workbook directly; no assemble_output / inject_caches / patch preservation step.
NEW_XLSX_HELPERS = {
    'hr': 'compatibility/helpers/hr_build.py',
    'manufacturing': 'compatibility/helpers/manufacturing_build_enhanced.py',
    'retail': 'compatibility/helpers/retail_build_enhance.py',
    'hr_metro': 'compatibility/helpers/hr_metro_build.py',
    'retail_online': 'compatibility/helpers/retail_online_build.py',
}
CLASSIFICATION = 'canonicalized_deterministic_replay_derived_from_successful_agent_executions'
XLSX_CONTRACT = 'spreadsheet-semantic-v3'
XLSX_ADAPTATIONS = [
    'Canonicalized deterministic replay derived from successful Agent executions: '
    'P01 copies the official input to the final output workbook; P02 writes and runs the '
    'dataset-specific enhancement helper through the generic atomic XLSX helper, enhancing the '
    'final output workbook directly; P03 recalculates that final output with the original Skill '
    'recalc.py so LibreOffice saves the computed caches in place, then checks formulas/cached '
    'values and exports metrics; P04 runs the versioned semantic verifier, writes the summary '
    'and checks deliverables.',
    'Enhancement helper writes the final output workbook directly; original recalc.py recalcs it in place.',
    'No assemble_output.py / inject_caches.py / patch.py preservation call in the formal recipe.',
    'Semantic verification uses the versioned spreadsheet-semantic-v3 contract; format/serialization differences are diagnostics.',
]
XLSX_OUTPUTS = {'hr': 'hr_enhanced.xlsx', 'manufacturing': 'manufacturing_enhanced.xlsx',
                'retail': 'retail_enhanced.xlsx',
                'hr_metro': 'hr_metro_enhanced.xlsx',
                'retail_online': 'retail_online_enhanced.xlsx'}
MANUAL_XLSX = {'hr_metro': 'hr', 'retail_online': 'retail'}
ALIGNED_XLSX = {'hr_metro', 'retail_online', 'industry_productivity'}
XLSX_ENTRYPOINT_REPLACEMENTS = {
    'compatibility/tools/inspect_xlsx.py source': 'compatibility/tools/inspect_xlsx_structure.py',
    'compatibility/tools/inspect_xlsx.py copy': 'compatibility/tools/inspect_xlsx_features.py',
    'compatibility/tools/inspect_xlsx.py formula': 'compatibility/tools/inspect_xlsx_formulas.py',
    'compatibility/tools/inspect_xlsx.py cached': 'compatibility/tools/inspect_xlsx_cached_values.py',
    'compatibility/tools/export_xlsx.py': 'compatibility/tools/export_xlsx_csv.py',
    'compatibility/tools/summary.py': 'compatibility/tools/write_xlsx_summary.py',
    'input/verify_office.py': 'input/verify_xlsx_enhanced.py',
}


def _align_xlsx_entrypoints(root, recipe):
    """Align executable basenames, preserving helper bytes and shared logic.

    Old recipes and canonical traces remain historical evidence. The original
    verifier entry is retained for callers using its previous name. Helpers
    stay in their original directory because Industry derives ROOT from it.
    """
    if recipe.get('dataset_id') not in ALIGNED_XLSX:
        return
    root = Path(root)
    helper_renames = {}
    for step in recipe['key_operations']:
        for item in step['source_tool_calls']:
            call = item.get('tool_call', {})
            if call.get('function_name') != 'write' or item.get('role') != 'write_helper':
                continue
            args = call['arguments']
            old_path = Path(args['path'])
            new_path = old_path.with_name('enhance_workbook.py')
            args['path'] = str(new_path)
            for key in list(recipe.get('helper_hashes', {})):
                if old_path.as_posix().endswith('/' + key):
                    new_key = str(Path(key).with_name('enhance_workbook.py'))
                    if new_key != key:
                        helper_renames[key] = new_key
                        recipe['helper_hashes'][new_key] = recipe['helper_hashes'].pop(key)
    for step in recipe['key_operations']:
        for item in step['source_tool_calls']:
            call = item.get('tool_call', {})
            if call.get('function_name') != 'exec':
                continue
            args = call['arguments']
            command = args.get('command', '')
            for old, new in {**XLSX_ENTRYPOINT_REPLACEMENTS, **helper_renames}.items():
                command = command.replace(old, new)
            args['command'] = command
    seed = root / recipe['seed_path']
    original = seed / 'input/verify_office.py'
    alias = seed / 'input/verify_xlsx_enhanced.py'
    shutil.copyfile(original, alias)
    recipe['input_hashes']['input/verify_xlsx_enhanced.py'] = digest(alias)
    recipe['script_naming'] = {
        'version': 'document-oe-script-names-v1',
        'scope': 'Executable entry basenames only; generic implementation and dataset helper bytes retained.',
        'historical_profiles_renamed': False,
    }


def preserved_suite_entries(previous, selected_datasets):
    selected = set(selected_datasets)
    return [entry for entry in previous.get('entries', [])
            if entry.get('dataset_id') not in selected]


def _refresh_public_suite(root):
    """Refresh derived hashes without changing the public dataset selection."""
    path = root / 'suites/default.json'
    if not path.is_file():
        return
    public = read(path)
    for key, filename in (
            ('legacy_lock_sha256', 'legacy/document-oe/lock.json'),
            ('compatible_contract_sha256', 'compatibility/contract.json'),
            ('compatible_suite_sha256', 'compatibility/suite.json')):
        public[key] = digest(root / filename)
    for entry in public['datasets']:
        entry['recipe_sha256'] = digest(root / relative(entry['recipe_path']))
    write(path, public)


def _copy_current_verifier_package(root, source, *, pdf=False):
    """Install the versioned measured-loop and independent verifier modules."""
    package = source / 'office_trace_bench'
    package.mkdir(parents=True, exist_ok=True)
    names = ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py')
    for name in names:
        shutil.copyfile(root / 'compatibility/verification' / name, package / name)
    if pdf:
        shutil.copyfile(root / 'compatibility/verification/pdf_loop.py', package / 'pdf_loop.py')
        # Keep the compatible-source verifier artifact aligned for tools that
        # invoke the independent subset contract directly.
        shutil.copyfile(root / 'compatibility/verification/pdf.py', package / 'compatible_pdf.py')


def _update_industry_productivity_recipe(root, index):
    """Refresh the existing canonical-derived recipe without touching its parent trace."""
    recipe_path = root / 'compatibility/recipes/industry_productivity.json'
    recipe = read(recipe_path)
    parent_path = root / recipe['parent']['recipe_path']
    parent = read(parent_path)
    if (recipe['parent'].get('run_id') != parent['run_id']
            or recipe['parent'].get('canonical_sha256') != parent['canonical_sha256']):
        raise ValueError('industry_productivity parent run/canonical binding mismatch')
    _copy_current_verifier_package(root, root / recipe['source_path'])
    for step in recipe['key_operations']:
        for item in step['source_tool_calls']:
            call = item.get('tool_call', {})
            if call.get('function_name') == 'exec':
                arguments = call.get('arguments', {})
                command_text = arguments.get('command', '')
                arguments['command'] = command_text.replace(
                    'compatibility/tools/industry_productivity_summary.py',
                    'compatibility/tools/summary.py')
    recipe['adaptations'] = [
        text.replace('metadata-only summary', 'shared workbook-reading summary')
        for text in recipe.get('adaptations', [])]
    if 'derivation' in recipe:
        recipe['derivation']['adaptations'] = [
            text.replace('metadata-only summary', 'shared workbook-reading summary')
            for text in recipe['derivation'].get('adaptations', [])]
    recipe['workflow_contract'] = XLSX_CONTRACT
    recipe['loop_verification_contract'] = 'document-oe-compatible-xlsx-loop-v1'
    _align_xlsx_entrypoints(root, recipe)
    write(recipe_path, recipe)
    index['entries'].append({
        'dataset_id': 'industry_productivity', 'variant_id': 'industry_productivity',
        'record_count': None,
        'recipe_path': str(recipe_path.relative_to(root)),
        'recipe_sha256': digest(recipe_path),
    })


def _compile_manual_xlsx(root, dataset, index):
    """Compile a manually authored workflow without inventing Agent ancestry."""
    reference = MANUAL_XLSX[dataset]
    asset = root / 'assets/cases' / dataset / 'source/datasets/xlsx' / dataset
    source_manifest = asset / 'manifest.json'
    manifest = read(source_manifest)
    if (manifest.get('origin_type') != 'manual_workflow_adaptation'
            or manifest.get('agent_generated') is not False
            or manifest.get('requires_llm') is not False):
        raise ValueError(f'{dataset}: manual origin markers missing')
    prompt_name = relative(manifest['prompt_template'])
    prompt_path = root / 'compatibility/prompts' / prompt_name
    if digest(prompt_path) != manifest.get('prompt_template_sha256'):
        raise ValueError(f'{dataset}: manifest does not bind the current shared prompt')
    # Frozen preparation artifact (produced by scripts/prepare_xlsx_selection.py
    # outside the measured loop).  It is copied only when the asset manifest
    # binds it; helpers refuse to run without a valid artifact, so an
    # unbound-but-compiled recipe fails fast at P02 instead of falling back to
    # in-loop aggregation.
    prepared_rel = 'input/prepared_selection.json'
    prepared_asset = asset / prepared_rel
    prepared_bound = prepared_rel in manifest.get('input_files', {})
    if prepared_bound:
        if not prepared_asset.is_file() or digest(prepared_asset) != manifest['input_files'][prepared_rel]:
            raise ValueError(f'{dataset}: prepared selection artifact missing or hash mismatch in assets')
    case = root / 'compatibility/cases' / dataset
    seed, source = case / 'seed', case / 'source'
    seed_input = seed / 'input'
    seed_input.mkdir(parents=True, exist_ok=True)
    for child in seed_input.iterdir():
        if child.is_file():
            child.unlink()
    shutil.copyfile(asset / 'input' / manifest['workbook'], seed_input / manifest['workbook'])
    if prepared_bound:
        shutil.copyfile(prepared_asset, seed_input / 'prepared_selection.json')
    write(seed_input / 'dataset_manifest.json', manifest)
    (seed_input / 'verify_office.py').write_text(_loop_verifier_wrapper())

    template = root / 'compatibility/cases' / reference / 'source'
    shutil.copytree(template, source, dirs_exist_ok=True,
                     ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
    # The template keeps its historical prompt. Manual datasets explicitly bind
    # the shared compatible prompt instead; never inherit the template's bytes.
    source_prompt = source / 'prompts' / prompt_name
    source_prompt.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(prompt_path, source_prompt)
    dataset_dir = source / 'datasets/xlsx' / dataset
    dataset_dir.mkdir(parents=True, exist_ok=True)
    for name in ('manifest.json', 'expected.json', 'complexity.json',
                 'content-summary.json', 'provenance.json'):
        shutil.copyfile(asset / name, dataset_dir / name)
    (dataset_dir / 'input').mkdir(parents=True, exist_ok=True)
    shutil.copyfile(asset / 'input' / manifest['workbook'],
                    dataset_dir / 'input' / manifest['workbook'])
    if prepared_bound:
        shutil.copyfile(prepared_asset, dataset_dir / 'input' / 'prepared_selection.json')
    _copy_current_verifier_package(root, source)
    runtime_wrapper = source / 'runtime/verify_office.py'
    runtime_wrapper.parent.mkdir(parents=True, exist_ok=True)
    runtime_wrapper.write_text(_loop_verifier_wrapper())

    workspace = ('/root/.openclaw/workspace/tool-modeling/OFFICE-XLSX-HR-METRO'
                 if dataset == 'hr_metro' else
                 '/root/.openclaw/workspace/tool-modeling/OFFICE-XLSX-RETAIL-ONLINE')
    helper = NEW_XLSX_HELPERS[dataset]
    output_book = XLSX_OUTPUTS[dataset]
    steps = [
        {'operation_id': f'XLSX-P0{i}-{phase}', 'source_tool_calls': []}
        for i, phase in enumerate(('inspect_prepare', 'build', 'process_publish', 'verify_deliver'), 1)]
    def add(stage, function, arguments, role):
        steps[stage - 1]['source_tool_calls'].append({
            'adaptation': 'manual-workflow-adaptation-v1', 'role': role,
            'tool_call': {'function_name': function, 'arguments': arguments}})
    def command(stage, text, role):
        add(stage, 'exec', {'command': f'cd {workspace} && {text}'}, role)
    add(1, 'read', {'path': '/root/.openclaw/skills/xlsx/SKILL.md'}, 'read_skill')
    command(1, 'python3 /opt/office-trace-replay/compatibility/tools/inspect_xlsx.py source', 'inspect_source')
    command(1, f'cp input/{shlex.quote(manifest["workbook"])} output/{output_book}', 'copy_official_input')
    command(1, 'python3 /opt/office-trace-replay/compatibility/tools/inspect_xlsx.py copy', 'inspect_copy')
    add(2, 'write', {'path': workspace + '/' + Path(helper).name,
                     'content': (root / helper).read_text()}, 'write_helper')
    command(2, f'python3 /opt/office-trace-replay/compatibility/tools/run_xlsx_helper_atomic.py {Path(helper).name} output/{output_book}', 'build_analysis')
    command(3, f'python3 /root/.openclaw/skills/xlsx/scripts/recalc.py output/{output_book} __DOCUMENT_RECALC_TIMEOUT__ > output/formula_recalc.json 2> output/formula_recalc.stderr.log', 'recalculate_in_place')
    command(3, 'python3 /opt/office-trace-replay/compatibility/tools/inspect_xlsx.py formula', 'inspect_formulas')
    command(3, 'python3 /opt/office-trace-replay/compatibility/tools/inspect_xlsx.py cached', 'inspect_cached_values')
    command(3, 'python3 /opt/office-trace-replay/compatibility/tools/export_xlsx.py', 'export_csv')
    command(4, 'python3 input/verify_office.py', 'loop_business_verification')
    command(4, 'python3 /opt/office-trace-replay/compatibility/tools/summary.py', 'write_summary')
    command(4, f'test -s output/{output_book} && test -s output/metrics.csv && test -s output/scenario.csv && test -s output/xlsx_enhancement_summary.json && test -s output/xlsx_business_diagnostics.json', 'check_deliverables')

    prompt = 'compatibility/prompts/xlsx.txt'
    provenance = dataset_dir / 'provenance.json'
    recipe = {
        'schema_version': 'scene-key-operations-v2', 'case_kind': 'xlsx',
        'case_id': manifest['runtime_case'], 'dataset_id': dataset, 'variant_id': dataset,
        'operation_count': 4,
        'workflow': {'success_path': [step['operation_id'] for step in steps]},
        'key_operations': steps, 'workspace': workspace,
        'seed_path': str(seed.relative_to(root)), 'source_path': str(source.relative_to(root)),
        'record_count': None, 'prepared_record_count': None,
        'operation_timeout': 900, 'recalc_timeout': 600, 'requires_llm': False,
        'origin_type': 'manual_workflow_adaptation', 'agent_generated': False,
        'classification': 'manual_workflow_adaptation',
        'manual_source': {
            'record_path': str(provenance.relative_to(root)), 'record_sha256': digest(provenance),
            'manifest_path': str((dataset_dir / 'manifest.json').relative_to(root)),
            'manifest_sha256': digest(dataset_dir / 'manifest.json'),
            'helper_path': helper, 'helper_sha256': digest(root / helper),
            'references': read(provenance)['references'],
        },
        'contract_sha256': index['contract_sha256'],
        'adaptations': [
            'Manual deterministic adaptation based on an existing workflow shape; no Agent run or canonical trajectory exists for this dataset.',
            'P01 copies the complete official input; P02 runs one dataset helper against the final output; P03 recalculates in place and exports; P04 runs the versioned loop verifier and summary.',
            'The original workbook sheets and rows are preserved; analysis is additive and narrowly scoped.',
        ],
        'prompt_path': prompt, 'prompt_sha256': digest(root / prompt),
        'input_hashes': {str(p.relative_to(seed)): digest(p)
                         for p in seed.rglob('*') if p.is_file()},
        'workflow_contract': XLSX_CONTRACT,
        'loop_verification_contract': manifest['loop_verification_contract'],
        'helper_hashes': {Path(helper).name: digest(root / helper)},
    }
    _align_xlsx_entrypoints(root, recipe)
    target = root / 'compatibility/recipes' / (dataset + '.json')
    write(target, recipe)
    index['entries'] = [e for e in index['entries'] if e['dataset_id'] != dataset]
    index['entries'].append({'dataset_id': dataset, 'variant_id': dataset,
                             'record_count': None,
                             'recipe_path': str(target.relative_to(root)),
                             'recipe_sha256': digest(target),
                             'origin_type': 'manual_workflow_adaptation'})
    print('compiled manual compatible', dataset)


def _loop_verifier_wrapper():
    return '''#!/usr/bin/env python3
"""Run the measured-loop verifier for this staged dataset."""
import os
from pathlib import Path
import sys

project = Path(os.environ.get('OFFICE_TRACE_PROJECT', '/workspace'))
sys.path.insert(0, str(project))
import office
from office_trace_bench.contracts import dataset_path, read_json
from office_trace_bench.verify import verify

if __name__ == '__main__':
    workspace = Path(__file__).resolve().parents[1]
    selection = read_json(workspace / 'input/dataset_manifest.json')
    report = verify(dataset_path(selection['kind'], selection['dataset_id']), workspace,
                    workspace / 'output/business_verification.json', scope='loop',
                    deferred_outputs=('xlsx_enhancement_summary.json',)
                    if selection['kind'] == 'xlsx' else ())
    print(report['status'])
    raise SystemExit(0 if report['status'] == 'success' else 1)
'''


def _freeze_base_sheet_dimensions(workbook_path, complexity_path):
    """Freeze declaration, serialized extent and populated extent at prepare time.

    A read-only worksheet's ``max_row`` can come from the OOXML dimension
    declaration rather than a serialized cell.  The complexity sidecar already
    records the serialized extent.  For modest inputs, one preparation-time
    value pass records the effective data extent and boundary cells; oversized
    workbooks use the hash-bound serialized extent without another full scan.
    """
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string

    complexity = read(complexity_path) if Path(complexity_path).is_file() else {}
    by_name = {item['name']: item for item in complexity.get('sheets', [])}
    modest = bool(complexity) and complexity.get('populated_cell_count', 0) <= 100_000
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    dimensions, details = {}, {}
    try:
        for sheet in wb.worksheets:
            declared = [sheet.max_row or 0, sheet.max_column or 0]
            item = by_name.get(sheet.title, {})
            serialized = [item.get('row_count', declared[0]),
                          item.get('column_count', declared[1])]
            boundary_cells = []
            if modest:
                max_row = max_col = 0
                row_boundary = []
                col_boundary = []
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None or cell.value == '':
                            continue
                        max_row = max(max_row, cell.row)
                        max_col = max(max_col, cell.column)
                        row_boundary.append(cell.coordinate)
                        col_boundary.append(cell.coordinate)
                effective = [max_row, max_col]
                if max_row:
                    boundary_cells = sorted({
                        address for address in row_boundary
                        if int(''.join(ch for ch in address if ch.isdigit())) == max_row
                    } | {
                        address for address in col_boundary
                        if column_index_from_string(
                            ''.join(ch for ch in address if ch.isalpha())) == max_col
                    })
            else:
                # The frozen complexity metadata is hash-bound to the source.
                # Do not scan a large workbook merely to refresh loop bounds.
                effective = serialized
            dimensions[sheet.title] = effective
            details[sheet.title] = {
                'declared_dimension': declared,
                'serialized_cell_range': serialized,
                'effective_data_range': effective,
                'effective_range_method': ('preparation-time populated-cell extent'
                                           if modest else ('hash-bound serialized extent; populated extent not rescanned'
                                                           if complexity else 'worksheet declaration fallback; populated extent unavailable')),
                'boundary_cells': boundary_cells,
            }
    finally:
        wb.close()
    return dimensions, details


def main(root=ROOT, datasets=None):
    root = Path(root)
    selective = datasets is not None
    datasets = list(DATASETS if datasets is None else datasets)
    index = {'schema_version': 'document-oe-compatible-suite-v1', 'contract_sha256': digest(root / 'compatibility/contract.json'),
             'legacy_lock_sha256': digest(root / 'legacy/document-oe/lock.json'), 'entries': []}
    previous_index = root / 'compatibility/suite.json'
    if selective and previous_index.is_file():
        old = read(previous_index)
        index['entries'] = preserved_suite_entries(old, datasets)
    for dataset in datasets:
        if dataset in MANUAL_XLSX:
            _compile_manual_xlsx(root, dataset, index)
            continue
        if dataset == 'industry_productivity':
            # This dataset is derived from an accepted canonical Agent run,
            # not an old recipes/<dataset>.json parent. Refresh it separately
            # after compiling the legacy-parent adaptations below.
            continue
        parent_path = root / 'recipes' / (dataset + '.json')
        parent = read(parent_path)
        old_seed = root / parent['seed_path']
        old_source = root / parent['source_path']
        original_manifest = read(old_seed / 'input/dataset_manifest.json')
        kind = original_manifest['kind']
        for count in ([1, 10] if kind == 'pdf' else [None]):
            variant = dataset + (f'-n{count}' if count else '')
            case = root / 'compatibility/cases' / variant
            seed = case / 'seed'
            source = case / 'source'
            seed.mkdir(parents=True, exist_ok=True)
            shutil.copytree(old_seed / 'input', seed / 'input', dirs_exist_ok=True)
            # A private source copy permits a versioned PDF subset contract,
            # without changing any canonical source/records/manifest bytes.
            shutil.copytree(old_source, source, dirs_exist_ok=True, ignore=shutil.ignore_patterns('__pycache__'))
            manifest = copy.deepcopy(original_manifest)
            adaptations = ['Use original native executor, original Skill tools and original runtime/timing boundaries.']
            if kind == 'pdf':
                record_path = seed / 'input' / manifest['records']
                records = read(record_path)
                if len(records[manifest['records_key']]) != 10:
                    raise ValueError('compatible PDF preparation requires the frozen ten-record fixture')
                ds = source / 'datasets/pdf' / dataset
                expected = read(ds / manifest['expected_file'])
                summary = copy.deepcopy(manifest['summary_contract'])
                summary.update(applicant_count=count, filled_pdf_count=count,
                               rendered_page_count=manifest['requirements']['page_count'] * (count + 1),
                               fill_script_invocations=count, render_script_invocations=count + 1,
                               prepared_record_count=10, field_values_count=10)
                expected['summary_values'] = summary
                write(ds / manifest['expected_file'], expected)
                manifest['expected_sha256'] = digest(ds / manifest['expected_file'])
                manifest['summary_contract'] = summary
                manifest['pdf_execution'] = {
                    'record_count': count,
                    'record_ids': [r['id'] for r in records[manifest['records_key']][:count]],
                    'selection': 'first records in frozen input order',
                }
                manifest['compatibility_variant'] = {'id': variant, 'method': 'prepare all records; execute selected prefix',
                    'parent_manifest_sha256': digest(old_seed / 'input/dataset_manifest.json')}
                # The original records, records contract, parent verifier and
                # canonical trace remain unchanged. Only this private source
                # dispatches to a verifier that distinguishes prepared/executed.
                _copy_current_verifier_package(root, source, pdf=True)
                manifest['compatibility_variant']['verifier_sha256'] = digest(source / 'office_trace_bench/compatible_pdf.py')
                wrapper = _loop_verifier_wrapper()
                (seed / 'input/verify_office.py').write_text(wrapper)
                runtime_wrapper = source / 'runtime/verify_office.py'
                runtime_wrapper.parent.mkdir(parents=True, exist_ok=True)
                runtime_wrapper.write_text(wrapper)
                write(ds / 'manifest.json', manifest)
                write(seed / 'input/dataset_manifest.json', manifest)
                adaptations += [f'Prepare all ten original mappings; explicitly fill/render first {count} records.',
                                'Same field-ID precheck for both new PDFs; complete mapping checks remain in final verification.',
                                'One subprocess batch driver alternates fill/render per selected record, matching the old driver pattern.',
                                'Private verifier checks all prepared mappings and exactly the selected PDFs/renders.']
            else:
                ds = source / 'datasets/xlsx' / dataset
                manifest['workload_contract'] = XLSX_CONTRACT
                # Private XLSX verifier: versioned semantic contract with strict
                # business checks and diagnostic-only format/serialization diffs.
                for name in ('verify.py', 'spreadsheet.py', 'contracts.py', 'workbooks.py'):
                    shutil.copyfile(root / 'compatibility/verification' / name,
                                    source / 'office_trace_bench' / name)
                # Explicit, checkable rule for empty-string normalization: only
                # the original sheets listed here may serialize empty-string
                # markers to empty cells; elsewhere it is a data change.
                workbook = root / 'compatibility/cases' / variant / 'seed/input' / manifest['workbook']
                manifest.setdefault('requirements', {})
                complexity = ds / 'complexity.json'
                expected_complexity = manifest.get('analysis_metadata', {}).get('complexity.json')
                if expected_complexity and digest(complexity) != expected_complexity:
                    raise ValueError(f'{dataset}: complexity metadata hash mismatch')
                dimensions, dimension_details = _freeze_base_sheet_dimensions(
                    workbook, complexity)
                manifest['requirements']['base_sheets'] = list(dimensions)
                manifest['requirements']['base_dimensions'] = dimensions
                manifest['requirements']['base_dimension_details'] = dimension_details
                manifest['requirements']['empty_strings_policy'] = {
                    'sheets': list(dimensions),
                    'rule': 'Empty-string cells on these original sheets are visual '
                            'fillers and may serialize to empty cells; empty-string '
                            'loss elsewhere is a data change.'}
                manifest['loop_verification_contract'] = 'document-oe-compatible-xlsx-loop-v1'
                wrapper = _loop_verifier_wrapper()
                (seed / 'input/verify_office.py').write_text(wrapper)
                runtime_wrapper = source / 'runtime/verify_office.py'
                runtime_wrapper.parent.mkdir(parents=True, exist_ok=True)
                runtime_wrapper.write_text(wrapper)
                adaptations += XLSX_ADAPTATIONS
                write(ds / 'manifest.json', manifest)
                write(seed / 'input/dataset_manifest.json', manifest)

            ws = parent['workspace']
            phases = ['inspect_prepare', 'build', 'process_publish', 'verify_deliver']
            steps = [{'operation_id': f'{kind.upper()}-P0{i}-{phase}', 'source_tool_calls': []}
                     for i, phase in enumerate(phases, 1)]
            def add(stage, function, arguments, role):
                steps[stage - 1]['source_tool_calls'].append({
                    'adaptation': 'document-oe-compatible-workflow-v2', 'role': role,
                    'tool_call': {'function_name': function, 'arguments': arguments}})
            def command(stage, cmd, role):
                add(stage, 'exec', {'command': f'cd {ws} && {cmd}'}, role)
            def write_helper(name, text):
                add(2, 'write', {'path': ws + '/' + name, 'content': text}, 'write_helper')

            add(1, 'read', {'path': f'/root/.openclaw/skills/{kind}/SKILL.md'}, 'read_skill')
            if kind == 'pdf':
                add(1, 'read', {'path': '/root/.openclaw/skills/pdf/forms.md'}, 'read_skill')
                add(1, 'read', {'path': ws + '/input/verify_office.py'}, 'read_verifier')
                tools = '/root/.openclaw/skills/pdf/scripts/'
                for name in ['check_fillable_fields.py', 'extract_form_field_info.py', 'fill_fillable_fields.py', 'convert_pdf_to_images.py']:
                    add(1, 'read', {'path': tools + name}, 'read_tool')
                command(1, 'mkdir -p output/field_values output/filled output/rendered/template', 'prepare')
                command(1, f'python3 {tools}check_fillable_fields.py input/{manifest["form"]} > output/check_fillable_fields.log 2>&1', 'check_fields')
                command(1, 'cat output/check_fillable_fields.log', 'inspect_log')
                command(1, f'python3 {tools}extract_form_field_info.py input/{manifest["form"]} output/form_field_info.json', 'extract_fields')
                add(1, 'read', {'path': ws + '/output/form_field_info.json'}, 'inspect_schema')
                command(1, f'python3 {tools}convert_pdf_to_images.py input/{manifest["form"]} output/rendered/template', 'render_template')
            else:
                command(1, 'cat input/verify_office.py', 'read_verifier')
                command(1, f'cp input/{manifest["workbook"]} output/{manifest["output_workbook"]}', 'copy_input')
                for mode in ['source', 'copy']:
                    command(1, f'python3 /opt/office-trace-replay/compatibility/tools/inspect_xlsx.py {mode}', 'inspect_workbook')

            for helper in HELPERS[dataset]:
                if kind == 'xlsx':
                    write_helper(helper, (root / NEW_XLSX_HELPERS[dataset]).read_text())
                else:
                    write_helper(helper, (old_seed / helper).read_text())
            if kind == 'xlsx':
                output_book = XLSX_OUTPUTS[dataset]
                command(2, f'python3 /opt/office-trace-replay/compatibility/tools/run_xlsx_helper_atomic.py {BUILD[dataset]} output/{output_book}', 'build')
            else:
                command(2, 'python3 ' + BUILD[dataset], 'build')
            if kind == 'pdf':
                command(2, "python3 -c \"import json,pathlib; ids={v['field_id'] for v in json.load(open('output/form_field_info.json'))}; rows=[json.load(open(p)) for p in pathlib.Path('output/field_values').glob('*.json')]; assert all(v['field_id'] in ids for r in rows for v in r); print('mapping checked')\"", 'validate_mapping')
                write_helper('run_batch_fill_render.py', (root / 'compatibility/tools/pdf_batch.py').read_text())
                command(3, 'python3 run_batch_fill_render.py', 'fill_render_batch')
                command(3, 'python3 /opt/office-trace-replay/compatibility/tools/summary.py', 'write_summary')
                command(4, 'python3 input/verify_office.py', 'semantic_verification')
                command(4, 'find output -type f | sort', 'check_deliverables')
            else:
                output_book = XLSX_OUTPUTS[dataset]
                command(3, f'python3 /root/.openclaw/skills/xlsx/scripts/recalc.py output/{output_book} __DOCUMENT_RECALC_TIMEOUT__ > output/formula_recalc.json 2> output/formula_recalc.stderr.log', 'recalculate')
                # HR build emits its index itself; other indexes were authored
                # in the parent trace. Only locations are restored, not answers.
                if dataset != 'hr':
                    for step in parent['key_operations']:
                        for item in step['source_tool_calls']:
                            c = item['tool_call']
                            if c['function_name'] == 'write' and c['arguments']['path'].endswith('/workbook_features.json'):
                                add(3, 'write', copy.deepcopy(c['arguments']), 'implementation_index')
                                break
                            if c['function_name'] == 'exec' and "cat > output/workbook_features.json <<'EOF'" in c['arguments']['command']:
                                text = c['arguments']['command'].split("cat > output/workbook_features.json <<'EOF'\n", 1)[1].split('\nEOF', 1)[0]
                                add(3, 'write', {'path': ws + '/output/workbook_features.json', 'content': text + '\n'}, 'implementation_index')
                for mode in ['formula', 'cached']:
                    command(3, f'python3 /opt/office-trace-replay/compatibility/tools/inspect_xlsx.py {mode}', 'inspect_workbook')
                command(3, 'cat output/formula_recalc.json', 'inspect_log')
                command(3, 'python3 /opt/office-trace-replay/compatibility/tools/export_xlsx.py', 'export_csv')
                command(4, 'python3 input/verify_office.py', 'semantic_verification')
                command(4, 'python3 /opt/office-trace-replay/compatibility/tools/summary.py', 'write_summary')
                command(4, f'test -s output/{output_book} && test -s output/metrics.csv && test -s output/scenario.csv && test -s output/xlsx_enhancement_summary.json', 'check_deliverables')

            recipe = {'schema_version': 'scene-key-operations-v2', 'case_kind': kind,
                      'case_id': manifest['runtime_case'], 'dataset_id': dataset, 'variant_id': variant,
                      'operation_count': 4, 'workflow': {'success_path': [s['operation_id'] for s in steps]},
                      'key_operations': steps, 'workspace': ws, 'seed_path': str(seed.relative_to(root)),
                      'source_path': str(source.relative_to(root)), 'record_count': count,
                      'prepared_record_count': 10 if kind == 'pdf' else None,
                      'operation_timeout': 900 if kind == 'xlsx' else 300,
                      'recalc_timeout': 600 if kind == 'xlsx' else 180,
                      'requires_llm': False, 'classification': CLASSIFICATION,
                      'parent': {'recipe_path': str(parent_path.relative_to(root)), 'recipe_sha256': digest(parent_path),
                                 **parent['source']},
                      'contract_sha256': index['contract_sha256'], 'adaptations': adaptations,
                      'prompt_path': f'compatibility/prompts/{kind}.txt',
                      'prompt_sha256': digest(root / f'compatibility/prompts/{kind}.txt'),
                      'input_hashes': {str(p.relative_to(seed)): digest(p) for p in (seed / 'input').rglob('*') if p.is_file()}}
            if kind == 'xlsx':
                recipe['workflow_contract'] = XLSX_CONTRACT
                helper_path = NEW_XLSX_HELPERS[dataset]
                recipe['loop_verification_contract'] = manifest['loop_verification_contract']
                recipe['helper_hashes'] = {BUILD[dataset]: digest(root / helper_path)}
            target = root / 'compatibility/recipes' / (variant + '.json')
            write(target, recipe)
            index['entries'].append({'dataset_id': dataset, 'variant_id': variant, 'record_count': count,
                                      'recipe_path': str(target.relative_to(root)), 'recipe_sha256': digest(target)})
            print('compiled compatible', variant, sum(len(s['source_tool_calls']) for s in steps))
    if 'industry_productivity' in datasets and (root / 'compatibility/recipes/industry_productivity.json').is_file():
        _update_industry_productivity_recipe(root, index)
    write(root / 'compatibility/suite.json', index)
    _refresh_public_suite(root)


def cli(argv=None):
    parser = argparse.ArgumentParser(
        description='Generate compatible recipes and source snapshots.',
        epilog='With no --dataset, all registered datasets are generated. '
               'Use --dataset repeatedly for a selective generation.')
    parser.add_argument('--dataset', action='append', choices=DATASETS, metavar='DATASET',
                        help='generate only this registered dataset (repeatable)')
    args = parser.parse_args(argv)
    main(datasets=args.dataset)
    return 0


if __name__ == '__main__':
    raise SystemExit(cli())
