
import openpyxl, time
t=time.time()
wb=openpyxl.load_workbook('input/egrid2023_data_rev2.xlsx', data_only=True, read_only=True)
print(wb.sheetnames)
ws=wb['US23']
for r in ws.iter_rows(min_row=1,max_row=3,max_col=167):
    pass
print('dims', ws.max_row, ws.max_column)
for addr in ['G3','M3','U3','DZ3','DX3','EA3','EB3','ED3','EE3']:
    print(addr, repr(ws[addr].value))
ws2=wb['ST23']
for r in [5,7,12,37,41,47]:
    print('ST23 row',r, [ws2.cell(row=r,column=c).value for c in (1,2,3,9)])
print('load', time.time()-t)
wb.close()
