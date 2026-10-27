
import openpyxl, zipfile
wb=openpyxl.load_workbook('output/energy_report.xlsx')
print(wb.sheetnames)
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column, 'charts',len(ws._charts),'merged',len(ws.merged_cells.ranges),'tables',len(ws.tables),'freeze',ws.freeze_panes)
ws=wb['Energy_Summary']
for a in ['A1','B5','B6','B7','B8','B9','B12','B13','B14','B15','B16','B20','B24','E20','E25']:
    c=ws[a]; print(a, c.data_type, repr(c.value), c.font.name, c.font.color.rgb if c.font.color else None)
print('cf ranges',[str(g.sqref) for g in ws.conditional_formatting])
print('comments',[c.coordinate for r in ws.iter_rows() for c in r if c.comment])
print('charts',[type(c).__name__ for c in ws._charts])
z=zipfile.ZipFile('output/energy_report.xlsx'); print([n for n in z.namelist() if 'table' in n or 'chart' in n or 'drawing' in n])
print('Contents D9', wb['Contents']['D9'].value, wb['Contents']['C9'].value)
