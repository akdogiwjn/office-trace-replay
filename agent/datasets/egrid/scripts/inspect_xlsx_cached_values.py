
import openpyxl, zipfile, re
z=zipfile.ZipFile('output/energy_report.xlsx')
c1=z.read('xl/charts/chart1.xml').decode(); c2=z.read('xl/charts/chart2.xml').decode()
print(re.findall(r'<c:f>(.*?)</c:f>', c1)); print(re.findall(r'<c:f>(.*?)</c:f>', c2))
ws=openpyxl.load_workbook('output/energy_report.xlsx',data_only=True)['Energy_Summary']
for r in range(5,17): print(r,[ws.cell(row=r,column=c).value for c in range(1,7)])
for r in range(20,26): print(r,[ws.cell(row=r,column=c).value for c in range(1,6)])
