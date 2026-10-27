"""Task-layout checks for the new Agent tasks; no prescribed business helper.

Reuse the existing scoped spreadsheet contract. Check the small declared Summary
through OOXML without another full openpyxl load, and never execute its code.
Original TLC/OPM verification does not enter this module.
"""
import copy
import re
import posixpath
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from .contracts import read_json, sha256, write_json

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
CHART_NS = {'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
REL_ID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'

# External audit adaptation only. The actual input bytes are verified before
# this source-hash-bound policy is applied; the Agent-facing manifest, local
# verifier, execution snapshot and frozen industry verifier remain untouched.
INDEPENDENT_EMPTY_STRING_POLICIES = {
    ('hr_metro', 'faf00eee127763c8c082d8fb862717ef2ad3a75f16f8975dc1e70935f27bcbb0'):
        ('MSA_M2025_dl',),
}


def independent_manifest(manifest):
    digest = manifest['input_files'].get('input/' + manifest['workbook'])
    sheets = INDEPENDENT_EMPTY_STRING_POLICIES.get((manifest['dataset_id'], digest))
    if not sheets:
        return manifest, None
    if not set(sheets).issubset(manifest['requirements']['base_sheets']):
        raise ValueError('external empty-string policy names an unknown source sheet')
    effective = copy.deepcopy(manifest)
    policy = effective['requirements'].setdefault('empty_strings_policy', {})
    policy['sheets'] = sorted(set(policy.get('sheets', [])) | set(sheets))
    return effective, {'version': 'hr-literal-empty-strings-v1',
                       'source_sha256': digest, 'sheets': list(sheets),
                       'equivalence': 'literal empty string to empty cell only',
                       'agent_manifest_modified': False}


def validate_task_manifest(manifest):
    req, layout = manifest['requirements'], manifest['task_layout']
    features = layout['features']
    if manifest.get('csv_mode') == 'copy_prepared':
        exports = manifest.get('publish_files', [])
        if (len(exports) != 2 or len({item['input'] for item in exports}) != 2
                or len({item['output'] for item in exports}) != 2):
            raise ValueError('copy-prepared tasks require exactly two distinct frozen CSV exports')
        for item in exports:
            from .contracts import safe_relative
            safe_relative(item['input'])
            safe_relative(item['output'])
            if (manifest['input_files'].get('input/' + item['input']) != item['sha256']
                    or item['output'] not in manifest['required_outputs']):
                raise ValueError('prepared CSV must bind its input hash and required output')
    if features['summary_sheet'] in req['base_sheets']:
        raise ValueError('task Summary must be a new worksheet')
    for section in ('metrics', 'charts'):
        if set(features[section]) != {item['id'] for item in req[section]}:
            raise ValueError('task layout inventory mismatch: ' + section)
    scene = features['scenario']
    if set(scene['assumptions']) != {row['name'] for row in req['scenario']['rows']}:
        raise ValueError('task scenario layout mismatch')
    if set(scene['projections']) != {p['id'] for p in req['scenario']['projections']}:
        raise ValueError('task projection layout mismatch')
    locations = [*features['metrics'].values(), scene['selector'],
                 *scene['projections'].values()]
    for row in req['scenario']['rows']:
        supplied = scene['assumptions'][row['name']]
        if set(supplied) != set(row['assumptions']):
            raise ValueError('task assumption layout mismatch')
        locations.extend(supplied.values())
    points = [(item['sheet'], item['cell']) for item in locations]
    if len(points) != len(set(points)):
        raise ValueError('duplicate task feature address')
    for sheet, address in points:
        if sheet != features['summary_sheet'] or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*', address):
            raise ValueError('invalid task feature address')
    if set(layout['metric_source_references']) != set(features['metrics']):
        raise ValueError('task metric source inventory mismatch')
    if set(layout['chart_requirements']) != set(features['charts']):
        raise ValueError('task chart specification inventory mismatch')
    for spec in layout['chart_requirements'].values():
        if spec['kind'] not in ('column', 'line', 'pie'):
            raise ValueError('unsupported task chart type')
        from openpyxl.utils.cell import range_boundaries
        for field in ('categories_range', 'values_range'):
            lo, bottom, hi, top = range_boundaries(spec[field])
            if lo != hi or top - bottom + 1 != len(spec['category_labels']):
                raise ValueError('task chart data range/label count mismatch')


def _range_operands(formula, *, value_dependencies=False):
    """Narrow only statically understood lookup return columns.

    VLOOKUP still depends on its key, search column, return column and other
    arguments. INDEX with a fixed column depends on that column plus its index
    arguments. Dynamic columns/complex tables retain the original conservative
    ranges. This is not a general Excel evaluator.
    """
    from openpyxl.formula import Tokenizer
    from openpyxl.formula.tokenizer import Token
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    items = Tokenizer(formula).items
    replacements = {}
    if value_dependencies:
        for start, token in enumerate(items):
            if (token.type, token.subtype) != (Token.FUNC, Token.OPEN):
                continue
            name = token.value[:-1].upper().removeprefix('_XLFN.')
            if name not in ('VLOOKUP', 'INDEX'):
                continue
            args, depth, begin = [], 0, start + 1
            for index in range(begin, len(items)):
                item = items[index]
                if item.subtype == Token.OPEN:
                    depth += 1
                elif item.subtype == Token.CLOSE:
                    if depth == 0:
                        args.append((begin, index))
                        break
                    depth -= 1
                elif (item.type, item.subtype) == (Token.SEP, Token.ARG) and depth == 0:
                    args.append((begin, index))
                    begin = index + 1
            else:
                continue
            if len(args) not in ((3, 4) if name == 'VLOOKUP' else (3,)):
                continue
            table_arg = args[1 if name == 'VLOOKUP' else 0]
            table = [i for i in range(*table_arg) if items[i].type != Token.WSPACE]
            column = [items[i] for i in range(*args[2]) if items[i].type != Token.WSPACE]
            if (len(table) != 1 or len(column) != 1
                    or (items[table[0]].type, items[table[0]].subtype) != (Token.OPERAND, Token.RANGE)
                    or (column[0].type, column[0].subtype) != (Token.OPERAND, Token.NUMBER)):
                continue
            value = items[table[0]].value
            prefix, area = (value.rsplit('!', 1) if '!' in value else ('', value))
            prefix = prefix + '!' if prefix else ''
            try:
                c1, r1, c2, r2 = range_boundaries(area.replace('$', ''))
                number = float(column[0].value)
                if (None in (c1, r1, c2, r2) or not number.is_integer()
                        or not 1 <= number <= c2 - c1 + 1):
                    continue
                selected = get_column_letter(c1 + int(number) - 1)
                ranges = [f'{prefix}{selected}{r1}:{selected}{r2}']
                if name == 'VLOOKUP':
                    first = get_column_letter(c1)
                    ranges.insert(0, f'{prefix}{first}{r1}:{first}{r2}')
                replacements[table[0]] = ranges
            except (ValueError, OverflowError):
                continue
    for index, item in enumerate(items):
        if (item.type, item.subtype) == (Token.OPERAND, Token.RANGE):
            yield from replacements.get(index, [item.value])


def reference_contains(formula, sheet, coordinate, current_sheet, *, value_dependencies=False):
    """Inspect range operands, not strings or substring matches such as A1/A10."""
    from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
    if not isinstance(formula, str) or not formula.startswith('='):
        return False
    for operand in _range_operands(formula, value_dependencies=value_dependencies):
        target, area = (operand.rsplit('!', 1) if '!' in operand
                        else (current_sheet, operand))
        target = target.strip("'").replace("''", "'")
        if target.casefold() != sheet.casefold():
            continue
        try:
            lo, bottom, hi, top = range_boundaries(area.replace('$', ''))
            row, col = coordinate_to_tuple(coordinate.replace('$', ''))
            if None not in (lo, bottom, hi, top) and lo <= col <= hi and bottom <= row <= top:
                return True
        except ValueError:
            continue  # Defined names/INDIRECT are not proof of this direct reference.
    return False


def source_reached(formulas, address, source_sheet, source_cell, summary_sheet, seen=None,
                   *, value_dependencies=False):
    """Follow only the small added Summary, never enumerate source-table cells."""
    seen = set() if seen is None else seen
    if address in seen:
        return False
    seen.add(address)
    formula = formulas.get(address)
    if reference_contains(formula, source_sheet, source_cell, summary_sheet,
                          value_dependencies=value_dependencies):
        return True
    return any(reference_contains(formula, summary_sheet, other, summary_sheet,
                                  value_dependencies=value_dependencies)
               and source_reached(formulas, other, source_sheet, source_cell, summary_sheet, seen,
                                  value_dependencies=value_dependencies)
               for other in formulas if other not in seen)


def _part_target(part, target):
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(
        posixpath.join(posixpath.dirname(part), target))


def _part_relationships(package, part):
    path = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
    if path not in package.namelist():
        return {}
    return {item.attrib['Id']: _part_target(part, item.attrib['Target'])
            for item in ET.fromstring(package.read(path))}


def _summary_charts(package, sheet_part, tree):
    drawing = tree.find('s:drawing', NS)
    if drawing is None:
        return []
    drawing_part = _part_relationships(package, sheet_part)[drawing.attrib[REL_ID]]
    chart_parts = _part_relationships(package, drawing_part)
    result = []
    for node in ET.fromstring(package.read(drawing_part)).findall('.//c:chart', CHART_NS):
        result.append(ET.fromstring(package.read(chart_parts[node.attrib[REL_ID]])))
    return result


def summary_formulas(path, sheet_name, label_addresses=()):
    with ZipFile(path) as package:
        workbook = ET.fromstring(package.read('xl/workbook.xml'))
        sheets = list(workbook.find('s:sheets', NS))
        names = [sheet.attrib['name'] for sheet in sheets]
        rels = {item.attrib['Id']: item.attrib['Target'] for item in ET.fromstring(
            package.read('xl/_rels/workbook.xml.rels'))}
        sheet = next(item for item in sheets if item.attrib['name'] == sheet_name)
        key = sheet.attrib[REL_ID]
        target = rels[key]
        target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        tree = ET.fromstring(package.read(target))
        formulas, labels, shared_labels = {}, {}, {}
        for cell in tree.findall('.//s:sheetData/s:row/s:c', NS):
            address = cell.attrib['r']
            formula = cell.find('s:f', NS)
            if formula is not None:
                formulas[address] = '=' + (formula.text or '')
            if address not in label_addresses:
                continue
            value = cell.find('s:v', NS)
            value = value.text if value is not None else None
            if cell.attrib.get('t') == 'inlineStr':
                value = ''.join(item.text or '' for item in cell.findall('.//s:t', NS))
            elif cell.attrib.get('t') == 's' and value is not None:
                shared_labels[address] = int(value)
            labels[address] = value
        if shared_labels:
            indices = set(shared_labels.values())
            strings = {}
            with package.open('xl/sharedStrings.xml') as stream:
                index = 0
                for _, item in ET.iterparse(stream, events=('end',)):
                    if item.tag != '{' + NS['s'] + '}si':
                        continue
                    if index in indices:
                        strings[index] = ''.join(t.text or '' for t in item.findall('.//s:t', NS))
                    item.clear()
                    if len(strings) == len(indices):
                        break
                    index += 1
            labels.update({address: strings[index] for address, index in shared_labels.items()})
        # LibreOffice may merge identical conditional-format rules into one
        # group covering multiple regions. Count regions with rules, as the
        # dataset-local Agent verifiers do, without expanding cells.
        from openpyxl.utils.cell import range_boundaries
        conditional_regions = {range_boundaries(area.replace('$', ''))
                               for group in tree.findall('s:conditionalFormatting', NS)
                               if group.findall('s:cfRule', NS)
                               for area in group.attrib.get('sqref', '').split()}
        return names, formulas, labels, _summary_charts(package, target, tree), len(conditional_regions)


def _range_identity(value):
    if not value or '!' not in value:
        return None
    sheet, area = value.rsplit('!', 1)
    return sheet.strip("'").replace("''", "'").casefold(), area.replace('$', '').upper()


def check_layout(manifest, workspace, checks):
    layout = manifest['task_layout']
    wanted = layout['features']
    output = Path(workspace) / 'output'
    actual = read_json(output / 'workbook_features.json')
    for section in ('summary_sheet', 'metrics', 'scenario', 'charts'):
        checks.add('task_layout:' + section, actual.get(section) == wanted[section])
    sheet = wanted['summary_sheet']
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    label_addresses = set()
    for spec in layout['chart_requirements'].values():
        col, bottom, _, top = range_boundaries(spec['categories_range'])
        label_addresses.update(get_column_letter(col) + str(row) for row in range(bottom, top + 1))
    names, formulas, labels, charts, conditional_regions = summary_formulas(
        output / manifest['output_workbook'], sheet, label_addresses)
    checks.add('format:conditional_rules',
               conditional_regions >= manifest['requirements']['minimum_conditional_formats'],
               {'covered_unique_regions': conditional_regions,
                'required': manifest['requirements']['minimum_conditional_formats']})
    checks.add('task_layout:sheet_order', names == [sheet, *manifest['requirements']['base_sheets']], names)
    for metric, references in layout['metric_source_references'].items():
        address = wanted['metrics'][metric]['cell']
        checks.add('task_source:' + metric, all(source_reached(
            formulas, address, ref['sheet'], ref['cell'], sheet) for ref in references))
    for address, references in layout.get('supporting_formula_references', {}).items():
        checks.add('task_formula:' + address, all(reference_contains(
            formulas.get(address), ref['sheet'], ref['cell'], sheet) for ref in references))
    for address, references in layout.get('forbidden_formula_references', {}).items():
        checks.add('task_independence:' + address, not any(source_reached(
            formulas, address, ref['sheet'], ref['cell'], sheet,
            value_dependencies=True) for ref in references))
    for ident, spec in layout['chart_requirements'].items():
        index = wanted['charts'][ident]['index']
        if index >= len(charts):
            checks.add('task_chart:' + ident, False, 'missing chart')
            continue
        chart = charts[index]
        tag = {'column': 'barChart', 'line': 'lineChart', 'pie': 'pieChart'}[spec['kind']]
        plot = chart.find('.//c:plotArea', CHART_NS)
        graph = plot.find('c:' + tag, CHART_NS) if plot is not None else None
        type_ok = graph is not None
        if type_ok and spec['kind'] == 'column':
            direction = graph.find('c:barDir', CHART_NS)
            type_ok = direction is not None and direction.attrib.get('val') == 'col'
        checks.add('task_chart_type:' + ident, type_ok)
        series = graph.findall('c:ser', CHART_NS) if graph is not None else []
        cats = [item.text for row in series for item in row.findall('c:cat/*/c:f', CHART_NS)]
        vals = [item.text for row in series for item in row.findall('c:val/c:numRef/c:f', CHART_NS)]
        checks.add('task_chart_categories:' + ident, len(series) == 1 and len(cats) == 1
                   and _range_identity(cats[0]) == (sheet.casefold(), spec['categories_range']))
        checks.add('task_chart_values:' + ident, len(series) == 1 and len(vals) == 1
                   and _range_identity(vals[0]) == (sheet.casefold(), spec['values_range']))
        col, bottom, _, top = range_boundaries(spec['categories_range'])
        actual_labels = [labels.get(get_column_letter(col) + str(row)) for row in range(bottom, top + 1)]
        checks.add('task_chart_labels:' + ident,
                   [str(value) for value in actual_labels] == [str(value) for value in spec['category_labels']],
                   actual_labels)


def verify_task(manifest_path, manifest, workspace, report_path, *, deferred_outputs, scope):
    from .industry_verification.verify import Checks
    from .industry_verification.spreadsheet import verify_semantic_xlsx, verify_semantic_xlsx_loop
    validate_task_manifest(manifest)
    class TaskChecks(Checks):
        defer_conditional_groups = True
        merged_payload_policy = None

        def add(self, name, ok, detail=None):
            if name.startswith('csv_matches_workbook:') and manifest.get('csv_mode') == 'copy_prepared':
                # These tasks publish immutable reference exports, as TLC does;
                # they do not export live cached CSVs. The semantic verifier's
                # published:* checks enforce both copies' exact frozen bytes.
                self.diag('not_applicable:cached_csv_export',
                          {'reason': 'Reference CSV copies are validated by published-file SHA256 checks.'})
                return
            if name == 'format:conditional_rules' and self.defer_conditional_groups:
                # The frozen semantic verifier counts serialized groups.
                # Preserve that observation; check_layout supplies the strict
                # region-based task check. Do not alter the historical module.
                self.diag('serialization:conditional_group_check',
                          {'legacy_group_threshold_met': bool(ok), 'detail': detail})
                return
            if self.merged_payload_policy is not None:
                from .merged_payload_audit import normalize_source_check
                normalized, result_detail = normalize_source_check(self.merged_payload_policy, name, ok, detail)
                if normalized and not ok:
                    self.diag('serialization:merged_non_anchor_cleanup:' + name[len('source:data:'):],
                              result_detail)
                ok, detail = normalized, result_detail
            super().add(name, ok, detail)

    checks = TaskChecks(manifest)
    checks.data.update(validation_scope=scope,
                       validation_contract=('dataset-specific-agent-loop-v1' if scope == 'loop'
                                            else 'dataset-specific-agent-independent-v2'))
    workspace = Path(workspace)
    try:
        if scope == 'independent':
            for relative, digest in manifest['input_files'].items():
                path = workspace / relative
                checks.add('input:' + relative, path.is_file() and sha256(path) == digest)
        if not checks.data['failures']:
            expected = read_json(Path(manifest_path).parent / manifest['expected_file'])
            verifier = verify_semantic_xlsx_loop if scope == 'loop' else verify_semantic_xlsx
            effective, policy = independent_manifest(manifest) if scope == 'independent' else (manifest, None)
            if policy:
                checks.data['external_audit_policy'] = policy
            if scope == 'independent' and (workspace / 'output' / manifest['output_workbook']).is_file():
                from .merged_payload_audit import merged_cleanup_policy
                merged_policy = merged_cleanup_policy(manifest, workspace, reference_contains)
                if merged_policy is not None:
                    checks.merged_payload_policy = merged_policy
                    checks.data['merged_payload_audit_policy'] = merged_policy
                    for sheet, proof in merged_policy['sheets'].items():
                        checks.add('source:merged_cleanup_guard:' + sheet, proof['ok'], proof)
            verifier(effective, expected, workspace, checks, deferred_outputs)
            checks.defer_conditional_groups = False
            if (workspace / 'output' / manifest['output_workbook']).is_file():
                check_layout(manifest, workspace, checks)
    except Exception as exc:
        checks.add('verification_exception', False, f'{type(exc).__name__}: {exc}')
    result = checks.finish()
    if report_path is not None:
        write_json(report_path, result)
    return result
