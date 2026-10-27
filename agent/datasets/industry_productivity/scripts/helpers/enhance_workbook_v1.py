#!/usr/bin/env python3
"""Enhance the BLS detailed-industries workbook with an executive summary sheet."""
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "input" / "labor-productivity-detailed-industries.xlsx"
DST = ROOT / "output" / "industry_productivity_enhanced.xlsx"

NAVY = "FF1F3864"
BLUE = "FF2F5597"
LIGHT_BLUE = "FFD9E2F3"
PALE_BLUE = "FFEEF3FA"
GRAY = "FFF2F2F2"
MID_GRAY = "FFD0CECE"
WHITE = "FFFFFFFF"
BORDER_C = "FFBFBFBF"
TEXT_DARK = "FF262626"

BLUE_FONT = "FF0000FF"
GREEN_FONT = "FF008000"
BLACK_FONT = "FF000000"

thin = Side(style="thin", color=BORDER_C)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MED = Side(style="medium", color=BLUE)


def style_range(ws, ref, font=None, fill=None, align=None, border=BOX, fmt=None):
    for row in ws[ref]:
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if align is not None:
                cell.alignment = align
            if border is not None:
                cell.border = border
            if fmt is not None:
                cell.number_format = fmt


def main():
    if not SRC.is_file():
        raise SystemExit(f"missing source {SRC}")
    shutil.copyfile(SRC, DST)

    wb = load_workbook(DST)
    wb.calculation.fullCalcOnLoad = True
    if "Productivity_Summary" in wb.sheetnames:
        del wb["Productivity_Summary"]
    ws = wb.create_sheet("Productivity_Summary", 0)
    ws.sheet_view.showGridLines = False

    arial = lambda **kw: Font(name="Arial", **kw)

    # ---------- title block ----------
    ws.merge_cells("A1:F1")
    ws["A1"] = "Industry Labor Productivity — Executive Summary"
    ws["A1"].font = arial(size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    ws["A2"] = ("Food manufacturing (NAICS 311), All workers, 2025. Index values (2017=100) are "
                "published BLS indexes, not physical hours or output levels.")
    ws["A2"].font = arial(size=9, italic=True, color="FF595959")
    ws["A2"].fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 26

    # ---------- KPI block ----------
    ws.merge_cells("A4:F4")
    ws["A4"] = "KEY PERFORMANCE INDICATORS — 2025 (INDEX, 2017=100)"
    ws["A4"].font = arial(size=11, bold=True, color=WHITE)
    ws["A4"].fill = PatternFill(start_color=BLUE, end_color=BLUE, fill_type="solid")
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[4].height = 22

    kpis = [
        (5, "Hours worked", "=Annual!AT489"),
        (6, "Real sectoral output", "=Annual!AT488"),
        (7, "Labor productivity", "=Annual!AT487"),
        (8, "Output per worker", "=Annual!AT490"),
        (9, "Source consistency status",
         '=IF(AND(Annual!AT489>0,Annual!AT490>0,ABS(Annual!AT488/Annual!AT489*100-Annual!AT487)<0.01),"PASS","REVIEW")'),
    ]
    for r, label, formula in kpis:
        ws.cell(row=r, column=1, value=label).font = arial(size=10, bold=True, color=TEXT_DARK)
        ws.cell(row=r, column=1).fill = PatternFill(start_color=GRAY, end_color=GRAY, fill_type="solid")
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.cell(row=r, column=1).border = BOX
        c = ws.cell(row=r, column=2, value=formula)
        c.font = arial(size=10, bold=True, color=GREEN_FONT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
        if r <= 8:
            c.number_format = "0.000"
        else:
            c.number_format = "General"
            c.fill = PatternFill(start_color=PALE_BLUE, end_color=PALE_BLUE, fill_type="solid")
        ws.row_dimensions[r].height = 20

    # ---------- scenario block ----------
    ws.merge_cells("A11:F11")
    ws["A11"] = "SCENARIO ANALYSIS — ILLUSTRATIVE SENSITIVITIES (NOT FORECASTS)"
    ws["A11"].font = arial(size=11, bold=True, color=WHITE)
    ws["A11"].fill = PatternFill(start_color=BLUE, end_color=BLUE, fill_type="solid")
    ws["A11"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[11].height = 22

    ws["A12"] = "Selected scenario"
    ws["A12"].font = arial(size=10, bold=True, color=TEXT_DARK)
    ws["A12"].fill = PatternFill(start_color=GRAY, end_color=GRAY, fill_type="solid")
    ws["A12"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws["A12"].border = BOX

    sel = ws["B12"]
    sel.value = "Base"
    sel.font = arial(size=10, bold=True, color=BLUE_FONT)
    sel.fill = PatternFill(start_color="FFFFF2CC", end_color="FFFFF2CC", fill_type="solid")
    sel.alignment = Alignment(horizontal="center", vertical="center")
    sel.border = BOX

    dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False)
    dv.promptTitle = "Scenario"
    dv.prompt = "Choose Base, Upside or Downside."
    dv.errorTitle = "Invalid scenario"
    dv.error = "Pick Base, Upside or Downside."
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add("B12")

    heads = ["Scenario", "Productivity multiplier", "Hours multiplier"]
    for j, h in enumerate(heads, start=4):
        c = ws.cell(row=12, column=j, value=h)
        c.font = arial(size=9, bold=True, color=WHITE)
        c.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[12].height = 30

    table = [("Base", 1.00, 1.00), ("Upside", 1.08, 1.05), ("Downside", 0.93, 0.95)]
    for i, (name, prod, hours) in enumerate(table, start=13):
        d = ws.cell(row=i, column=4, value=name)
        d.font = arial(size=10, color=TEXT_DARK)
        d.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        d.border = BOX
        if name == "Base":
            d.fill = PatternFill(start_color=PALE_BLUE, end_color=PALE_BLUE, fill_type="solid")
        for j, val in ((5, prod), (6, hours)):
            c = ws.cell(row=i, column=j, value=val)
            c.font = arial(size=10, color=BLUE_FONT)
            c.number_format = '0.00"x"'
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BOX
        ws.row_dimensions[i].height = 19

    local = [
        (13, "Productivity multiplier", "=VLOOKUP($B$12,$D$13:$F$15,2,FALSE)", '0.00"x"'),
        (14, "Hours multiplier", "=VLOOKUP($B$12,$D$13:$F$15,3,FALSE)", '0.00"x"'),
        (15, "Projected hours-worked index", "=B5*B14", "0.000"),
        (16, "Projected real-output index", "=B6*B13*B14", "0.000"),
    ]
    for r, label, formula, fmt in local:
        lc = ws.cell(row=r, column=1, value=label)
        lc.font = arial(size=10, bold=True, color=TEXT_DARK)
        lc.fill = PatternFill(start_color=GRAY, end_color=GRAY, fill_type="solid")
        lc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        lc.border = BOX
        c = ws.cell(row=r, column=2, value=formula)
        c.font = arial(size=10, bold=(r >= 15), color=BLACK_FONT)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
        if r >= 15:
            c.fill = PatternFill(start_color=PALE_BLUE, end_color=PALE_BLUE, fill_type="solid")
        ws.row_dimensions[r].height = 19

    # ---------- chart data ----------
    ws.merge_cells("A18:F18")
    ws["A18"] = "CHART DATA — LIVE LINKS TO THE ANNUAL WORKSHEET"
    ws["A18"].font = arial(size=11, bold=True, color=WHITE)
    ws["A18"].fill = PatternFill(start_color=BLUE, end_color=BLUE, fill_type="solid")
    ws["A18"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[18].height = 22

    for ref, text in (("A19", "Subindustry (NAICS)"), ("B19", "2025 productivity index"),
                      ("D19", "Year"), ("E19", "Food productivity index")):
        c = ws[ref]
        c.value = text
        c.font = arial(size=9, bold=True, color=WHITE)
        c.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[19].height = 30

    subs = [("Food", "=Annual!AT487"), ("Beverages and tobacco", "=Annual!AT510"),
            ("Textile mills", "=Annual!AT533"), ("Textile products", "=Annual!AT556"),
            ("Apparel", "=Annual!AT579")]
    for i, (label, formula) in enumerate(subs, start=20):
        a = ws.cell(row=i, column=1, value=label)
        a.font = arial(size=10, color=TEXT_DARK)
        a.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        a.border = BOX
        b = ws.cell(row=i, column=2, value=formula)
        b.font = arial(size=10, color=GREEN_FONT)
        b.number_format = "0.000"
        b.alignment = Alignment(horizontal="center", vertical="center")
        b.border = BOX
        ws.row_dimensions[i].height = 19

    trend = [("2021", "=Annual!AP487"), ("2022", "=Annual!AQ487"), ("2023", "=Annual!AR487"),
             ("2024", "=Annual!AS487"), ("2025", "=Annual!AT487")]
    for i, (label, formula) in enumerate(trend, start=20):
        d = ws.cell(row=i, column=4, value=label)
        d.font = arial(size=10, color=TEXT_DARK)
        d.alignment = Alignment(horizontal="center", vertical="center")
        d.number_format = "@"
        d.border = BOX
        e = ws.cell(row=i, column=5, value=formula)
        e.font = arial(size=10, color=GREEN_FONT)
        e.number_format = "0.000"
        e.alignment = Alignment(horizontal="center", vertical="center")
        e.border = BOX

    # ---------- notes ----------
    ws.merge_cells("A26:F26")
    ws["A26"] = ("Notes: Annual and MachineReadable are wide and long views of the same published "
                 "BLS series, not independent observations to aggregate. Scenario multipliers are "
                 "illustrative sensitivities, not forecasts.")
    ws["A26"].font = arial(size=8, italic=True, color="FF595959")
    ws["A26"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[26].height = 28

    # ---------- charts ----------
    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = "2025 Labor Productivity Index — Selected Manufacturing Subindustries"
    bar.y_axis.title = "Index (2017=100)"
    bar.x_axis.title = "NAICS"
    bar.height = 8.5
    bar.width = 17
    data = Reference(ws, min_col=2, min_row=19, max_row=24)
    cats = Reference(ws, min_col=1, min_row=20, max_row=24)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.legend = None
    bar.y_axis.majorGridlines = None
    ws.add_chart(bar, "G4")

    line = LineChart()
    line.style = 12
    line.title = "Food Manufacturing Labor Productivity Index, 2021–2025"
    line.y_axis.title = "Index (2017=100)"
    line.x_axis.title = "Year"
    line.height = 8.5
    line.width = 17
    data2 = Reference(ws, min_col=5, min_row=19, max_row=24)
    cats2 = Reference(ws, min_col=4, min_row=20, max_row=24)
    line.add_data(data2, titles_from_data=True)
    line.set_categories(cats2)
    line.legend = None
    line.y_axis.majorGridlines = None
    ws.add_chart(line, "G20")

    # ---------- conditional formatting ----------
    red_fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    green_fill = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
    amber_fill = PatternFill(start_color="FFFFEB9C", end_color="FFFFEB9C", fill_type="solid")
    ws.conditional_formatting.add("B5:B8", ColorScaleRule(
        start_type="num", start_value=80, start_color="FFF8696B",
        mid_type="num", mid_value=100, mid_color="FFFFEB84",
        end_type="num", end_value=120, end_color="FF63BE7B"))
    ws.conditional_formatting.add("B9", FormulaRule(
        formula=['$B$9="PASS"'], fill=green_fill,
        font=Font(name="Arial", bold=True, color="FF006100")))
    ws.conditional_formatting.add("B9", FormulaRule(
        formula=['$B$9="REVIEW"'], fill=red_fill,
        font=Font(name="Arial", bold=True, color="FF9C0006")))
    ws.conditional_formatting.add("E13:F15", ColorScaleRule(
        start_type="num", start_value=0.9, start_color="FFF8696B",
        mid_type="num", mid_value=1.0, mid_color="FFFFEB84",
        end_type="num", end_value=1.1, end_color="FF63BE7B"))
    ws.conditional_formatting.add("B15:B16", CellIsRule(
        operator="greaterThan", formula=["100"], fill=green_fill))

    # ---------- layout ----------
    for col, w in (("A", 46), ("B", 24), ("C", 3), ("D", 24), ("E", 24), ("F", 20)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"

    # ---------- comments ----------
    ws["B5"].comment = Comment(
        "Source: BLS Annual labor productivity and costs, Food manufacturing (NAICS 311), "
        "All workers, 2025 hours-worked index (2017=100), cell Annual!AT489. "
        "Index values are not physical hours.", "Analyst", width=360, height=110)
    ws["B7"].comment = Comment(
        "Source: Annual!AT487, Food manufacturing labor productivity index (2017=100). "
        "Consistency check in B9 compares output/hours with the published productivity index.",
        "Analyst", width=360, height=100)
    ws["B9"].comment = Comment(
        "PASS when Annual!AT489 and Annual!AT490 are positive AND "
        "ABS(Annual!AT488/Annual!AT489*100-Annual!AT487)<0.01; otherwise REVIEW. "
        "Tolerance covers published three-decimal rounding.", "Analyst", width=380, height=110)
    ws["B12"].comment = Comment(
        "Illustrative sensitivity selector, default Base. Names denote higher/lower assumptions, "
        "not forecasts. Dropdown: Base, Upside, Downside.", "Analyst", width=340, height=90)
    ws["E13"].comment = Comment(
        "Hardcoded editable scenario inputs (blue font). Base 1.00/1.00, Upside 1.08/1.05, "
        "Downside 0.93/0.95. Applied to 2025 indexes for illustration only.",
        "Analyst", width=360, height=100)
    ws["B20"].comment = Comment(
        "Chart values link to Annual!AT487, AT510, AT533, AT556 and AT579 "
        "(NAICS 311-315 labor productivity, 2025).", "Analyst", width=340, height=90)

    wb.save(DST)
    print("saved", DST, DST.stat().st_size)
    print("sheets", wb.sheetnames)
    print("charts", len(ws._charts))


if __name__ == "__main__":
    sys.exit(main())
