from openpyxl import load_workbook
wb = load_workbook('input/labor-productivity-detailed-industries.xlsx', read_only=False, data_only=True)
ws = wb['Annual']
print('row3:', [(c, ws.cell(row=3,column=c).value) for c in range(1,47)])
print('row4:', [(c, ws.cell(row=4,column=c).value) for c in range(1,47)])
for r in range(480, 500):
    print(r, [ws.cell(row=r,column=c).value for c in range(1,8)], 'AT=', ws.cell(row=r,column=46).value)
print('--- scan Food manufacturing NAICS 311')
for r in range(1, ws.max_row+1):
    v = ws.cell(row=r,column=2).value
    if str(v).strip() == '311':
        print(r, [ws.cell(row=r,column=c).value for c in range(1,8)])
print('--- other NAICS')
for tgt in ('312','313','314','315'):
    rows=[r for r in range(1, ws.max_row+1) if str(ws.cell(row=r,column=2).value).strip()==tgt]
    print(tgt, rows[:10])
