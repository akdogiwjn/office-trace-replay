from openpyxl import load_workbook
out = load_workbook('output/industry_productivity_enhanced.xlsx', data_only=True, read_only=False)
r = out['ReadMe']
for row in r.iter_rows(min_row=1,max_row=14,max_col=2):
    print([c.value for c in row], [c.hyperlink.target if c.hyperlink else None for c in row])
a = out['Annual']
for row in (487,488,489,490,510,533,556,579):
    print(row, a.cell(row=row,column=2).value, a.cell(row=row,column=3).value, a.cell(row=row,column=6).value, 'AT=',a.cell(row=row,column=46).value)
for rr in (491,492,493,510,533,556,579):
    print('check',rr, a.cell(row=rr,column=1).value)
s = out['Productivity_Summary']
print('charts', len(s._charts), [c.anchor._from.col if hasattr(c.anchor,'_from') else None for c in s._charts])
print('freeze', s.freeze_panes, 'gridlines', s.sheet_view.showGridLines)
for cell in ['B5','B6','B7','B8','B9','B12','B13','B14','B15','B16']:
    print(cell, s[cell].value, s[cell].font.name, s[cell].font.color.rgb if s[cell].font.color else None)
print('widths', {k:v.width for k,v in s.column_dimensions.items()})
