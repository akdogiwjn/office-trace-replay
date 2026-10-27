from openpyxl import load_workbook
wb = load_workbook('input/labor-productivity-detailed-industries.xlsx', read_only=False, data_only=True)
print(wb.sheetnames)
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column, 'charts', len(ws._charts), 'merged', len(ws.merged_cells.ranges))
ws = wb['Annual']
for r in range(1, 12):
    print(r, [ws.cell(row=r, column=c).value for c in range(1, 8)])
print('--- header rows around 480-495 col 40-46')
for r in range(478, 496):
    print(r, [ws.cell(row=r, column=c).value for c in range(36, 47)])
