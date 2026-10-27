#!/usr/bin/env python3
"""Dataset-local XLSX verification, adapted from the frozen TLC verifier.

One normal formula load, then one read-only cached load. No source workbook
load, full-file hash, recursive dependency walk, project import or dispatcher.
Deep source preservation is checked separately after the Agent completes.
"""
import argparse
import gc
import json
import math
from pathlib import Path

from openpyxl import load_workbook

CONFIG = {'dataset_id': 'hr_metro',
 'output_workbook': 'workforce_report.xlsx',
 'features': {'summary_sheet': 'Workforce_Summary',
              'metrics': {'total_employment': {'sheet': 'Workforce_Summary', 'cell': 'B5'},
                          'annual_mean_wage': {'sheet': 'Workforce_Summary', 'cell': 'B6'},
                          'hourly_median_wage': {'sheet': 'Workforce_Summary', 'cell': 'B7'},
                          'annual_median_wage': {'sheet': 'Workforce_Summary', 'cell': 'B8'},
                          'reconciliation_status': {'sheet': 'Workforce_Summary', 'cell': 'B9'}},
              'scenario': {'selector': {'sheet': 'Workforce_Summary', 'cell': 'B12'},
                           'assumptions': {'Base': {'wage_multiplier': {'sheet': 'Workforce_Summary',
                                                                        'cell': 'E13'},
                                                    'employment_multiplier': {'sheet': 'Workforce_Summary',
                                                                              'cell': 'F13'}},
                                           'Upside': {'wage_multiplier': {'sheet': 'Workforce_Summary',
                                                                          'cell': 'E14'},
                                                      'employment_multiplier': {'sheet': 'Workforce_Summary',
                                                                                'cell': 'F14'}},
                                           'Downside': {'wage_multiplier': {'sheet': 'Workforce_Summary',
                                                                            'cell': 'E15'},
                                                        'employment_multiplier': {'sheet': 'Workforce_Summary',
                                                                                  'cell': 'F15'}}},
                           'projections': {'projected_employment': {'sheet': 'Workforce_Summary',
                                                                    'cell': 'B15'},
                                           'projected_annual_mean_wage': {'sheet': 'Workforce_Summary',
                                                                          'cell': 'B16'}}},
              'charts': {'major_employment': {'sheet': 'Workforce_Summary', 'index': 0},
                         'major_mean_wage': {'sheet': 'Workforce_Summary', 'index': 1}}},
 'base_sheets': ['MSA_M2025_dl', 'Field Descriptions', 'UpdateTime', 'Filler'],
 'base_dimensions': {'MSA_M2025_dl': [150024, 32],
                     'Field Descriptions': [48, 3],
                     'UpdateTime': [2, 1],
                     'Filler': [0, 0]},
 'base_chart_counts': {'MSA_M2025_dl': 0, 'Field Descriptions': 0, 'UpdateTime': 0, 'Filler': 0},
 'formula_cells': ['B5',
                   'B6',
                   'B7',
                   'B8',
                   'B9',
                   'B13',
                   'B14',
                   'B15',
                   'B16',
                   'B20',
                   'B21',
                   'B22',
                   'B23',
                   'B24',
                   'E20',
                   'E21',
                   'E22',
                   'E23',
                   'E24'],
 'numeric_formula_cells': ['B5',
                           'B6',
                           'B7',
                           'B8',
                           'B13',
                           'B14',
                           'B15',
                           'B16',
                           'B20',
                           'B21',
                           'B22',
                           'B23',
                           'B24',
                           'E20',
                           'E21',
                           'E22',
                           'E23',
                           'E24'],
 'minimum_formula_count': 19,
 'minimum_comments': 2,
 'minimum_conditional_formats': 2,
 'chart_count': 2,
 'scenario_default': 'Base',
 'lookup_values': {'B13': 1, 'B14': 1},
 'minimum_bytes': 1,
 'metric_values': {'total_employment': 9492000,
                   'annual_mean_wage': 86310,
                   'hourly_median_wage': 29.53,
                   'annual_median_wage': 61430.0,
                   'reconciliation_status': 'PASS'},
 'projection_values': {'projected_employment': 9492000, 'projected_annual_mean_wage': 86310}}
ERRORS = {'#REF!', '#DIV/0!', '#VALUE!', '#N/A', '#NAME?', '#NUM!', '#NULL!'}


def same_number(actual, expected):
    if isinstance(expected, str):
        return actual == expected
    return (isinstance(actual, (int, float)) and not isinstance(actual, bool)
            and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=1e-8, abs_tol=1e-6))


def color_endswith(cell, rgb):
    color = cell.font.color
    return color is not None and color.type == 'rgb' and str(color.rgb).upper().endswith(rgb)


def verify_formula_workbook(path, add):
    book = load_workbook(path, data_only=False, read_only=False)
    try:
        summary = CONFIG['features']['summary_sheet']
        base = CONFIG['base_sheets']
        add('sheet_order', book.sheetnames == [summary, *base], book.sheetnames)
        for name in base:
            sheet = book[name] if name in book else None
            rows, cols = CONFIG['base_dimensions'][name]
            add('base_sheet_scale:' + name, sheet is not None
                and (sheet.max_row or 0) >= rows and (sheet.max_column or 0) >= cols)
        sheet = book[summary]
        addresses = CONFIG['formula_cells']
        formulas = {address: sheet[address].value for address in addresses}
        add('required_formulas', all(sheet[address].data_type == 'f' for address in addresses), formulas)
        metrics = CONFIG['features']['metrics']
        # As in TLC: cross-sheet presence, not recursive business dependency analysis.
        add('cross_sheet_formulas', all('!' in str(sheet[p['cell']].value) for p in metrics.values()))
        count = sum(cell.data_type == 'f' for row in sheet.iter_rows() for cell in row)
        add('formula_count', count >= CONFIG['minimum_formula_count'], count)
        selector = CONFIG['features']['scenario']['selector']['cell']
        add('scenario_default', sheet[selector].value == CONFIG['scenario_default'])
        add('scenario_data_validation', any(v.type == 'list' and selector in v.sqref
                                           for v in sheet.data_validations.dataValidation))
        add('summary_charts', len(sheet._charts) >= CONFIG['chart_count'], len(sheet._charts))
        add('base_charts_preserved', all(len(book[n]._charts) >= c
                                        for n, c in CONFIG['base_chart_counts'].items() if c))
        # LibreOffice may coalesce groups: count distinct covered ranges, not groups.
        ranges = {str(area) for group in sheet.conditional_formatting
                  if sheet.conditional_formatting[group] for area in group.sqref.ranges}
        add('conditional_formatting', len(ranges) >= CONFIG['minimum_conditional_formats'], len(ranges))
        comments = [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.comment]
        add('source_comments', len(comments) >= CONFIG['minimum_comments'], comments)
        add('freeze_panes', sheet.freeze_panes is not None)
        add('font_color_convention', color_endswith(sheet[selector], '0000FF')
            and all(color_endswith(sheet[p['cell']], '008000') for p in metrics.values())
            and all(color_endswith(sheet[address], '000000') for address in CONFIG['lookup_values']))
        add('no_external_links', not book._external_links)
    finally:
        book.close()
        del book
        gc.collect()


def verify_cached_values(path, add):
    book = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = book[CONFIG['features']['summary_sheet']]
        # One bounded scan of the small new Summary; never scan a large data sheet.
        values = {cell.coordinate: cell.value for row in sheet.iter_rows() for cell in row if cell.value is not None}
        add('cached_numeric_values', all(same_number(values.get(a), values.get(a))
                                       for a in CONFIG['numeric_formula_cells']))
        for ident, spec in CONFIG['features']['metrics'].items():
            add('cached_metric:' + ident, same_number(values.get(spec['cell']), CONFIG['metric_values'][ident]))
        for address, expected in CONFIG['lookup_values'].items():
            add('scenario_cached_multiplier:' + address, same_number(values.get(address), expected))
        for ident, spec in CONFIG['features']['scenario']['projections'].items():
            add('scenario_cached:' + ident, same_number(values.get(spec['cell']), CONFIG['projection_values'][ident]))
        errors = [a for a, v in values.items() if isinstance(v, str) and v in ERRORS]
        add('zero_cached_formula_errors', not errors, errors[:12])
    finally:
        book.close()


def verify_workspace(workspace, report_path=None, *, workbook_path=None, recalc_path=None):
    workspace = Path(workspace)
    output = workspace / 'output'
    path = Path(workbook_path) if workbook_path is not None else output / CONFIG['output_workbook']
    checks, failures = {}, []
    def add(name, ok, detail=None):
        checks[name] = {'ok': bool(ok), 'detail': detail}
        if not ok:
            failures.append(name)
    try:
        add('workbook_file', path.is_file() and path.stat().st_size >= CONFIG['minimum_bytes'])
        if not failures:
            verify_formula_workbook(path, add)
            verify_cached_values(path, add)
        recalc = json.loads((Path(recalc_path) if recalc_path is not None else output / 'formula_recalc.json').read_text())
        add('recalc_report', recalc.get('status') == 'success' and recalc.get('total_errors') == 0
            and int(recalc.get('total_formulas') or 0) >= CONFIG['minimum_formula_count'], recalc)
    except Exception as exc:
        add('verification_exception', False, f'{type(exc).__name__}: {exc}')
    result = {'dataset_id': CONFIG['dataset_id'], 'kind': 'xlsx',
              'validation_scope': 'loop', 'validation_contract': 'dataset-local-tlc-style-v1',
              'status': 'failed' if failures else 'success', 'checks': checks, 'failures': failures}
    if report_path is not None:
        Path(report_path).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('recalc_json', type=Path)
    parser.add_argument('report_json', type=Path)
    args = parser.parse_args()
    report = verify_workspace(Path(__file__).resolve().parents[1], args.report_json,
                              workbook_path=args.workbook, recalc_path=args.recalc_json)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report['status'] == 'success' else 1


if __name__ == '__main__':
    raise SystemExit(main())
