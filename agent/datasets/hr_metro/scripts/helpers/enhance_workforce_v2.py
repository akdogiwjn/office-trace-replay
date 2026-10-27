#!/usr/bin/env python3
"""Build Workforce_Summary on the original BLS OEWS May 2025 metro workbook (v2)."""
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "input" / "MSA_M2025_dl.xlsx"
OUT = ROOT / "output" / "workforce_report.xlsx"
SHEET = "Workforce_Summary"
DATA = "MSA_M2025_dl"

NAVY = "FF1F3864"
BLUE = "FF2F5597"
LIGHT = "FFD9E2F3"
PALE = "FFF2F2F2"
WHITE = "FFFFFFFF"
GREY = "FFD9D9D9"
BLACK = "FF000000"
GREEN = "FF008000"
INPUT_BLUE = "FF0000FF"
YELLOW = "FFFFF2CC"
SLATE = "FF44546A"


def solid(color):
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def arial(**kw):
    kw.setdefault("name", "Arial")
    kw.setdefault("color", BLACK)
    return Font(**kw)


THIN = Side(style="thin", color=GREY)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_range(ws, ref, *, font=None, fill=None, align=None, border=None, fmt=None):
    block = ws[ref]
    if not isinstance(block, tuple):
        block = ((block,),)
    elif block and not isinstance(block[0], tuple):
        block = (block,)
    for row in block:
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if align is not None:
                cell.alignment = align
            if border is not None:
                cell.border = border
            if fmt is not None:
                cell.number_format = fmt


def banner(ws, row, text):
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"] = text
    style_range(ws, f"A{row}:F{row}", font=arial(size=11, bold=True, color=WHITE),
                fill=solid(BLUE), align=Alignment(horizontal="left", vertical="center", indent=1))
    ws.row_dimensions[row].height = 20


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SRC, OUT)
    wb = load_workbook(OUT)

    ws = wb.create_sheet(SHEET, 0)
    ws.sheet_view.showGridLines = False

    for col, width in {"A": 44, "B": 18, "C": 22, "D": 44, "E": 18, "F": 22}.items():
        ws.column_dimensions[col].width = width
    for row, height in {1: 34, 2: 20, 18: 8, 19: 30}.items():
        ws.row_dimensions[row].height = height

    # ---- Title block -------------------------------------------------
    ws.merge_cells("A1:F1")
    ws["A1"] = "Metropolitan Workforce & Wage Dashboard"
    style_range(ws, "A1:F1", font=arial(size=16, bold=True, color=WHITE),
                fill=solid(NAVY), align=Alignment(horizontal="left", vertical="center", indent=1))

    ws.merge_cells("A2:F2")
    ws["A2"] = ("BLS Occupational Employment and Wage Statistics, May 2025   |   "
                "New York-Newark-Jersey City, NY-NJ (AREA 35620)   |   "
                "NAICS 000000, ownership 1235, all occupations")
    style_range(ws, "A2:F2", font=arial(size=10, italic=True, color=SLATE),
                fill=solid(LIGHT), align=Alignment(horizontal="left", vertical="center", indent=1))

    # ---- 1. Published benchmarks -------------------------------------
    banner(ws, 3, "PUBLISHED BENCHMARKS  —  ALL OCCUPATIONS")
    for addr, label in {"A4": "Metric", "B4": "Value", "C4": "Unit",
                        "D4": "Source reference"}.items():
        ws[addr] = label
    style_range(ws, "A4:D4", font=arial(size=10, bold=True, color=NAVY),
                fill=solid(LIGHT), border=BOX,
                align=Alignment(horizontal="left", vertical="center", indent=1))

    rows = [
        (5, "Employment, all occupations", f"={DATA}!L93966", "employees", "#,##0"),
        (6, "Annual mean wage, all occupations", f"={DATA}!S93966", "USD / year", "$#,##0"),
        (7, "Hourly median wage, all occupations", f"={DATA}!W93966", "USD / hour", "$#,##0.00"),
        (8, "Annual median wage, all occupations", f"={DATA}!AB93966", "USD / year", "$#,##0"),
    ]
    for row, label, formula, unit, fmt in rows:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=formula).number_format = fmt
        ws.cell(row=row, column=3, value=unit)
        ws.cell(row=row, column=4, value=formula[1:])
        style_range(ws, f"A{row}", font=arial(size=10, color="FF262626"), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", indent=1))
        style_range(ws, f"B{row}", font=arial(size=11, bold=True, color=GREEN), border=BOX,
                    align=Alignment(horizontal="right", vertical="center"), fmt=fmt)
        style_range(ws, f"C{row}", font=arial(size=9, color="FF595959"), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", indent=1))
        style_range(ws, f"D{row}", font=arial(size=9, color="FF808080"), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", indent=1))
        if row % 2 == 1:
            for col in "ACD":
                ws[f"{col}{row}"].fill = solid(PALE)

    ws["A9"] = "Source consistency status"
    ws["B9"] = (f'=IF(AND({DATA}!L93966>0,{DATA}!S93966>0,{DATA}!W93966>=0,'
                f'{DATA}!AB93966>=0),"PASS","REVIEW")')
    ws["C9"] = "status"
    ws["D9"] = f"{DATA}!L93966 / S93966 / W93966 / AB93966"
    style_range(ws, "A9", font=arial(size=10, color="FF262626"), border=BOX,
                align=Alignment(horizontal="left", vertical="center", indent=1))
    style_range(ws, "B9", font=arial(size=11, bold=True, color=GREEN), border=BOX,
                align=Alignment(horizontal="right", vertical="center"))
    style_range(ws, "C9", font=arial(size=9, color="FF595959"), border=BOX,
                align=Alignment(horizontal="left", vertical="center", indent=1))
    style_range(ws, "D9", font=arial(size=9, color="FF808080"), border=BOX,
                align=Alignment(horizontal="left", vertical="center", indent=1))

    ws["B5"].comment = Comment(
        "Source: BLS OEWS May 2025, MSA_M2025_dl!L93966.\n"
        "AREA 35620 (New York-Newark-Jersey City, NY-NJ), NAICS 000000, "
        "OWN_CODE 1235, All Occupations (row 93966).\n"
        "Published employment value; no payroll total is constructed from employment x wage.",
        "Workforce Analyst", width=320, height=130)
    ws["B9"].comment = Comment(
        "PASS only when L93966 and S93966 are positive and W93966 and AB93966 are nonnegative.\n"
        "All four references point directly at the single published all-occupations row; "
        "suppression markers are never replaced and the table is never rescanned.",
        "Workforce Analyst", width=320, height=110)

    # ---- 2. Scenario sensitivities -----------------------------------
    banner(ws, 10, "SCENARIO SENSITIVITIES  —  ILLUSTRATIVE, NOT FORECASTS")
    ws.merge_cells("A11:F11")
    ws["A11"] = ("Higher / lower assumptions for wage and employment only. "
                 "Select a scenario in B12; the editable assumptions sit in D13:F15.")
    style_range(ws, "A11:F11", font=arial(size=9, italic=True, color=SLATE),
                align=Alignment(horizontal="left", vertical="center", indent=1))

    ws["A12"] = "Selected scenario"
    ws["B12"] = "Base"
    style_range(ws, "A12", font=arial(size=10, bold=True, color=NAVY),
                align=Alignment(horizontal="left", vertical="center", indent=1))
    style_range(ws, "B12", font=arial(size=11, bold=True, color=INPUT_BLUE),
                fill=solid(YELLOW), border=BOX,
                align=Alignment(horizontal="center", vertical="center"))

    for addr, label in {"D12": "Scenario", "E12": "Wage multiplier",
                        "F12": "Employment multiplier"}.items():
        ws[addr] = label
    style_range(ws, "D12:F12", font=arial(size=9, bold=True, color=NAVY), fill=solid(LIGHT),
                border=BOX, align=Alignment(horizontal="center", vertical="center", wrap_text=True))

    scenarios = [("Base", 1.00, 1.00), ("Upside", 1.08, 1.05), ("Downside", 0.93, 0.95)]
    for i, (name, wage, emp) in enumerate(scenarios):
        r = 13 + i
        ws.cell(row=r, column=4, value=name)
        ws.cell(row=r, column=5, value=wage).number_format = "0.00"
        ws.cell(row=r, column=6, value=emp).number_format = "0.00"
        style_range(ws, f"D{r}", font=arial(size=10, color=INPUT_BLUE), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", indent=1))
        style_range(ws, f"E{r}:F{r}", font=arial(size=10, color=INPUT_BLUE), border=BOX,
                    align=Alignment(horizontal="center", vertical="center"))
        if i % 2 == 1:
            for col in "DEF":
                ws[f"{col}{r}"].fill = solid(PALE)

    ws["A13"] = "Wage multiplier applied"
    ws["B13"] = '=INDEX($E$13:$E$15,MATCH($B$12,$D$13:$D$15,0))'
    ws["A14"] = "Employment multiplier applied"
    ws["B14"] = '=INDEX($F$13:$F$15,MATCH($B$12,$D$13:$D$15,0))'
    ws["A15"] = "Projected employment (employees)"
    ws["B15"] = "=B5*B14"
    ws["A16"] = "Projected annual mean wage (USD/year)"
    ws["B16"] = "=B6*B13"

    for r in range(13, 17):
        style_range(ws, f"A{r}", font=arial(size=10, color="FF262626"), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", indent=1))
        style_range(ws, f"B{r}", font=arial(size=11, bold=True, color=BLACK), border=BOX,
                    align=Alignment(horizontal="right", vertical="center"))
        ws.row_dimensions[r].height = 20
    ws["B13"].number_format = "0.00"
    ws["B14"].number_format = "0.00"
    ws["B15"].number_format = "#,##0"
    ws["B16"].number_format = "$#,##0"
    for r in (15, 16):
        ws[f"A{r}"].font = arial(size=10, bold=True, color=NAVY)
        ws[f"B{r}"].fill = solid(LIGHT)
        ws[f"A{r}"].fill = solid(LIGHT)

    ws["B13"].comment = Comment(
        "Scenario multipliers are hardcoded illustrative sensitivities held in D13:F15 "
        "(Base 1.00/1.00, Upside 1.08/1.05, Downside 0.93/0.95).\n"
        "They denote higher/lower assumptions only, not forecasts; B13/B14 look them up from B12.",
        "Workforce Analyst", width=320, height=110)

    dv = DataValidation(type="list", formula1='"Base,Upside,Downside"', allow_blank=False,
                        showErrorMessage=True, errorTitle="Invalid scenario",
                        error="Choose Base, Upside or Downside.",
                        promptTitle="Scenario", prompt="Select Base, Upside or Downside.")
    dv.add(ws["B12"])
    ws.add_data_validation(dv)

    # ---- 3. Major occupation groups ----------------------------------
    banner(ws, 17, "MAJOR OCCUPATION GROUPS  —  EMPLOYMENT AND ANNUAL MEAN WAGE")
    ws["A19"] = "Major occupation group"
    ws["B19"] = "Employment (employees)"
    ws["D19"] = "Major occupation group"
    ws["E19"] = "Annual mean wage (USD/year)"
    style_range(ws, "A19:B19", font=arial(size=9, bold=True, color=NAVY), fill=solid(LIGHT),
                border=BOX, align=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    style_range(ws, "D19", font=arial(size=9, bold=True, color=NAVY), fill=solid(LIGHT),
                border=BOX, align=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    style_range(ws, "E19", font=arial(size=9, bold=True, color=NAVY), fill=solid(LIGHT),
                border=BOX, align=Alignment(horizontal="right", vertical="center", wrap_text=True, indent=1))

    groups = [
        ("Office and Administrative Support Occupations", "L94460", "S94460"),
        ("Healthcare Support Occupations", "L94336", "S94336"),
        ("Sales and Related Occupations", "L94437", "S94437"),
        ("Business and Financial Operations Occupations", "L94004", "S94004"),
        ("Management Occupations", "L93967", "S93967"),
    ]
    for i, (name, emp_ref, wage_ref) in enumerate(groups):
        r = 20 + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=f"={DATA}!{emp_ref}").number_format = "#,##0"
        ws.cell(row=r, column=4, value=name)
        ws.cell(row=r, column=5, value=f"={DATA}!{wage_ref}").number_format = "$#,##0"
        style_range(ws, f"A{r}", font=arial(size=10, color="FF262626"), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
        style_range(ws, f"B{r}", font=arial(size=10, bold=True, color=GREEN), border=BOX,
                    align=Alignment(horizontal="right", vertical="center"))
        style_range(ws, f"D{r}", font=arial(size=10, color="FF262626"), border=BOX,
                    align=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
        style_range(ws, f"E{r}", font=arial(size=10, bold=True, color=GREEN), border=BOX,
                    align=Alignment(horizontal="right", vertical="center"))
        if i % 2 == 1:
            for col in ("A", "B", "D", "E"):
                ws[f"{col}{r}"].fill = solid(PALE)
        ws.row_dimensions[r].height = 30

    ws["A20"].comment = Comment(
        "Explicitly selected major groups, in the supplied order. Each value is a live link to "
        "MSA_M2025_dl column L (TOT_EMP); wages link to column S (A_MEAN). "
        "No rescanning or ranking of the full table is used.",
        "Workforce Analyst", width=320, height=110)

    # ---- Conditional formatting on three distinct regions ------------
    ws.conditional_formatting.add(
        "B9", CellIsRule(operator="equal", formula=['"PASS"'],
                         fill=solid("FFC6EFCE"),
                         font=Font(name="Arial", color="FF006100", bold=True)))
    ws.conditional_formatting.add(
        "B9", CellIsRule(operator="equal", formula=['"REVIEW"'],
                         fill=solid("FFFFEB9C"),
                         font=Font(name="Arial", color="FF9C5700", bold=True)))
    ws.conditional_formatting.add(
        "B20:B24", DataBarRule(start_type="num", start_value=0, end_type="max",
                               color="FF2F5597", showValue=True))
    ws.conditional_formatting.add(
        "E20:E24", DataBarRule(start_type="num", start_value=0, end_type="max",
                               color="FF8EA9DB", showValue=True))
    ws.conditional_formatting.add(
        "B13:B14", ColorScaleRule(start_type="num", start_value=0.9, start_color="FFF8696B",
                                  mid_type="num", mid_value=1.0, mid_color="FFFFEB84",
                                  end_type="num", end_value=1.1, end_color="FF63BE7B"))

    # ---- Charts ------------------------------------------------------
    emp = BarChart()
    emp.type = "col"
    emp.style = 10
    emp.title = "Employment by major occupation group"
    emp.y_axis.title = "Employees"
    emp.y_axis.numFmt = "#,##0"
    emp.legend = None
    emp.height = 9.5
    emp.width = 21
    emp.gapWidth = 60
    emp.add_data(Reference(ws, min_col=2, min_row=19, max_row=24), titles_from_data=True)
    emp.set_categories(Reference(ws, min_col=1, min_row=20, max_row=24))
    ws.add_chart(emp, "A27")

    wage = BarChart()
    wage.type = "col"
    wage.style = 12
    wage.title = "Annual mean wage by major occupation group"
    wage.y_axis.title = "USD per year"
    wage.y_axis.numFmt = "$#,##0"
    wage.legend = None
    wage.height = 9.5
    wage.width = 21
    wage.gapWidth = 60
    wage.add_data(Reference(ws, min_col=5, min_row=19, max_row=24), titles_from_data=True)
    wage.set_categories(Reference(ws, min_col=4, min_row=20, max_row=24))
    ws.add_chart(wage, "A47")

    ws.freeze_panes = "A5"
    ws.sheet_properties.tabColor = NAVY
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(OUT)
    print(f"saved {OUT} size={OUT.stat().st_size}")


if __name__ == "__main__":
    main()
