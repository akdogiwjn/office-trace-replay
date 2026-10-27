from openpyxl import load_workbook
wb=load_workbook('output/workforce_report.xlsx')
print(wb.sheetnames)
ws=wb['Workforce_Summary']
for a in ['B5','B6','B7','B8','B9','B12','B13','B14','B15','B16','B20','B24','E20','E24']:
    c=ws[a]; col=c.font.color.rgb if c.font.color and c.font.color.type=='rgb' else None
    print(a, c.data_type, repr(c.value), c.number_format, c.font.name, col, c.font.sz)
print('dv', [(v.type, str(v.sqref), v.formula1) for v in ws.data_validations.dataValidation])
print('cf', [str(g.sqref) for g in ws.conditional_formatting])
print('charts', len(ws._charts), [ (c.title is not None) for c in ws._charts])
print('comments', [c.coordinate for r in ws.iter_rows() for c in r if c.comment])
print('freeze', ws.freeze_panes, 'merged', ws.merged_cells.ranges)
print('formula count', sum(c.data_type=='f' for r in ws.iter_rows() for c in r))
print('base', [(s.title, s.max_row, s.max_column, len(s._charts), len(s.merged_cells.ranges)) for s in wb.worksheets if s.title!='Workforce_Summary'])
print('ext', wb._external_links, 'names', list(wb.defined_names))
