from openpyxl import load_workbook
import time
t=time.time()
wb=load_workbook('input/MSA_M2025_dl.xlsx')
print('load', time.time()-t)
print(wb.sheetnames)
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column, 'charts', len(ws._charts), 'images', len(ws._images), 'merged', len(ws.merged_cells.ranges), 'freeze', ws.freeze_panes)
ws=wb['MSA_M2025_dl']
print('--- header rows 1..8 ---')
for r in range(1,9):
    print(r, [ws.cell(row=r,column=c).value for c in range(1,33)])
print('--- row 93966 ---')
for c in range(1,33):
    cell=ws.cell(row=93966,column=c)
    print(c, cell.coordinate, repr(cell.value), cell.number_format)
