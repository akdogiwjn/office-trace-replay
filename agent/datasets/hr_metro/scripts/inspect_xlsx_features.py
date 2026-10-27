from openpyxl import load_workbook
wb=load_workbook('input/MSA_M2025_dl.xlsx')
ws=wb['MSA_M2025_dl']
for r in [93966,93967,94004,94336,94437,94460]:
    print(r, ws.cell(r,10).value, 'L=',ws.cell(r,12).value, 'S=',ws.cell(r,19).value, 'W=',ws.cell(r,23).value, 'AB=',ws.cell(r,28).value)
print('defined names', list(wb.defined_names))
print('external links', wb._external_links)
for s in wb.worksheets:
    print(s.title, 'cf', len(s.conditional_formatting), 'dv', len(s.data_validations.dataValidation), 'tables', list(s.tables), 'autofilter', s.auto_filter.ref, 'print_area', s.print_area)
print('props', wb.properties.title, wb.properties.creator)
c=ws.cell(1,1); print('A1 font', c.font.name, c.font.sz, c.font.b, c.font.color, 'fill', c.fill.fill_type, c.fill.fgColor.rgb, 'align', c.alignment.horizontal)
c=ws.cell(2,12); print('L2 font', c.font.name, c.font.sz, c.number_format)
fd=wb['Field Descriptions']
print('merged', fd.merged_cells.ranges)
for r in range(1,10):
    print(r,[fd.cell(r,c).value for c in range(1,4)])
print('col widths MSA', {k:v.width for k,v in list(ws.column_dimensions.items())[:6]})
print('row heights', {k:v.height for k,v in list(ws.row_dimensions.items())[:4]})
print('default font', wb._fonts[0].name if wb._fonts else None, len(wb._fonts))
