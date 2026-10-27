#!/usr/bin/env python3
"""Inspect cached formula values in Executive_Summary."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook = load_workbook(args.workbook, data_only=True)
    try:
        worksheet = workbook["Executive_Summary"]
        for row in range(5, 17):
            cell = worksheet.cell(row=row, column=2)
            print(f"B{row}: {repr(cell.value)}")
        return 0
    finally:
        workbook.close()


if __name__ == "__main__":
    raise SystemExit(main())
