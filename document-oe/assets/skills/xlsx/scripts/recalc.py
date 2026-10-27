"""Recalculate XLSX formulas with LibreOffice and verify cached values."""

from __future__ import annotations

import json
import os
import platform
import subprocess
import sys
import tempfile
from pathlib import Path

from office.soffice import get_soffice_env
from openpyxl import load_workbook


MACRO_FILENAME = "Module1.xba"
REQUIRED_CACHED_CELLS = ("B5", "B6", "B7", "B8", "B9", "B13", "B14", "B15", "B16")
EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")

RECALCULATE_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""


def _failed(status: str, error: str, returncode: int | None = None, **details) -> dict:
    result = {"status": status, "error": error}
    if returncode is not None:
        result["returncode"] = returncode
    result.update(details)
    return result


def _profile_argument(profile: Path) -> str:
    return f"-env:UserInstallation={profile.resolve().as_uri()}"


def setup_libreoffice_macro(profile: Path) -> tuple[bool, str]:
    """Initialize an isolated LibreOffice profile and install the recalc macro."""
    profile.mkdir(parents=True, exist_ok=True)
    init_cmd = [
        "soffice",
        _profile_argument(profile),
        "--headless",
        "--norestore",
        "--terminate_after_init",
    ]
    try:
        initialized = subprocess.run(
            init_cmd,
            capture_output=True,
            text=True,
            timeout=60,
            env=get_soffice_env(),
        )
    except subprocess.TimeoutExpired:
        return False, "LibreOffice profile initialization exceeded 60 seconds"
    if initialized.returncode != 0:
        return False, initialized.stderr or initialized.stdout or "LibreOffice profile initialization failed"

    macro_dir = profile / "user" / "basic" / "Standard"
    macro_dir.mkdir(parents=True, exist_ok=True)
    try:
        (macro_dir / MACRO_FILENAME).write_text(RECALCULATE_MACRO, encoding="utf-8")
    except OSError as exc:
        return False, f"failed to install LibreOffice macro: {exc}"
    return True, ""


def _inspect_workbook(filename: Path) -> dict:
    error_details = {error: [] for error in EXCEL_ERRORS}
    total_errors = 0
    missing_cached_cells: list[str] = []

    cached = load_workbook(filename, data_only=True, read_only=True)
    try:
        has_executive_summary = "Executive_Summary" in cached.sheetnames
        if has_executive_summary:
            executive = cached["Executive_Summary"]
            missing_cached_cells = [cell for cell in REQUIRED_CACHED_CELLS if executive[cell].value is None]

        for sheet_name in cached.sheetnames:
            sheet = cached[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        for error in EXCEL_ERRORS:
                            if error in cell.value:
                                error_details[error].append(f"{sheet_name}!{cell.coordinate}")
                                total_errors += 1
                                break
    finally:
        cached.close()

    formulas = load_workbook(filename, data_only=False, read_only=True)
    try:
        formula_count = sum(
            1
            for sheet in formulas.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        )
    finally:
        formulas.close()

    if missing_cached_cells:
        return _failed(
            "failed",
            "formula cache values are missing",
            missing_cached_cells=missing_cached_cells,
            total_errors=total_errors,
            total_formulas=formula_count,
        )
    if total_errors:
        return _failed(
            "errors_found",
            "cached workbook contains formula errors",
            total_errors=total_errors,
            total_formulas=formula_count,
            error_summary={
                error: {"count": len(locations), "locations": locations[:20]}
                for error, locations in error_details.items()
                if locations
            },
        )
    return {
        "status": "success",
        "total_errors": 0,
        "total_formulas": formula_count,
        "error_summary": {},
        "cached_cells_verified": list(REQUIRED_CACHED_CELLS) if has_executive_summary else [],
    }


def recalc(filename: str, timeout: int = 30) -> dict:
    workbook = Path(filename).resolve()
    if not workbook.is_file():
        return _failed("failed", f"file does not exist: {workbook}")
    if timeout <= 0:
        return _failed("failed", "timeout must be positive")

    before_mtime = workbook.stat().st_mtime_ns
    with tempfile.TemporaryDirectory(prefix="document-bench-lo-profile-") as profile_name:
        profile = Path(profile_name)
        configured, error = setup_libreoffice_macro(profile)
        if not configured:
            return _failed("failed", error)

        command = [
            "soffice",
            _profile_argument(profile),
            "--headless",
            "--norestore",
            "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application",
            str(workbook),
        ]
        if platform.system() == "Linux":
            command = ["timeout", "--kill-after=30s", str(timeout), *command]

        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                text=True,
                env=get_soffice_env(),
            )
        except OSError as exc:
            return _failed("failed", f"failed to start LibreOffice: {exc}")

    if completed.returncode in {124, 137}:
        return _failed(
            "timeout",
            f"LibreOffice recalculation exceeded {timeout} seconds",
            returncode=completed.returncode,
        )
    if completed.returncode != 0:
        return _failed(
            "failed",
            completed.stderr or completed.stdout or "unknown LibreOffice error",
            returncode=completed.returncode,
        )

    after_mtime = workbook.stat().st_mtime_ns
    if after_mtime <= before_mtime:
        return _failed(
            "failed",
            "workbook was not saved after LibreOffice recalculation",
            before_mtime_ns=before_mtime,
            after_mtime_ns=after_mtime,
        )

    try:
        result = _inspect_workbook(workbook)
    except Exception as exc:
        return _failed("failed", f"failed to inspect recalculated workbook: {exc}")
    result["returncode"] = completed.returncode
    result["before_mtime_ns"] = before_mtime
    result["after_mtime_ns"] = after_mtime
    return result


def main() -> int:
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        return 2
    try:
        timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30
    except ValueError:
        print(json.dumps(_failed("failed", "timeout must be an integer"), indent=2))
        return 2

    result = recalc(sys.argv[1], timeout)
    print(json.dumps(result, indent=2))
    return 0 if result.get("status") == "success" else 1


if __name__ == "__main__":
    raise SystemExit(main())
