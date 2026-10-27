#!/usr/bin/env python3
"""Inspect XLSX charts, formatting, validations, and links for benchmark P01."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

try:
    # Replace/add this module later with the QEMU-specific marker implementation.
    from document_roi import roi_begin, roi_end
except ImportError:
    def roi_begin(_roi_id: str) -> None:
        pass

    def roi_end(_roi_id: str) -> None:
        pass


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook = None
    roi_begin("XLSX_P01_FEATURES")
    try:
        workbook = load_workbook(args.workbook, data_only=False)

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            if sheet._charts:
                print(f"\n=== {sheet_name} charts ({len(sheet._charts)}) ===")
                for index, chart in enumerate(sheet._charts):
                    print(f"  Chart {index}: type={chart.__class__.__name__}, title={chart.title}")

        reconciliation = workbook["Reconciliation"]
        for cell_ref in ["B10", "B14", "B5", "B24", "B25"]:
            cell = reconciliation[cell_ref]
            print(
                f"\nReconciliation!{cell_ref}: font={cell.font.name}/{cell.font.size}/"
                f"{cell.font.color}, fill={cell.fill}, numfmt={cell.number_format}"
            )

        hourly = workbook["Hourly_Summary"]
        for cell_ref in ["B34", "B29"]:
            cell = hourly[cell_ref]
            print(
                f"Hourly_Summary!{cell_ref}: font={cell.font.name}/{cell.font.size}/"
                f"{cell.font.color}, fill={cell.fill}, numfmt={cell.number_format}"
            )

        raw = workbook["Raw_Sample"]
        print(f"\nRaw_Sample: rows={raw.max_row}, cols={raw.max_column}")
        for column_index in range(1, min(raw.max_column + 1, 10)):
            value = raw.cell(row=1, column=column_index).value
            if value:
                print(f"  Col {column_index}={value}")

        print(f"\nExternal links: {len(workbook._external_links)}")
        print(f"Defined names: {workbook.defined_names}")

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            if sheet.data_validations and sheet.data_validations.dataValidation:
                print(f"\nData validations in {sheet_name}:")
                for validation in sheet.data_validations.dataValidation:
                    print(
                        f"  sqref={validation.sqref}, type={validation.type}, "
                        f"formula1={validation.formula1}"
                    )

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            if sheet.conditional_formatting:
                print(f"\nConditional formatting in {sheet_name}:")
                for formatting in sheet.conditional_formatting:
                    print(f"  range={formatting}, rules={len(formatting.rules)}")
        return 0
    finally:
        if workbook is not None:
            workbook.close()
        roi_end("XLSX_P01_FEATURES")


if __name__ == "__main__":
    raise SystemExit(main())
