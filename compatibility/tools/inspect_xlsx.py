"""Generic inspection with the old openpyxl load/inspect pattern."""
import json
from pathlib import Path
import sys
from openpyxl import load_workbook

ws = Path.cwd()
manifest = json.loads((ws / 'input/dataset_manifest.json').read_text())
path = ws / 'input' / manifest['workbook'] if sys.argv[1] == 'source' else ws / 'output' / manifest['output_workbook']
wb = load_workbook(path, data_only=sys.argv[1] == 'cached')
try:
    for sheet in wb:
        print(sheet.title, sheet.max_row, sheet.max_column, 'charts', len(sheet._charts),
              'merged', len(sheet.merged_cells.ranges))
    if sys.argv[1] in ('formula', 'cached'):
        f = json.loads((ws / 'output/workbook_features.json').read_text())
        for ident, ref in f['metrics'].items():
            print(ident, wb[ref['sheet']][ref['cell']].value)
finally:
    wb.close()
