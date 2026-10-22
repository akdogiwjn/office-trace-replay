"""Create a compact sales view while preserving every copied transaction row.

The full-period aggregation and the 2010 monthly totals are frozen in a
preparation artifact (``input/prepared_selection.json``) produced by
``scripts/prepare_xlsx_selection.py`` outside the measured loop.  This helper
never rescans the transaction table and writes no whole-table array formulas:
it checks the small artifact's structure (schema, dataset, workload version),
confirms the recorded header positions against the first worksheet row, copies
the frozen values into one small sheet explicitly labelled as derived data,
and links the KPIs, scenario projection and chart to that sheet.  Without a
structurally valid artifact it refuses to run rather than aggregating on the
spot.

The full input/artifact hash binding is a pre-execution preflight
responsibility (``replay.freeze.validate_dataset`` via the runner's
``validate_selected``, also exercised by ``--validate-only`` and the container
preflight); this helper deliberately does not re-hash the complete XLSX every
iteration.  The formal supported entry is the unified runner, which validates
that binding before any workload command runs.
"""
from datetime import date
from pathlib import Path
import json
import sys

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import column_index_from_string

TARGET = "retail_online_enhanced.xlsx"
PREPARED_NAME = "prepared_selection.json"
PREPARED_SCHEMA = "xlsx-prepared-selection-v1"
WORKLOAD_VERSION = "manual-workflow-adaptation-v2"
SOURCE_SHEET = "Online Retail"


def _text(value):
    return str(value).strip() if value is not None else ""


def load_prepared_selection(target_path):
    """Load the frozen preparation artifact and check its structure."""
    prepared_path = Path(target_path).parent.parent / 'input' / PREPARED_NAME
    if not prepared_path.is_file():
        raise FileNotFoundError(
            f'frozen preparation artifact missing: {prepared_path}; generate it first with '
            'scripts/prepare_xlsx_selection.py --dataset retail_online (approved preparation '
            'outside the measured loop)')
    data = json.loads(prepared_path.read_text(encoding='utf-8'))
    problems = []
    if data.get('schema_version') != PREPARED_SCHEMA:
        problems.append(f"unsupported artifact schema {data.get('schema_version')!r}")
    if data.get('dataset_id') != 'retail_online':
        problems.append(f"artifact dataset {data.get('dataset_id')!r} is not retail_online")
    if data.get('workload_version') != WORKLOAD_VERSION:
        problems.append(f"artifact workload version {data.get('workload_version')!r} is not {WORKLOAD_VERSION!r}")
    if problems:
        raise ValueError('prepared selection rejected: ' + '; '.join(problems))
    return data


def _verify_source_binding(source, prepared):
    """Confirm the recorded header positions against this workbook (first row only)."""
    selection = prepared['selection']
    if selection['sheet'] != SOURCE_SHEET or source.title != SOURCE_SHEET:
        raise ValueError('prepared selection worksheet binding mismatch')
    headers = {_text(cell.value): cell.column for cell in source[1]}
    for item in selection['columns'].values():
        column = headers.get(item['header'])
        if column is None or column_index_from_string(item['column']) != column:
            raise ValueError(
                f"header position for {item['header']!r} does not match the frozen prepared selection")


def _parse_range(area):
    first, _, last = area.partition(':')
    last = last or first
    def split(address):
        letters = ''.join(ch for ch in address if ch.isalpha()).upper()
        digits = ''.join(ch for ch in address if ch.isdigit())
        return letters, int(digits)
    first_col, first_row = split(first)
    last_col, last_row = split(last)
    return first_col, first_row, last_col, last_row


def build(path):
    path = Path(path)
    prepared = load_prepared_selection(path)
    wb = load_workbook(path)
    source = wb[SOURCE_SHEET]
    _verify_source_binding(source, prepared)
    selection = prepared['selection']
    totals = selection['totals']
    monthly = selection['monthly_2010_sales_gbp']
    layout = selection['derived_layout']
    derived_name = layout['sheet']
    anchors = layout['metric_anchors']
    chart_layout = layout['chart_series']['monthly_2010_sales']

    derived = wb.create_sheet(derived_name)
    derived['A1'] = 'Derived transaction totals (frozen preparation data)'
    derived['A1'].font = Font(bold=True)
    derived['A2'] = ('Derived once from the complete original Online Retail worksheet by '
                     'scripts/prepare_xlsx_selection.py outside the replay loop; not an '
                     'official source worksheet. Original transaction rows are unchanged.')
    derived['A2'].comment = Comment('Frozen preparation artifact values; not recomputed in the measured loop.', 'OpenAI')
    labels = {
        'gross_eligible_sales_gbp': 'Gross eligible sales (GBP)',
        'eligible_line_count': 'Eligible transaction lines',
        'eligible_quantity_units': 'Eligible quantity (items)',
    }
    for ident, anchor in anchors.items():
        anchor_col, anchor_row = _parse_range(anchor)[:2]
        derived.cell(anchor_row, 1, labels[ident])
        cell = derived.cell(anchor_row, column_index_from_string(anchor_col),
                            float(totals[ident]))
        cell.number_format = '£#,##0.00' if ident == 'gross_eligible_sales_gbp' else '#,##0'

    derived['A8'] = 'Preparation diagnostics'
    derived['A8'].font = Font(bold=True)
    derived['A9'] = 'Eligible lines missing CustomerID'
    derived['B9'] = totals['eligible_lines_missing_customer_id']
    derived['A10'] = 'Excluded rows (counts)'
    for offset, (reason, count) in enumerate(totals['exclusion_counts'].items()):
        derived.cell(11 + offset, 1, reason)
        derived.cell(11 + offset, 2, int(count))
    currency_row = 11 + len(totals['exclusion_counts'])
    derived.cell(currency_row, 1, 'Currency')
    derived.cell(currency_row, 2, selection['currency'])

    month_col, month_first_row, _, month_last_row = _parse_range(chart_layout['category_range'])
    value_col, value_first_row, _, value_last_row = _parse_range(chart_layout['value_range'])
    if month_first_row != value_first_row or month_last_row != value_last_row:
        raise ValueError('derived layout month and value ranges are inconsistent')
    if month_last_row - month_first_row + 1 != len(monthly):
        raise ValueError('derived layout month range does not cover twelve monthly values')
    header_row = month_first_row - 1
    if header_row <= currency_row:
        raise ValueError('derived layout month header overlaps the diagnostics block')
    derived.cell(header_row, column_index_from_string(month_col), 'Month (2010)')
    derived.cell(header_row, column_index_from_string(value_col), 'Eligible sales (GBP)')
    for offset, amount in enumerate(monthly):
        month_cell = derived.cell(month_first_row + offset,
                                  column_index_from_string(month_col), date(2010, offset + 1, 1))
        month_cell.number_format = 'mmm yyyy'
        value_cell = derived.cell(month_first_row + offset,
                                  column_index_from_string(value_col), float(amount))
        value_cell.number_format = '£#,##0.00'

    summary = wb.create_sheet('Sales Overview')
    summary.append(['Online Retail Sales Snapshot', 'Value'])
    metric_formulas = [
        ('gross_eligible_sales_gbp', 'Gross eligible sales (GBP)', '£#,##0.00'),
        ('eligible_line_count', 'Eligible transaction lines', '#,##0'),
        ('eligible_quantity_units', 'Eligible quantity (items)', '#,##0'),
    ]
    metrics = {}
    for row_no, (ident, label, fmt) in enumerate(metric_formulas, 2):
        summary.cell(row_no, 1, label)
        cell = summary.cell(row_no, 2, f"='{derived_name}'!{anchors[ident]}")
        cell.number_format = fmt
        cell.font = Font(color='008000', bold=True)
        cell.comment = Comment('Frozen full-period preparation value on the derived summary sheet; '
                               'repeated InvoiceNo values remain separate rows.', 'OpenAI')
        metrics[ident] = {'sheet': summary.title, 'cell': cell.coordinate}

    summary['A6'], summary['B6'] = 'Scenario', 'Base'
    summary['B6'].font = Font(color='0000FF', bold=True)
    validation = DataValidation(type='list', formula1='"Base,Higher,Lower"', allow_blank=False)
    summary.add_data_validation(validation)
    validation.add(summary['B6'])
    summary['A7'], summary['A8'], summary['A9'] = 'Base multiplier', 'Higher multiplier', 'Lower multiplier'
    summary['B7'], summary['B8'], summary['B9'] = 1.0, 1.05, 0.95
    for cell in ('B7', 'B8', 'B9'):
        summary[cell].font = Font(color='0000FF')
        summary[cell].comment = Comment('Illustrative assumption; not a forecast.', 'OpenAI')
    summary['A10'] = 'Projected eligible sales (GBP)'
    summary['B10'] = '=IF(B6="Base",B2*B7,IF(B6="Higher",B2*B8,B2*B9))'
    summary['B10'].number_format = '£#,##0.00'
    summary['B10'].font = Font(color='008000', bold=True)
    summary['B10'].comment = Comment('All-period eligible sales scaled by the selected scenario assumption.', 'OpenAI')
    summary['A11'] = 'Row policy'
    summary['C11'] = ('One source row is one transaction line; cancellation invoices and nonpositive '
                      'quantity/price are excluded from positive-sales metrics. Original rows remain unchanged.')
    summary['C11'].comment = Comment('Missing CustomerID is retained and counted diagnostically; no FX conversion is applied.', 'OpenAI')
    summary['A12'] = 'Derived data note'
    summary['C12'] = (f'KPI and chart values reference the {derived_name} sheet, which holds frozen '
                      'preparation data derived once from the complete original worksheet.')
    summary.freeze_panes = 'A2'
    summary.conditional_formatting.add('B2:B4', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='E2F0D9')))
    summary.conditional_formatting.add('B10', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='E2F0D9')))

    chart = BarChart()
    chart.type = 'col'
    chart.title = 'Eligible Transaction-Line Sales by Month (2010)'
    chart.y_axis.title = 'GBP'
    chart.x_axis.title = 'Month'
    chart.add_data(Reference(derived, min_col=column_index_from_string(value_col),
                             min_row=header_row, max_row=value_last_row), titles_from_data=True)
    chart.set_categories(Reference(derived, min_col=column_index_from_string(month_col),
                                   min_row=month_first_row, max_row=month_last_row))
    summary.add_chart(chart, 'A14')

    features = {
        'summary_sheet': summary.title, 'metrics': metrics,
        'derived_sheet': derived_name,
        'scenario': {'selector': {'sheet': summary.title, 'cell': 'B6'},
                     'assumptions': {
                         'Base': {'sales_multiplier': {'sheet': summary.title, 'cell': 'B7'}},
                         'Higher': {'sales_multiplier': {'sheet': summary.title, 'cell': 'B8'}},
                         'Lower': {'sales_multiplier': {'sheet': summary.title, 'cell': 'B9'}}},
                     'projections': {'projected_sales_gbp': {'sheet': summary.title, 'cell': 'B10'}}},
        'charts': {'monthly_2010_sales': {'sheet': summary.title, 'index': 0}},
    }
    output = path.parent / TARGET
    wb.save(output)
    wb.close()
    (path.parent / 'workbook_features.json').write_text(json.dumps(features, indent=2) + '\n')
    (path.parent / 'xlsx_business_diagnostics.json').write_text(json.dumps({
        'dataset_id': 'retail_online', 'derived_sheet': derived_name,
        'derived_data_note': ('KPI and chart values reference the frozen preparation artifact; the '
                              'complete original transaction worksheet is preserved unchanged.'),
        'eligible_line_count': totals['eligible_line_count'],
        'gross_eligible_sales_gbp': totals['gross_eligible_sales_gbp'],
        'eligible_quantity_units': totals['eligible_quantity_units'],
        'distinct_invoice_count': selection['diagnostics']['distinct_invoice_count'],
        'distinct_invoice_note': selection['diagnostics']['distinct_invoice_note'],
        'eligible_lines_missing_customer_id': totals['eligible_lines_missing_customer_id'],
        'exclusion_counts': totals['exclusion_counts'],
        'currency': selection['currency'], 'original_rows_preserved': True}, indent=2) + '\n')
    return output


if __name__ == '__main__':
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(TARGET)
    build(target)
