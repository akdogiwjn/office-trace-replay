"""Build the small metro HR analysis on the complete copied source workbook.

The metro selection (the unique all-occupations total row and the five largest
unsuppressed major occupation groups) is frozen in a preparation artifact
(``input/prepared_selection.json``) produced by
``scripts/prepare_xlsx_selection.py`` outside the measured loop.  This helper
never rescans the business table: it checks the small artifact's structure
(schema, dataset, workload version), confirms the recorded header positions
and spot-checks the recorded KPI/chart cell values against this workbook, then
references the recorded cells directly.  Without a structurally valid artifact
it refuses to run rather than aggregating on the spot.

The full input/artifact hash binding is a pre-execution preflight
responsibility (``replay.freeze.validate_dataset`` via the runner's
``validate_selected``, also exercised by ``--validate-only`` and the container
preflight); this helper deliberately does not re-hash the complete XLSX every
iteration.  The formal supported entry is the unified runner, which validates
that binding before any workload command runs.
"""
from pathlib import Path
import json
import math
import sys

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TARGET = "hr_metro_enhanced.xlsx"
PREPARED_NAME = "prepared_selection.json"
PREPARED_SCHEMA = "xlsx-prepared-selection-v1"
WORKLOAD_VERSION = "manual-workflow-adaptation-v2"
SOURCE_SHEET = "MSA_M2025_dl"


def _text(value):
    return str(value).strip() if value is not None else ""


def _same_number(left, right):
    if isinstance(left, bool) or isinstance(right, bool):
        return False
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return math.isclose(float(left), float(right), rel_tol=1e-9, abs_tol=1e-9)
    return left == right


def load_prepared_selection(target_path):
    """Load the frozen preparation artifact and check its structure."""
    prepared_path = Path(target_path).parent.parent / 'input' / PREPARED_NAME
    if not prepared_path.is_file():
        raise FileNotFoundError(
            f'frozen preparation artifact missing: {prepared_path}; generate it first with '
            'scripts/prepare_xlsx_selection.py --dataset hr_metro (approved preparation '
            'outside the measured loop)')
    data = json.loads(prepared_path.read_text(encoding='utf-8'))
    problems = []
    if data.get('schema_version') != PREPARED_SCHEMA:
        problems.append(f"unsupported artifact schema {data.get('schema_version')!r}")
    if data.get('dataset_id') != 'hr_metro':
        problems.append(f"artifact dataset {data.get('dataset_id')!r} is not hr_metro")
    if data.get('workload_version') != WORKLOAD_VERSION:
        problems.append(f"artifact workload version {data.get('workload_version')!r} is not {WORKLOAD_VERSION!r}")
    if problems:
        raise ValueError('prepared selection rejected: ' + '; '.join(problems))
    return data


def _verify_source_binding(source, prepared):
    """Confirm the recorded header positions and cell values against this workbook."""
    from openpyxl.utils import get_column_letter
    selection = prepared['selection']
    if selection['sheet'] != SOURCE_SHEET or source.title != SOURCE_SHEET:
        raise ValueError('prepared selection worksheet binding mismatch')
    headers = {_text(cell.value): cell.column for cell in source[1]}
    for name, letter in selection['columns'].items():
        column = headers.get(name)
        if column is None or get_column_letter(column) != letter:
            raise ValueError(
                f'header position for {name!r} does not match the frozen prepared selection')
    for metric, item in selection['total']['kpis'].items():
        if not _same_number(source[item['cell']].value, item['value']):
            raise ValueError(
                f"recorded total KPI {metric} at {item['cell']} does not match this workbook")
    for row in selection['chart_rows']:
        if not _same_number(source[row['employment_cell']].value, row['employment']):
            raise ValueError(
                f"recorded chart employment at {row['employment_cell']} does not match this workbook")


def build(path):
    path = Path(path)
    prepared = load_prepared_selection(path)
    wb = load_workbook(path)
    source = wb[SOURCE_SHEET]
    _verify_source_binding(source, prepared)
    selection = prepared['selection']
    scope = selection['scope']
    total = selection['total']

    summary = wb.create_sheet('Metro Summary')
    summary.append(['New York Metro Employment and Wage Snapshot', 'Value'])
    metrics = [
        ('total_employment', 'Employment, all occupations', total['kpis']['total_employment']['cell'], '#,##0'),
        ('annual_mean_wage', 'Annual mean wage (USD/year)', total['kpis']['annual_mean_wage']['cell'], '$#,##0.00'),
        ('hourly_median_wage', 'Hourly median wage (USD/hour)', total['kpis']['hourly_median_wage']['cell'], '$#,##0.00'),
    ]
    feature_metrics = {}
    for row_no, (ident, label, source_cell, fmt) in enumerate(metrics, 2):
        summary.cell(row_no, 1, label)
        cell = summary.cell(row_no, 2, f"='{SOURCE_SHEET}'!{source_cell}")
        cell.number_format = fmt
        cell.font = Font(color='008000', bold=True)
        cell.comment = Comment('Published value from the selected official MSA all-occupations total row.', 'OpenAI')
        feature_metrics[ident] = {'sheet': summary.title, 'cell': cell.coordinate}

    summary['A6'] = 'Scenario'
    summary['B6'] = 'Base'
    summary['B6'].font = Font(color='0000FF', bold=True)
    dv = DataValidation(type='list', formula1='"Base,Increase"', allow_blank=False)
    summary.add_data_validation(dv)
    dv.add(summary['B6'])
    summary['A7'], summary['A8'] = 'Base multiplier', 'Increase multiplier'
    summary['B7'], summary['B8'] = 1.0, 1.05
    summary['B7'].font = summary['B8'].font = Font(color='0000FF')
    summary['B7'].comment = Comment('Illustrative assumption; not a forecast.', 'OpenAI')
    summary['B8'].comment = Comment('Illustrative assumption; not a forecast.', 'OpenAI')
    summary['A9'] = 'Projected annual mean wage'
    summary['B9'] = '=IF(B6="Base",B3*B7,B3*B8)'
    summary['B9'].number_format = '$#,##0.00'
    summary['B9'].font = Font(color='008000', bold=True)
    summary['B9'].comment = Comment('Published annual mean wage multiplied by the selected scenario assumption.', 'OpenAI')
    summary['B10'] = 'Area scope'
    summary['C10'] = f"{scope['AREA']} — {scope['AREA_TITLE']}; all industries, all ownerships"
    summary['C10'].comment = Comment('Only the exact total row and major occupation groups are used.', 'OpenAI')
    summary.freeze_panes = 'A2'
    summary.conditional_formatting.add('B2:B3', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='E2F0D9')))
    summary.conditional_formatting.add('B9', CellIsRule(operator='greaterThan', formula=['0'], fill=PatternFill('solid', fgColor='E2F0D9')))

    chart_data = wb.create_sheet('Chart Data')
    chart_data.append(['Occupation group', 'Employment'])
    for out_row, row in enumerate(selection['chart_rows'], 2):
        chart_data.cell(out_row, 1, f"='{SOURCE_SHEET}'!{row['occupation_title_cell']}")
        chart_data.cell(out_row, 2, f"='{SOURCE_SHEET}'!{row['employment_cell']}")
    chart = BarChart()
    chart.type = 'bar'
    chart.title = 'Largest Major Occupation Groups by Employment'
    chart.y_axis.title = 'Occupation group'
    chart.x_axis.title = 'Employees'
    chart.add_data(Reference(chart_data, min_col=2, min_row=1, max_row=6), titles_from_data=True)
    chart.set_categories(Reference(chart_data, min_col=1, min_row=2, max_row=6))
    summary.add_chart(chart, 'A12')

    features = {
        'summary_sheet': summary.title,
        'metrics': feature_metrics,
        'scenario': {'selector': {'sheet': summary.title, 'cell': 'B6'},
                     'assumptions': {'Base': {'wage_multiplier': {'sheet': summary.title, 'cell': 'B7'}},
                                     'Increase': {'wage_multiplier': {'sheet': summary.title, 'cell': 'B8'}}},
                     'projections': {'projected_mean_wage': {'sheet': summary.title, 'cell': 'B9'}}},
        'charts': {'major_employment': {'sheet': summary.title, 'index': 0}},
    }
    output = path.parent / TARGET
    wb.save(output)
    wb.close()
    (path.parent / 'workbook_features.json').write_text(json.dumps(features, indent=2) + '\n')
    (path.parent / 'xlsx_business_diagnostics.json').write_text(json.dumps({
        'dataset_id': 'hr_metro', 'area_code': scope['AREA'], 'area_title': scope['AREA_TITLE'],
        'scope': 'cross-industry all-ownerships; official total KPIs and major occupation chart only',
        'selection_source': 'frozen preparation artifact (input/prepared_selection.json); no in-loop table rescan',
        'total_row': total['row'],
        'selected_chart_rows': [row['occupation_code'] for row in selection['chart_rows']],
        **selection['diagnostics']}, indent=2) + '\n')
    return output


if __name__ == '__main__':
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(TARGET)
    build(target)
